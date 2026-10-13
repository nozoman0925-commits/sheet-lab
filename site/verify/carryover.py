"""記事『エクセル家計簿で前月繰越を翌月に引き継ぐ方法』の数式をLibreOfficeで再計算して期待値と照合する。
使い方: python3 site/verify/carryover.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os
from openpyxl import Workbook, load_workbook

# 各月: (シート名, [(区分, 金額), ...])  区分は "収入" / "支出"
def build(opening, months):
    wb = Workbook(); wb.remove(wb.active)
    prev = None
    for name, rows in months:
        ws = wb.create_sheet(name)
        for r, t in enumerate(["前月繰越", "収入合計", "支出合計", "月末残高"], 2):
            ws.cell(r, 1, t)
        if prev is None: ws["B2"] = opening
        else: ws["B2"] = f"='{prev}'!B5"
        ws["B3"] = '=SUMIF(A8:A27,"収入",B8:B27)'
        ws["B4"] = '=SUMIF(A8:A27,"支出",B8:B27)'
        ws["B5"] = "=B2+B3-B4"
        for i, (k, v) in enumerate(rows, 8):
            ws.cell(i, 1, k); ws.cell(i, 2, v)
        prev = name
    return wb

def expected(opening, months):
    out, bal = [], opening
    for name, rows in months:
        inc = sum(v for k, v in rows if k == "収入"); exp = sum(v for k, v in rows if k == "支出")
        out.append((name, bal, inc, exp, bal + inc - exp)); bal = bal + inc - exp
    return out

SCEN = {
 "通常3か月": (50000, [("4月", [("収入", 250000), ("支出", 180000), ("支出", 42000)]),
                      ("5月", [("収入", 250000), ("支出", 200000), ("支出", 61234)]),
                      ("6月", [("収入", 255000), ("支出", 190000)])]),
 "赤字月をはさむ": (10000, [("4月", [("収入", 100000), ("支出", 130000)]),
                         ("5月", [("収入", 100000), ("支出", 50000)]),
                         ("6月", [("支出", 20000)])]),
 "入力が空の月": (0, [("4月", []), ("5月", [("収入", 1000)]), ("6月", [])]),
}

def main():
    ng = 0
    for label, (opening, months) in SCEN.items():
        d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); build(opening, months).save(src)
        out = os.path.join(d, "out"); os.mkdir(out)
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
        wb = load_workbook(os.path.join(out, "t.xlsx"), data_only=True)
        for name, b, i, e, c in expected(opening, months):
            ws = wb[name]; got = tuple(ws.cell(r, 2).value for r in range(2, 6))
            ok = got == (b, i, e, c); ng += not ok
            print(("OK " if ok else "NG ") + f"{label} {name}: 繰越{got[0]!r} 収入{got[1]!r} 支出{got[2]!r} 残高{got[3]!r} (期待 {b},{i},{e},{c})")
    print("不一致", ng, "件"); sys.exit(1 if ng else 0)
main()
