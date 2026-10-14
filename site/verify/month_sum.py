"""記事『日付から月別に集計する』の式をLibreOfficeで再計算して期待値と照合する。
使い方: python3 site/verify/month_sum.py"""
import subprocess, sys, tempfile, os, random, datetime as dt
from openpyxl import Workbook, load_workbook

random.seed(7)
def gen(n, years):
    rows = []
    for _ in range(n):
        y = random.choice(years); m = random.randint(1,12); d = random.randint(1,28)
        rows.append((dt.date(y,m,d), random.randint(100, 20000)))
    return rows
MAIN = gen(60, [2025, 2026])
# 境界: 月初・月末・うるう日
EDGE = [(dt.date(2026,4,1),100),(dt.date(2026,4,30),200),(dt.date(2026,5,1),400),(dt.date(2026,3,31),800),
        (dt.date(2024,2,29),1600),(dt.date(2025,2,28),3200),(dt.date(2026,12,31),6400),(dt.date(2027,1,1),12800)]

def main():
    wb = Workbook(); wb.remove(wb.active)
    for name, rows in (("main", MAIN), ("edge", EDGE)):
        ws = wb.create_sheet(name)
        for i,(d,v) in enumerate(rows,2):
            ws.cell(i,1,d); ws.cell(i,2,v)
        last = len(rows)+1
        for k, (y, m) in enumerate([(2026,m) for m in range(1,13)]+[(2025,2),(2024,2)]):
            r = k+1
            ws.cell(r,4,y); ws.cell(r,5,m)
            ws.cell(r,6,f'=SUMIFS(B2:B{last},A2:A{last},">="&DATE(D{r},E{r},1),A2:A{last},"<"&EDATE(DATE(D{r},E{r},1),1))')
            ws.cell(r,7,f'=SUMPRODUCT((YEAR(A2:A{last})=D{r})*(MONTH(A2:A{last})=E{r})*B2:B{last})')
            ws.cell(r,8,f'=SUMPRODUCT((MONTH(A2:A{last})=E{r})*B2:B{last})')  # 年を見ない版
    d = tempfile.mkdtemp(); src = os.path.join(d,"t.xlsx"); wb.save(src)
    out = os.path.join(d,"o"); os.mkdir(out)
    subprocess.run(["soffice","--headless","--convert-to","xlsx","--outdir",out,src], check=True, capture_output=True)
    res = load_workbook(os.path.join(out,"t.xlsx"), data_only=True)
    ng = 0; n = 0; diff_year = 0
    for name, rows in (("main", MAIN), ("edge", EDGE)):
        ws = res[name]
        for k, (y, m) in enumerate([(2026,m) for m in range(1,13)]+[(2025,2),(2024,2)]):
            r = k+1
            e = sum(v for d,v in rows if d.year==y and d.month==m)
            e_any = sum(v for d,v in rows if d.month==m)
            a,b,c = ws.cell(r,6).value, ws.cell(r,7).value, ws.cell(r,8).value
            n += 1
            ok = (a==e and b==e and c==e_any); ng += not ok
            if e != e_any: diff_year += 1
            if name=="edge" or not ok: print(("OK " if ok else "NG ")+f"{name} {y}年{m}月: SUMIFS {a} / SUMPRODUCT {b} / 年なし {c} (期待 {e} / 年なし {e_any})")
    print(f"集計 {n} 件、不一致 {ng} 件(年を見ないと結果が変わる月: {diff_year} 件)"); sys.exit(1 if ng else 0)
main()
