"""記事『外注先への発注額・請求額・支払済から、未払い・発注残・要確認をExcelで出す式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/subcontract_balance.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random
from openpyxl import Workbook, load_workbook

N = 1000
G = ('=IF(A{r}&B{r}&C{r}&D{r}&E{r}="","",IF(A{r}="","発注番号が空欄",IF(COUNTIF($A$2:$A${n},A{r})>1,"発注番号が重複",'
     'IF(B{r}="","外注先が空欄",IF(NOT(ISNUMBER(C{r})),"発注額が数字でない",IF(OR(C{r}<1,C{r}<>INT(C{r})),"発注額は1以上の整数",'
     'IF(AND(D{r}<>"",NOT(ISNUMBER(D{r}))),"請求額が数字でない",IF(AND(ISNUMBER(D{r}),OR(D{r}<0,D{r}<>INT(D{r}))),"請求額は0以上の整数",'
     'IF(AND(E{r}<>"",NOT(ISNUMBER(E{r}))),"支払済が数字でない",IF(AND(ISNUMBER(E{r}),OR(E{r}<0,E{r}<>INT(E{r}))),"支払済は0以上の整数",'
     'IF(N(E{r})>N(D{r}),"支払済が請求額を超えている","")))))))))))')
F = ('=IF(OR(G{r}<>"",A{r}&B{r}&C{r}&D{r}&E{r}=""),"",IF(N(D{r})>C{r},"発注額超過",IF(N(D{r})=0,"未請求",'
     'IF(N(E{r})<N(D{r}),"未払あり",IF(N(D{r})<C{r},"支払済(発注残あり)","完了")))))')
H = '=IF(F{r}="","",N(D{r})-N(E{r}))'
I = '=IF(F{r}="","",C{r}-N(D{r}))'
SUM = {"K5": '=SUM($H$2:$H${n})', "K6": '=SUMIF($I$2:$I${n},">0")',
       "K7": '=COUNTIF($F$2:$F${n},"発注額超過")',
       "K8": '=IF($K$3="","",SUMIFS($H$2:$H${n},$B$2:$B${n},$K$3))'}

def build(rows, vendor):
    wb = Workbook(); ws = wb.active; ws.title = "外注"
    for i, h in enumerate(["発注番号", "外注先", "発注額(税抜)", "請求額", "支払済", "状態", "チェック", "未払額", "発注残"], 1): ws.cell(1, i, h)
    ws["J3"] = "外注先を指定"; ws["K3"] = vendor
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1): ws.cell(r, c, v)
        ws.cell(r, 7, G.format(r=r, n=N)); ws.cell(r, 6, F.format(r=r))
        ws.cell(r, 8, H.format(r=r)); ws.cell(r, 9, I.format(r=r))
    for k, f in SUM.items(): ws[k] = f.format(n=N)
    return wb

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
empty = lambda v: v in (None, "")
key = lambda v: v.lower() if isinstance(v, str) else v
Z = lambda v: v if num(v) else 0

def expect(rows):
    out = []
    for a, b, c, d, e in rows:
        if all(empty(x) for x in (a, b, c, d, e)): out.append(("", "", "", "")); continue
        err = None
        if empty(a): err = "発注番号が空欄"
        elif sum(1 for x in rows if not empty(x[0]) and key(x[0]) == key(a)) > 1: err = "発注番号が重複"
        elif empty(b): err = "外注先が空欄"
        elif not num(c): err = "発注額が数字でない"
        elif c < 1 or c != int(c): err = "発注額は1以上の整数"
        elif not empty(d) and not num(d): err = "請求額が数字でない"
        elif num(d) and (d < 0 or d != int(d)): err = "請求額は0以上の整数"
        elif not empty(e) and not num(e): err = "支払済が数字でない"
        elif num(e) and (e < 0 or e != int(e)): err = "支払済は0以上の整数"
        elif Z(e) > Z(d): err = "支払済が請求額を超えている"
        if err: out.append(("", "", "", err)); continue
        d, e = Z(d), Z(e)
        st = ("発注額超過" if d > c else "未請求" if d == 0 else "未払あり" if e < d
              else "支払済(発注残あり)" if d < c else "完了")
        out.append((st, d - e, c - d, ""))
    return out

def exp_sum(rows, ex, vendor):
    unpaid = sum(x[1] for x in ex if x[1] != "")
    remain = sum(x[2] for x in ex if x[2] != "" and x[2] > 0)
    over = sum(1 for x in ex if x[0] == "発注額超過")
    v = "" if empty(vendor) else sum(x[1] for r, x in zip(rows, ex) if x[1] != "" and key(r[1]) == key(vendor))
    return (unpaid, remain, over, v)

def read(wb, n):
    ws = wb["外注"]; res = []
    for r in range(2, n + 2):
        v = ["" if ws.cell(r, c).value is None else ws.cell(r, c).value for c in (6, 8, 9, 7)]
        res.append(tuple(v))
    s = tuple("" if ws[k].value is None else ws[k].value for k in ("K5", "K6", "K7", "K8"))
    return res, s

def check(rows, vendor):
    e = expect(rows); es = exp_sum(rows, e, vendor)
    g, gs = read(run_lo(build(rows, vendor)), len(rows))
    return g == e and gs == es, g, e, gs, es

S = lambda name, rows, vendor=None: (name, rows, vendor)
SCEN = [
 S("まだ請求なし(未請求)", [("H-001", "A工業", 500000, None, None)]),
 S("請求ありで未払(未払あり)", [("H-001", "A工業", 500000, 500000, None)]),
 S("一部だけ支払済(未払あり・発注残なし)", [("H-001", "A工業", 500000, 500000, 200000)]),
 S("出来高の一部を請求して支払済(発注残あり)", [("H-001", "A工業", 500000, 300000, 300000)]),
 S("全額請求して全額支払済(完了)", [("H-001", "A工業", 500000, 500000, 500000)]),
 S("請求が発注額を超過(要確認)", [("H-001", "A工業", 500000, 550000, 100000)]),
 S("請求0を明示入力(未請求)", [("H-001", "A工業", 500000, 0, 0)]),
 S("3件の集計(未払・発注残・超過)", [("H-001", "A工業", 500000, 500000, 200000), ("H-002", "B電気", 300000, 100000, 100000), ("H-003", "A工業", 200000, 250000, 0), ("H-004", "C設備", 400000, None, None)]),
 S("外注先を指定して未払を出す", [("H-001", "A工業", 500000, 500000, 200000), ("H-002", "B電気", 300000, 100000, 0), ("H-003", "A工業", 200000, 200000, 50000)], "A工業"),
 S("指定した外注先が一覧にない", [("H-001", "A工業", 500000, 500000, 200000)], "Z組"),
 S("発注番号が空欄", [(None, "A工業", 500000, None, None)]),
 S("発注番号が重複(大文字小文字違い)", [("H-001", "A工業", 500000, None, None), ("h-001", "B電気", 300000, None, None)]),
 S("外注先が空欄", [("H-001", None, 500000, None, None)]),
 S("発注額が文字『50万』", [("H-001", "A工業", "50万", None, None)]),
 S("発注額が全角『５００，０００』", [("H-001", "A工業", "５００，０００", None, None)]),
 S("発注額が空欄", [("H-001", "A工業", None, None, None)]),
 S("発注額が0", [("H-001", "A工業", 0, None, None)]),
 S("発注額が小数", [("H-001", "A工業", 1000.5, None, None)]),
 S("請求額が文字『済』", [("H-001", "A工業", 500000, "済", None)]),
 S("請求額がマイナス", [("H-001", "A工業", 500000, -1000, None)]),
 S("請求額が小数", [("H-001", "A工業", 500000, 1000.5, None)]),
 S("支払済が文字『払い済み』", [("H-001", "A工業", 500000, 500000, "払い済み")]),
 S("支払済がマイナス", [("H-001", "A工業", 500000, 500000, -1)]),
 S("支払済が請求額を超える", [("H-001", "A工業", 500000, 300000, 400000)]),
 S("請求が空欄なのに支払済に入力", [("H-001", "A工業", 500000, None, 100000)]),
 S("空行をはさむ", [(None, None, None, None, None), ("H-001", "A工業", 500000, 500000, 0)]),
]

def main():
    ng = 0
    for name, rows, vendor in SCEN:
        ok, g, e, gs, es = check(rows, vendor); ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: {g} 集計={gs}")
        if not ok: print("   期待", e, es)
    rnd = random.Random(20261027); rows = []
    vendors = ["A工業", "B電気", "C設備", "D塗装", "E防水"]
    for i in range(500):
        c = rnd.randint(1, 100) * 10000
        d = rnd.choice([None, 0, c, c, rnd.randint(1, c // 10000) * 10000, c + 10000])
        e = None if d in (None, 0) else rnd.choice([0, d, d, rnd.randint(0, d // 10000) * 10000])
        a = f"H-{i:03d}"; b = rnd.choice(vendors); x = rnd.random()
        if x < 0.02: a = None
        elif x < 0.04: a = rows[i - 1][0] if i else a
        elif x < 0.06: b = None
        elif x < 0.08: c = rnd.choice(["50万", None, 0, 1000.5, "５０００００", -3])
        elif x < 0.10: d = rnd.choice(["済", -5, 1000.5, "１０００"])
        elif x < 0.12: e = rnd.choice(["払い済み", -5, 1000.5, (d if num(d) else 0) + 10000])
        rows.append((a, b, c, d, e))
    ok, g, e, gs, es = check(rows, "B電気"); ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: {len(rows)}行 全項目一致={ok} 集計={gs}")
    if not ok:
        for i, (x, y) in enumerate(zip(g, e)):
            if x != y: print("  行", i + 2, rows[i], x, y); break
        print("  集計", gs, es)
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
