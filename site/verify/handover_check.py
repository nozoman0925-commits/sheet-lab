"""記事『引渡日から1年点検・2年点検の予定日と期限超過をExcelで自動表示する式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/handover_check.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random, calendar
from datetime import datetime, date, timedelta
from openpyxl import Workbook, load_workbook

N = 1000
H = ('=IF(A{r}&B{r}&C{r}&D{r}="","",IF(A{r}="","工事番号が空欄",IF(COUNTIF($A$2:$A${n},A{r})>1,"工事番号が重複",'
     'IF(NOT(ISNUMBER($K$1)),"基準日が日付でない",IF(NOT(ISNUMBER(B{r})),"引渡日が日付でない",IF(OR(B{r}<1,B{r}>73050),"引渡日が範囲外",'
     'IF(AND(C{r}<>"",NOT(ISNUMBER(C{r}))),"1年点検済日が日付でない",IF(AND(D{r}<>"",NOT(ISNUMBER(D{r}))),"2年点検済日が日付でない",'
     'IF(AND(D{r}<>"",C{r}=""),"1年点検が空欄のまま2年点検済み",IF(AND(C{r}<>"",C{r}<B{r}),"1年点検済日が引渡日より前",'
     'IF(AND(D{r}<>"",D{r}<B{r}),"2年点検済日が引渡日より前","")))))))))))')
E = '=IF(OR(H{r}<>"",A{r}&B{r}&C{r}&D{r}=""),"",IF(C{r}="","1年点検",IF(D{r}="","2年点検","完了")))'
F = '=IF(OR(H{r}<>"",A{r}&B{r}&C{r}&D{r}="",E{r}="完了"),"",IF(E{r}="1年点検",EDATE(B{r},12),EDATE(B{r},24)))'
G = ('=IF(OR(H{r}<>"",A{r}&B{r}&C{r}&D{r}=""),"",IF(E{r}="完了","完了",IF(F{r}<$K$1,"期限超過",'
     'IF(F{r}-$K$1<=30,"30日以内","予定あり"))))')

def build(rows, base):
    wb = Workbook(); ws = wb.active; ws.title = "引渡"
    for i, h in enumerate(["工事番号", "引渡日", "1年点検済日", "2年点検済日", "次の点検", "次の点検予定日", "状態", "チェック"], 1): ws.cell(1, i, h)
    ws["J1"] = "基準日"; ws["K1"] = base; ws["K1"].number_format = "yyyy/mm/dd"
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
            if isinstance(v, datetime): ws.cell(r, c).number_format = "yyyy/mm/dd"
        ws.cell(r, 8, H.format(r=r, n=N)); ws.cell(r, 5, E.format(r=r))
        ws.cell(r, 6, F.format(r=r)); ws.cell(r, 6).number_format = "yyyy/mm/dd"; ws.cell(r, 7, G.format(r=r))
    return wb

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

def edate(d, m):
    y, mo = divmod(d.year * 12 + d.month - 1 + m, 12); mo += 1
    return date(y, mo, min(d.day, calendar.monthrange(y, mo)[1]))

isdt = lambda v: isinstance(v, datetime)
key = lambda v: v.lower() if isinstance(v, str) else v
empty = lambda v: v in (None, "")

def expect(rows, base):
    out = []
    for a, b, c, d in rows:
        if all(empty(x) for x in (a, b, c, d)): out.append(("", "", "", "")); continue
        err = None
        if empty(a): err = "工事番号が空欄"
        elif sum(1 for x in rows if not empty(x[0]) and key(x[0]) == key(a)) > 1: err = "工事番号が重複"
        elif not isdt(base): err = "基準日が日付でない"
        elif b == 0 and not isinstance(b, bool) and not isinstance(b, str): err = "引渡日が範囲外"
        elif not isdt(b): err = "引渡日が日付でない"
        elif not empty(c) and not isdt(c): err = "1年点検済日が日付でない"
        elif not empty(d) and not isdt(d): err = "2年点検済日が日付でない"
        elif not empty(d) and empty(c): err = "1年点検が空欄のまま2年点検済み"
        elif not empty(c) and c < b: err = "1年点検済日が引渡日より前"
        elif not empty(d) and d < b: err = "2年点検済日が引渡日より前"
        if err: out.append(("", "", "", err)); continue
        if empty(c): nxt, due = "1年点検", edate(b.date(), 12)
        elif empty(d): nxt, due = "2年点検", edate(b.date(), 24)
        else: out.append(("完了", "", "完了", "")); continue
        diff = (due - base.date()).days
        st = "期限超過" if diff < 0 else "30日以内" if diff <= 30 else "予定あり"
        out.append((nxt, due, st, ""))
    return out

def read(wb, n):
    ws = wb["引渡"]; res = []
    for r in range(2, n + 2):
        v = [ws.cell(r, c).value for c in (5, 6, 7, 8)]
        v = ["" if x is None else x for x in v]
        if isinstance(v[1], datetime): v[1] = v[1].date()
        res.append(tuple(v))
    return res

def check(rows, base):
    e = expect(rows, base); g = read(run_lo(build(rows, base)), len(rows))
    return g == e, g, e

D = lambda y, m, d: datetime(y, m, d)
BASE = D(2026, 10, 26)
S = lambda name, rows, base=BASE: (name, rows, base)
SCEN = [
 S("引渡から10か月・未点検(予定あり)", [("A-001", D(2025,12,1), None, None)]),
 S("1年点検まで25日(30日以内)", [("A-001", D(2025,11,20), None, None)]),
 S("1年点検の予定日を過ぎた(期限超過)", [("A-001", D(2025,9,1), None, None)]),
 S("1年点検済・2年点検へ", [("A-001", D(2025,9,1), D(2026,9,5), None)]),
 S("両方済み(完了)", [("A-001", D(2024,9,1), D(2025,9,3), D(2026,9,2))]),
 S("うるう日の引渡(2024/2/29→2025/2/28)", [("A-001", D(2024,2,29), None, None), ("A-002", D(2024,2,29), D(2025,3,1), None)]),
 S("月末の引渡(2025/10/31→2026/10/31)", [("A-001", D(2025,10,31), None, None)]),
 S("予定日が基準日と同じ日(30日以内)", [("A-001", D(2025,10,26), None, None)]),
 S("予定日が基準日の前日(期限超過)", [("A-001", D(2025,10,25), None, None)]),
 S("予定日が基準日の30日後(30日以内)", [("A-001", D(2025,11,25), None, None)]),
 S("予定日が基準日の31日後(予定あり)", [("A-001", D(2025,11,26), None, None)]),
 S("引渡日が文字『2025/9/1』", [("A-001", "2025/9/1", None, None)]),
 S("引渡日が空欄", [("A-001", None, None, None)]),
 S("工事番号が空欄", [(None, D(2025,9,1), None, None)]),
 S("工事番号が重複(大文字小文字違い)", [("A-001", D(2025,9,1), None, None), ("a-001", D(2025,9,1), None, None)]),
 S("2年点検だけ入力(1年が空欄)", [("A-001", D(2024,9,1), None, D(2026,9,2))]),
 S("1年点検済日が引渡日より前(日付の逆転)", [("A-001", D(2025,9,1), D(2025,8,1), None)]),
 S("2年点検済日が引渡日より前", [("A-001", D(2024,9,1), D(2025,9,3), D(2024,1,1))]),
 S("点検済日が文字『済』", [("A-001", D(2025,9,1), "済", None)]),
 S("点検済日が全角『２０２６/９/５』", [("A-001", D(2025,9,1), "２０２６/９/５", None)]),
 S("引渡日が0(範囲外の数)", [("A-001", 0, None, None)]),
 S("基準日が空欄", [("A-001", D(2025,9,1), None, None)], None),
 S("基準日が文字", [("A-001", D(2025,9,1), None, None)], "今日"),
 S("空行", [(None, None, None, None), ("A-001", D(2025,9,1), None, None)]),
]

def main():
    ng = 0
    for name, rows, base in SCEN:
        ok, g, e = check(rows, base); ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: {g}")
        if not ok: print("   期待", e)
    rnd = random.Random(20261026); rows = []
    for i in range(500):
        b = datetime.fromordinal(date(2023, 1, 1).toordinal() + rnd.randint(0, 1200))
        c = d = None
        if rnd.random() < 0.5: c = b + timedelta(days=rnd.randint(330, 400))
        if c and rnd.random() < 0.5: d = c + timedelta(days=rnd.randint(330, 400))
        a = f"J-{i:03d}"; x = rnd.random()
        if x < 0.02: a = None
        elif x < 0.04: a = rows[i - 1][0] if i else a
        elif x < 0.06: b = rnd.choice(["9/1", None, 0])
        elif x < 0.08: c = rnd.choice(["済", "２０２６/９/５", b - timedelta(days=5) if isinstance(b, datetime) else None])
        elif x < 0.10: d = rnd.choice(["済", (b - timedelta(days=5)) if isinstance(b, datetime) else None])
        elif x < 0.12: c, d = None, (d or (b + timedelta(days=700) if isinstance(b, datetime) else None))
        rows.append((a, b, c, d))
    ok, g, e = check(rows, D(2026, 10, 26)); ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: {len(rows)}行 全項目一致={ok}")
    if not ok:
        for i, (x, y) in enumerate(zip(g, e)):
            if x != y: print("  行", i + 2, rows[i], x, y); break
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
