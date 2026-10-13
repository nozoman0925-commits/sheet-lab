"""記事『郵便番号のハイフンあり・なしをそろえる方法』の数式をLibreOfficeで再計算して期待値と照合する。
使い方: python3 site/verify/postal_code.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os
from openpyxl import Workbook, load_workbook

# (入力, 入力の型 'text'|'num', B=記号除去, C=6桁なら0補完, D=ハイフン付き7桁 or 要確認, E=ハイフンなし7桁 or 要確認)
CASES = [
    ("1234567",   "text", "1234567", "1234567", "123-4567", "1234567"),
    ("123-4567",  "text", "1234567", "1234567", "123-4567", "1234567"),
    ("123ー4567", "text", "1234567", "1234567", "123-4567", "1234567"),
    ("123－4567", "text", "1234567", "1234567", "123-4567", "1234567"),
    ("１２３－４５６７", "text", "1234567", "1234567", "123-4567", "1234567"),
    ("〒123-4567", "text", "1234567", "1234567", "123-4567", "1234567"),
    (" 123-4567 ", "text", "1234567", "1234567", "123-4567", "1234567"),
    ("123　4567",  "text", "1234567", "1234567", "123-4567", "1234567"),
    (1234567,     "num",  "1234567", "1234567", "123-4567", "1234567"),
    (234567,      "num",  "234567",  "0234567", "023-4567", "0234567"),   # 数値になって先頭の0が消えた6桁
    ("234567",    "text", "234567",  "234567",  "要確認",   "要確認"),     # 文字列の6桁は0を足さない
    ("12-34567",  "text", "1234567", "1234567", "123-4567", "1234567"),   # ハイフン位置が違っても7桁の数字になれば整形される
    ("123-45678", "text", "12345678", "12345678", "要確認", "要確認"),
    ("1E34567",   "text", "1E34567", "1E34567", "要確認",   "要確認"),
    ("123-456A",  "text", "123456A", "123456A", "要確認",   "要確認"),
    ("",          "text", "",        "",        "要確認",   "要確認"),
]
DASHES = ['"-"', '"－"', '"ー"', '"‐"', '"−"']
def norm(r):
    f = f'A{r}&""'
    for ch in ['"〒"', '" "', '"　"'] + DASHES:
        f = f'SUBSTITUTE({f},{ch},"")'
    return f'=ASC({f})'
C = '=IF(AND(ISNUMBER(A{r}),LEN(B{r})=6),"0"&B{r},B{r})'
OK7 = 'AND(LEN(C{r})=7,SUMPRODUCT(--ISNUMBER(FIND(MID(C{r},ROW($1:$7),1),"0123456789")))=7)'
D = '=IF(' + OK7 + ',LEFT(C{r},3)&"-"&MID(C{r},4,4),"要確認")'
E = '=IF(' + OK7 + ',C{r},"要確認")'

def main():
    wb = Workbook(); ws = wb.active
    for r, (inp, typ, *_ ) in enumerate(CASES, 1):
        if typ == "num": ws.cell(r, 1, inp)
        else:
            ws.cell(r, 1, inp if inp != "" else None); ws.cell(r, 1).number_format = "@"
        ws.cell(r, 2, norm(r)); ws.cell(r, 3, C.format(r=r)); ws.cell(r, 4, D.format(r=r)); ws.cell(r, 5, E.format(r=r))
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    res = load_workbook(os.path.join(out, "t.xlsx"), data_only=True).active
    ng = 0
    for r, (inp, typ, *exp) in enumerate(CASES, 1):
        for c, e in zip((2, 3, 4, 5), exp):
            got = res.cell(r, c).value
            got = "" if got is None else str(got)
            ok = got == e; ng += not ok
            print(("OK " if ok else "NG ") + f"{inp!r}({typ}) col{c}: 期待{e!r} 結果{got!r}")
    print("不一致", ng, "件"); sys.exit(1 if ng else 0)
main()
