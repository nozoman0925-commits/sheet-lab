"""記事『工事の出来高(契約金額×進捗率)と今回請求できる額をExcelで出す式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/progress_billing.py   (soffice が必要。不一致なら終了コード1)
進捗率は 0〜1 の数値(セルの表示形式は%)。出来高は契約金額(税抜)×進捗率を円未満切り捨て。"""
import subprocess, sys, tempfile, os, random
from openpyxl import Workbook, load_workbook

N = 2000

def build(projects, bills):
    """projects: [(番号, 契約金額, 進捗率)], bills: [(番号, 請求額)]"""
    wb = Workbook(); led = wb.active; led.title = "台帳"; bl = wb.create_sheet("請求")
    for i, h in enumerate(["工事番号", "契約金額(税抜)", "進捗率", "出来高", "請求済み", "今回請求できる額", "チェック"], 1):
        led.cell(1, i, h)
    for i, h in enumerate(["工事番号", "請求額(税抜)", "整えた番号", "チェック"], 1):
        bl.cell(1, i, h)
    for r, (no, amt, prog) in enumerate(projects, 2):
        led.cell(r, 1, no); led.cell(r, 2, amt); led.cell(r, 3, prog)
        led.cell(r, 7, (f'=IF(A{r}="","",IF(NOT(ISNUMBER(B{r})),"契約金額が未入力か数値でない",IF(B{r}<=0,"契約金額が0以下",'
                        f'IF(NOT(ISNUMBER(C{r})),"進捗率が未入力か数値でない",IF(OR(C{r}<0,C{r}>1),"進捗率が0〜100%の範囲外",'
                        f'IF(COUNTIF(A$2:A{r},A{r})>1,"工事番号が重複",IF(SUMIFS(請求!$B$2:$B${N},請求!$C$2:$C${N},A{r})>ROUNDDOWN(ROUND(B{r}*C{r},6),0),"請求済みが出来高を超えている","")))))))'))
        led.cell(r, 4, f'=IF(A{r}="","",IF(G{r}="請求済みが出来高を超えている",ROUNDDOWN(ROUND(B{r}*C{r},6),0),IF(G{r}<>"","",ROUNDDOWN(ROUND(B{r}*C{r},6),0))))')
        led.cell(r, 5, f'=IF(A{r}="","",SUMIFS(請求!$B$2:$B${N},請求!$C$2:$C${N},A{r}))')
        led.cell(r, 6, f'=IF(OR(A{r}="",G{r}<>""),"",D{r}-E{r})')
    for r, (no, amt) in enumerate(bills, 2):
        bl.cell(r, 1, no); bl.cell(r, 2, amt)
        bl.cell(r, 3, f'=ASC(TRIM(A{r}))')
        bl.cell(r, 4, (f'=IF(AND(A{r}="",B{r}=""),"",IF(A{r}="","工事番号が空欄",IF(ISNA(MATCH(C{r},台帳!$A$2:$A${N},0)),"工事番号なし",'
                       f'IF(NOT(ISNUMBER(B{r})),"請求額が数値でない",""))))'))
    led["I1"] = "要確認の件数"
    led["J1"] = f'=COUNTIF(G2:G{N},"?*")+COUNTIF(請求!D2:D{N},"?*")'
    return wb

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

def norm(v):
    if v is None: return ""
    return "".join(chr(ord(c) - 0xFEE0) if "！" <= c <= "～" else ("-" if c == "－" else c) for c in str(v).strip().replace("　", " ").strip())

def fw(s):
    return "".join(chr(ord(c) + 0xFEE0) if "!" <= c <= "~" and c != "-" else ("－" if c == "-" else c) for c in s)

def isnum(v): return isinstance(v, (int, float)) and not isinstance(v, bool)

def deg(prog):  # 進捗率(1/1000単位の整数)に直す。float誤差を避ける
    return round(prog * 1000)

def expected(projects, bills):
    keys = {norm(p[0]) for p in projects if p[0] not in (None, "")}
    paid = {k: 0 for k in keys}; brows = []; bbad = 0
    for no, amt in bills:
        if no in (None, "") and amt is None: brows.append(""); continue
        if no in (None, ""): c = "工事番号が空欄"
        elif norm(no) not in keys: c = "工事番号なし"
        elif not isnum(amt): c = "請求額が数値でない"
        else: c = ""
        brows.append(c)
        if c: bbad += 1
        elif True: paid[norm(no)] += amt
    rows = []; seen = set(); pbad = 0
    for no, amt, prog in projects:
        if no in (None, ""): rows.append(("", "", "", "")); continue
        k = norm(no); dup = (no in seen); seen.add(no)
        if not isnum(amt): c = "契約金額が未入力か数値でない"
        elif amt <= 0: c = "契約金額が0以下"
        elif not isnum(prog): c = "進捗率が未入力か数値でない"
        elif prog < 0 or prog > 1: c = "進捗率が0〜100%の範囲外"
        elif dup: c = "工事番号が重複"
        else:
            dek = amt * deg(prog) // 1000
            c = "請求済みが出来高を超えている" if paid[k] > dek else ""
        if c and c != "請求済みが出来高を超えている":
            rows.append(("", paid[k], "", c)); pbad += 1; continue
        dek = amt * deg(prog) // 1000
        if c: pbad += 1; rows.append((dek, paid[k], "", c))
        else: rows.append((dek, paid[k], dek - paid[k], ""))
    return rows, brows, bbad + pbad

SCEN = [
 ("基本: 100万円×35%・請求なし", [("K-001", 1000000, 0.35)], []),
 ("基本: 150万円×30%・着手金45万円を請求済み(今回0円)", [("K-002", 1500000, 0.30)], [("K-002", 450000)]),
 ("円未満の切り捨て(1,234,567円×33.3%)", [("K-003", 1234567, 0.333)], []),
 ("進捗率0%", [("K-004", 800000, 0)], []),
 ("進捗率100%・一部請求済み", [("K-005", 800000, 1)], [("K-005", 300000), ("K-005", 200000)]),
 ("請求済みが出来高を超えている(前払い超過)", [("K-006", 1000000, 0.2)], [("K-006", 300000)]),
 ("進捗率が100%超(50を入れて5000%になった想定)", [("K-007", 1000000, 50)], []),
 ("進捗率がマイナス", [("K-008", 1000000, -0.1)], []),
 ("進捗率に文字('半分')", [("K-009", 1000000, "半分")], []),
 ("進捗率が空欄", [("K-010", 1000000, None)], []),
 ("契約金額が空欄", [("K-011", None, 0.5)], []),
 ("契約金額が0", [("K-012", 0, 0.5)], []),
 ("契約金額に文字('百万円')", [("K-013", "百万円", 0.5)], []),
 ("請求に存在しない工事番号", [("K-014", 1000000, 0.5)], [("K-140", 100000)]),
 ("請求の工事番号が空欄", [("K-015", 1000000, 0.5)], [(None, 100000)]),
 ("請求額に文字", [("K-016", 1000000, 0.5)], [("K-016", "十万")]),
 ("全角・前後スペースの請求番号が同じ工事に集計される", [("K-017", 1000000, 0.6)], [("Ｋ－017", 100000), (" K-017 ", 200000)]),
 ("工事番号の重複(2行目を要確認に)", [("K-018", 1000000, 0.5), ("K-018", 2000000, 0.5)], []),
 ("まったくの空行", [(None, None, None)], [(None, None)]),
]

def main():
    ng = 0
    for name, projects, bills in SCEN:
        res = run_lo(build(projects, bills)); led, bl = res["台帳"], res["請求"]
        rows, brows, bad = expected(projects, bills)
        def v(c): return "" if c.value is None else c.value
        got = [(v(led.cell(r,4)), v(led.cell(r,5)), v(led.cell(r,6)), v(led.cell(r,7))) for r in range(2, len(projects)+2)]
        gb = [v(bl.cell(r,4)) for r in range(2, len(bills)+2)]
        ok = got == rows and gb == brows and led["J1"].value == bad
        ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: 出来高/請求済み/今回/チェック{got} 請求チェック{gb} 要確認{led['J1'].value}")
        if not ok: print("   期待", rows, brows, bad)
    rnd = random.Random(20261021)
    projects = []; bills = []
    for i in range(1, 301):
        no = f"K-{i:03d}"; amt = rnd.randint(50, 3000) * 1000 + rnd.choice([0, 0, 567, 321, 999])
        prog = rnd.randint(0, 1000) / 1000
        t = rnd.random()
        if t < 0.80: projects.append((no, amt, prog))
        elif t < 0.84: projects.append((no, amt, rnd.choice([1.5, 50, -0.2])))
        elif t < 0.88: projects.append((no, rnd.choice([None, 0, "未定"]), prog))
        elif t < 0.92: projects.append((no, amt, rnd.choice([None, "半分"])))
        else: projects.append((no, amt, prog))
    for _ in range(1000):
        p = rnd.choice(projects)[0]; amt = rnd.randint(1, 80) * 10000; t = rnd.random()
        if t < 0.80: bills.append((p, amt))
        elif t < 0.85: bills.append(("K-9%02d" % rnd.randint(0, 99), amt))
        elif t < 0.89: bills.append((None, amt))
        elif t < 0.94: bills.append((fw(p), amt))
        elif t < 0.97: bills.append((" " + p + " ", amt))
        else: bills.append((p, "未定"))
    res = run_lo(build(projects, bills)); led, bl = res["台帳"], res["請求"]
    rows, brows, bad = expected(projects, bills)
    def v(c): return "" if c.value is None else c.value
    got = [(v(led.cell(r,4)), v(led.cell(r,5)), v(led.cell(r,6)), v(led.cell(r,7))) for r in range(2, 302)]
    gb = [v(bl.cell(r,4)) for r in range(2, 1002)]
    ok = got == rows and gb == brows and led["J1"].value == bad
    ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: 工事300件・請求1000行 一致={ok} 要確認{bad}件")
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
