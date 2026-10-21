"""記事『請求と入金をExcelで照合する式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/payment_match.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random, datetime
from openpyxl import Workbook, load_workbook

N = 2000  # 式の範囲 2:2000
BASE = datetime.date(2026, 11, 30)  # 検証用に固定した基準日(TODAY関数そのものは検証しない)

def build(bills, pays):
    """bills: [(請求番号, 請求日, 入金期限, 請求額)], pays: [(入金日, 請求番号(セルの値そのまま), 入金額)]"""
    wb = Workbook(); bl = wb.active; bl.title = "請求"; py = wb.create_sheet("入金")
    for i, h in enumerate(["請求番号", "請求日", "入金期限", "請求額(税込)", "入金済み", "未入金", "状態", "遅れ日数", "チェック", "", "整えた番号"], 1):
        bl.cell(1, i, h)
    for i, h in enumerate(["入金日", "請求番号", "入金額", "整えた番号", "チェック"], 1):
        py.cell(1, i, h)
    bl["M1"], bl["N1"] = "基準日", BASE
    for r, (no, d, due, amt) in enumerate(bills, 2):
        bl.cell(r, 1, no); bl.cell(r, 2, d); bl.cell(r, 3, due); bl.cell(r, 4, amt)
    for r in range(2, len(bills) + 2):
        bl.cell(r, 11, f'=ASC(TRIM(A{r}))')
        bl.cell(r, 5, f'=IF(A{r}="","",SUMIFS(入金!$C$2:$C${N},入金!$D$2:$D${N},K{r}))')
        bl.cell(r, 6, f'=IF(A{r}="","",D{r}-E{r})')
        bl.cell(r, 7, (f'=IF(A{r}="","",IF(F{r}<0,"過入金",IF(F{r}=0,"入金済み",'
                       f'IF(AND(ISNUMBER(C{r}),$N$1>C{r}),"入金遅れ",IF(E{r}>0,"一部入金","未入金")))))'))
        bl.cell(r, 8, f'=IF(G{r}="入金遅れ",$N$1-C{r},"")')
        bl.cell(r, 9, (f'=IF(A{r}="","",IF(NOT(ISNUMBER(C{r})),"入金期限が日付でない",'
                       f'IF(AND(ISNUMBER(B{r}),C{r}<B{r}),"期限が請求日より前","")))'))
    for r, (d, no, amt) in enumerate(pays, 2):
        py.cell(r, 1, d); py.cell(r, 2, no); py.cell(r, 3, amt)
        py.cell(r, 4, f'=ASC(TRIM(B{r}))')
        py.cell(r, 5, (f'=IF(AND(B{r}="",C{r}=""),"",IF(B{r}="","請求番号が空欄",'
                       f'IF(COUNTIF(請求!$K$2:$K${N},D{r})=0,"請求番号なし","")))'))
    bl["P1"], bl["Q1"] = "未紐づけの入金額", f'=SUM(入金!C2:C{N})-SUM(E2:E{N})'
    return wb

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

def fw(s):
    return "".join(chr(ord(c) + 0xFEE0) if "!" <= c <= "~" and c != "-" else ("－" if c == "-" else c) for c in s)

def norm(v):
    if v is None: return ""
    return "".join(chr(ord(c) - 0xFEE0) if "！" <= c <= "～" else ("-" if c == "－" else c) for c in str(v).strip().replace("　", " ").strip())

def expected(bills, pays):
    paid = {norm(b[0]): 0 for b in bills if b[0] not in (None, "")}
    pflags = []; unl = 0
    for d, no, amt in pays:
        k = norm(no)
        if no in (None, "") and amt in (None, ""): pflags.append(""); continue
        if no in (None, ""): pflags.append("請求番号が空欄"); unl += amt or 0; continue
        if k not in paid: pflags.append("請求番号なし"); unl += amt or 0; continue
        paid[k] += amt or 0; pflags.append("")
    rows = []
    for no, d, due, amt in bills:
        if no in (None, ""): rows.append(("", "", "", "", "")); continue
        p = paid[norm(no)]; rest = amt - p
        isd = isinstance(due, datetime.date)
        if rest < 0: st = "過入金"
        elif rest == 0: st = "入金済み"
        elif isd and BASE > due: st = "入金遅れ"
        elif p > 0: st = "一部入金"
        else: st = "未入金"
        late = (BASE - due).days if st == "入金遅れ" else ""
        chk = "入金期限が日付でない" if not isd else ("期限が請求日より前" if isinstance(d, datetime.date) and due < d else "")
        rows.append((p, rest, st, late, chk))
    return rows, pflags, unl

D = datetime.date
SCEN = [
 ("全額入金", [("S-001", D(2026,10,31), D(2026,11,30), 1100000)], [(D(2026,11,25), "S-001", 1100000)]),
 ("2回に分けて入金(一部入金・期限内)", [("S-002", D(2026,11,20), D(2026,12,20), 500000)], [(D(2026,11,25), "S-002", 200000)]),
 ("未入金・期限内", [("S-003", D(2026,11,20), D(2026,12,20), 300000)], []),
 ("未入金・期限超過(入金遅れ)", [("S-004", D(2026,10,1), D(2026,10,31), 300000)], []),
 ("期限が基準日と同じ日(まだ遅れではない)", [("S-005", D(2026,10,31), D(2026,11,30), 300000)], []),
 ("一部入金のまま期限超過", [("S-006", D(2026,10,1), D(2026,10,31), 300000)], [(D(2026,10,25), "S-006", 100000)]),
 ("過入金(二重に振り込まれた)", [("S-007", D(2026,10,1), D(2026,10,31), 300000)], [(D(2026,10,25), "S-007", 300000), (D(2026,10,28), "S-007", 300000)]),
 ("振込手数料440円を引いて入金", [("S-008", D(2026,11,10), D(2026,12,10), 330000)], [(D(2026,11,20), "S-008", 329560)]),
 ("存在しない請求番号への入金", [("S-009", D(2026,11,10), D(2026,12,10), 330000)], [(D(2026,11,20), "S-090", 330000)]),
 ("請求番号が空欄の入金", [("S-010", D(2026,11,10), D(2026,12,10), 330000)], [(D(2026,11,20), None, 330000)]),
 ("全角・前後スペースの請求番号", [("S-011", D(2026,11,10), D(2026,12,10), 330000)], [(D(2026,11,20), "Ｓ－011", 100000), (D(2026,11,21), " S-011 ", 230000)]),
 ("入金期限が請求日より前", [("S-012", D(2026,11,20), D(2026,11,10), 330000)], []),
 ("入金期限が空欄", [("S-013", D(2026,11,20), None, 330000)], []),
 ("入金期限が文字", [("S-014", D(2026,11,20), "月末", 330000)], []),
]

def main():
    ng = 0
    for name, bills, pays in SCEN:
        res = run_lo(build(bills, pays)); bl, py = res["請求"], res["入金"]
        exp, pflags, unl = expected(bills, pays)
        got = []
        for r in range(2, len(bills) + 2):
            got.append(tuple((bl.cell(r, c).value if bl.cell(r, c).value is not None else "") for c in (5, 7, 8, 9)))
        gotv = [(bl.cell(r, 5).value, bl.cell(r, 6).value, bl.cell(r, 7).value, bl.cell(r, 8).value or "", bl.cell(r, 9).value or "") for r in range(2, len(bills) + 2)]
        gf = [(py.cell(r, 5).value or "") for r in range(2, len(pays) + 2)]
        ok = gotv == exp and gf == pflags and bl["Q1"].value == unl
        ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: 入金済み/未入金/状態/遅れ日数/チェック={gotv} 入金側チェック={gf} 未紐づけ={bl['Q1'].value}")
        if not ok: print("   期待", exp, pflags, unl)
    # --- ランダム照合: 請求300件・入金450行(誤入力を混ぜる) ---
    rnd = random.Random(20261022)
    bills = []
    for i in range(1, 301):
        d = D(2026, 8, 1) + datetime.timedelta(days=rnd.randint(0, 110)); due = d + datetime.timedelta(days=rnd.choice([14, 30, 45]))
        t = rnd.random()
        if t < 0.03: due = d - datetime.timedelta(days=rnd.randint(1, 10))
        elif t < 0.05: due = None
        elif t < 0.07: due = "月末"
        bills.append((f"S-{i:03d}", d, due, rnd.randint(1, 300) * 10000 + rnd.choice([0, 0, 1000, 33000])))
    pays = []
    for _ in range(450):
        t = rnd.random(); b = rnd.choice(bills)
        amt = rnd.choice([b[3], b[3], b[3] // 2, b[3] - 440, rnd.randint(1, 100) * 10000]); d = D(2026, 9, 1) + datetime.timedelta(days=rnd.randint(0, 90))
        if t < 0.80: pays.append((d, b[0], amt))
        elif t < 0.86: pays.append((d, "S-9%02d" % rnd.randint(0, 99), amt))
        elif t < 0.91: pays.append((d, None, amt))
        elif t < 0.96: pays.append((d, fw(b[0]), amt))
        else: pays.append((d, " " + b[0] + " ", amt))
    res = run_lo(build(bills, pays)); bl, py = res["請求"], res["入金"]
    exp, pflags, unl = expected(bills, pays)
    gotv = [(bl.cell(r, 5).value, bl.cell(r, 6).value, bl.cell(r, 7).value, bl.cell(r, 8).value or "", bl.cell(r, 9).value or "") for r in range(2, 302)]
    gf = [(py.cell(r, 5).value or "") for r in range(2, 452)]
    ok = gotv == exp and gf == pflags and bl["Q1"].value == unl
    ng += not ok
    from collections import Counter
    print(("OK " if ok else "NG ") + f"ランダム照合: 請求300件・入金450行 全項目一致={ok} 状態の内訳={dict(Counter(x[2] for x in exp))}")
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
