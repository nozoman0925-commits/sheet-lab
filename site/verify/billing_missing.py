"""記事『完成した工事の請求漏れ(完了したのに請求していない工事)をExcelで見つける式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/billing_missing.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random, datetime as dt
from openpyxl import Workbook, load_workbook

N = 1000
G = ('=IF(A{r}&B{r}&C{r}&D{r}&E{r}="","",IF(A{r}="","工事番号が空欄",IF(COUNTIF($A$2:$A${n},A{r})>1,"工事番号が重複",'
     'IF(AND(C{r}<>"",NOT(ISNUMBER(C{r}))),"完了日が日付でない",IF(AND(ISNUMBER(C{r}),OR(C{r}<1,C{r}<>INT(C{r}))),"完了日が日付でない",'
     'IF(AND(D{r}<>"",NOT(ISNUMBER(D{r}))),"請求日が日付でない",IF(AND(ISNUMBER(D{r}),OR(D{r}<1,D{r}<>INT(D{r}))),"請求日が日付でない",'
     'IF(E{r}="","請求額が空欄",IF(NOT(ISNUMBER(E{r})),"請求額が数字でない",IF(OR(E{r}<1,E{r}<>INT(E{r})),"請求額は1以上の整数",'
     'IF(AND(D{r}<>"",C{r}=""),"完了日が空欄なのに請求日がある",IF(AND(D{r}<>"",C{r}>D{r}),"請求日が完了日より前","")))))))))))')
F = ('=IF(OR(G{r}<>"",A{r}&B{r}&C{r}&D{r}&E{r}=""),"",IF(C{r}="","工事中",IF(D{r}<>"","請求済",'
     'IF(NOT(ISNUMBER($K$2)),"基準日を入力",IF(C{r}>$K$2,"完了予定",IF($K$2-C{r}>N($K$3),"請求漏れ","請求待ち"))))))')
H = '=IF(OR(F{r}="請求漏れ",F{r}="請求待ち"),$K$2-C{r},"")'
SUM = {"K5": '=COUNTIF($F$2:$F${n},"請求漏れ")', "K6": '=SUMIF($F$2:$F${n},"請求漏れ",$E$2:$E${n})',
       "K7": '=SUMIF($F$2:$F${n},"請求待ち",$E$2:$E${n})', "K8": '=COUNTIF($F$2:$F${n},"工事中")'}

BASE = dt.date(2026, 10, 27)
def d(n): return BASE - dt.timedelta(days=n)   # 基準日のn日前

def build(rows, base, grace):
    wb = Workbook(); ws = wb.active; ws.title = "請求漏れ"
    for i, h in enumerate(["工事番号", "工事名", "完了日", "請求日", "請求額(予定を含む)", "状態", "チェック", "経過日"], 1): ws.cell(1, i, h)
    ws["J2"] = "基準日"; ws["K2"] = base; ws["J3"] = "猶予日数"; ws["K3"] = grace
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1): ws.cell(r, c, v)
        ws.cell(r, 7, G.format(r=r, n=N)); ws.cell(r, 6, F.format(r=r))
        ws.cell(r, 8, H.format(r=r)).number_format = "0"
    for k, f in SUM.items(): ws[k] = f.format(n=N)
    return wb

def run_lo(wb):
    t = tempfile.mkdtemp(); src = os.path.join(t, "t.xlsx"); wb.save(src)
    out = os.path.join(t, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
isdate = lambda v: isinstance(v, dt.date) and not (isinstance(v, dt.datetime) and (v.hour or v.minute or v.second))
empty = lambda v: v in (None, "")
key = lambda v: v.lower() if isinstance(v, str) else v
ser = lambda v: v.toordinal()

def expect(rows, base, grace):
    out = []
    for a, b, c, dd, e in rows:
        if all(empty(x) for x in (a, b, c, dd, e)): out.append(("", "", "")); continue
        err = None
        if empty(a): err = "工事番号が空欄"
        elif sum(1 for x in rows if not empty(x[0]) and key(x[0]) == key(a)) > 1: err = "工事番号が重複"
        elif not empty(c) and not isdate(c): err = "完了日が日付でない"
        elif not empty(dd) and not isdate(dd): err = "請求日が日付でない"
        elif empty(e): err = "請求額が空欄"
        elif not num(e): err = "請求額が数字でない"
        elif e < 1 or e != int(e): err = "請求額は1以上の整数"
        elif not empty(dd) and empty(c): err = "完了日が空欄なのに請求日がある"
        elif not empty(dd) and ser(c) > ser(dd): err = "請求日が完了日より前"
        if err: out.append(("", "", err)); continue
        if empty(c): out.append(("工事中", "", "")); continue
        if not empty(dd): out.append(("請求済", "", "")); continue
        if not isdate(base): out.append(("基準日を入力", "", "")); continue
        diff = ser(base) - ser(c)
        g = grace if num(grace) else 0
        if diff < 0: out.append(("完了予定", "", ""))
        elif diff > g: out.append(("請求漏れ", diff, ""))
        else: out.append(("請求待ち", diff, ""))
    return out

def exp_sum(rows, ex):
    cnt = sum(1 for x in ex if x[0] == "請求漏れ")
    s1 = sum(r[4] for r, x in zip(rows, ex) if x[0] == "請求漏れ")
    s2 = sum(r[4] for r, x in zip(rows, ex) if x[0] == "請求待ち")
    return (cnt, s1, s2, sum(1 for x in ex if x[0] == "工事中"))

def read(wb, n):
    ws = wb["請求漏れ"]; res = []
    for r in range(2, n + 2):
        res.append(tuple("" if ws.cell(r, c).value is None else ws.cell(r, c).value for c in (6, 8, 7)))
    return res, tuple("" if ws[k].value is None else ws[k].value for k in ("K5", "K6", "K7", "K8"))

def check(rows, base, grace):
    e = expect(rows, base, grace); es = exp_sum(rows, e)
    g, gs = read(run_lo(build(rows, base, grace)), len(rows))
    return g == e and gs == es, g, e, gs, es

S = lambda name, rows, base=BASE, grace=7: (name, rows, base, grace)
SCEN = [
 S("完了日が空欄(工事中)", [("K-001", "山田邸", None, None, 800000)]),
 S("完了日が基準日より先(完了予定)", [("K-001", "山田邸", d(-3), None, 800000)]),
 S("完了日が基準日と同じ日(請求待ち・経過0日)", [("K-001", "山田邸", d(0), None, 800000)]),
 S("完了から3日(猶予内=請求待ち)", [("K-001", "山田邸", d(3), None, 800000)]),
 S("完了からちょうど7日(猶予内=請求待ち)", [("K-001", "山田邸", d(7), None, 800000)]),
 S("完了から8日(請求漏れ)", [("K-001", "山田邸", d(8), None, 800000)]),
 S("完了から60日(請求漏れ)", [("K-001", "山田邸", d(60), None, 800000)]),
 S("請求済(完了から30日後に請求)", [("K-001", "山田邸", d(40), d(10), 800000)]),
 S("請求日が完了日と同じ日(請求済)", [("K-001", "山田邸", d(5), d(5), 800000)]),
 S("請求日が完了日より前(要確認)", [("K-001", "山田邸", d(5), d(9), 800000)]),
 S("完了日が空欄なのに請求日あり", [("K-001", "山田邸", None, d(2), 800000)]),
 S("完了日が全角の文字『２０２６/１０/５』", [("K-001", "山田邸", "２０２６/１０/５", None, 800000)]),
 S("請求日が文字『済』", [("K-001", "山田邸", d(20), "済", 800000)]),
 S("完了日に時刻つき(小数)", [("K-001", "山田邸", dt.datetime(2026, 10, 10, 12, 0), None, 800000)]),
 S("工事番号が空欄", [(None, "山田邸", d(20), None, 800000)]),
 S("工事番号が重複(大文字小文字違い)", [("K-001", "山田邸", d(20), None, 800000), ("k-001", "佐藤邸", d(3), None, 300000)]),
 S("請求額が空欄", [("K-001", "山田邸", d(20), None, None)]),
 S("請求額が文字『80万』", [("K-001", "山田邸", d(20), None, "80万")]),
 S("請求額が全角『８００，０００』", [("K-001", "山田邸", d(20), None, "８００，０００")]),
 S("請求額が0", [("K-001", "山田邸", d(20), None, 0)]),
 S("請求額が小数", [("K-001", "山田邸", d(20), None, 1000.5)]),
 S("猶予日数が空欄(0日として扱う)で完了から1日", [("K-001", "山田邸", d(1), None, 800000)], BASE, None),
 S("猶予日数が空欄で完了日が基準日と同じ日", [("K-001", "山田邸", d(0), None, 800000)], BASE, None),
 S("猶予日数が文字(0日として扱う)で完了から1日", [("K-001", "山田邸", d(1), None, 800000)], BASE, "1週間"),
 S("基準日が空欄(請求待ち判定の行は『基準日を入力』)", [("K-001", "山田邸", d(20), None, 800000), ("K-002", "佐藤邸", d(20), d(5), 300000)], None),
 S("基準日が文字", [("K-001", "山田邸", d(20), None, 800000)], "今日"),
 S("集計(請求漏れ2件・請求待ち1件・工事中1件)", [("K-001", "山田邸", d(30), None, 800000), ("K-002", "佐藤邸", d(9), None, 300000), ("K-003", "鈴木邸", d(2), None, 150000), ("K-004", "高橋邸", None, None, 500000), ("K-005", "田中邸", d(40), d(30), 200000)]),
 S("空行をはさむ", [(None, None, None, None, None), ("K-001", "山田邸", d(8), None, 800000)]),
]

def main():
    ng = 0
    for name, rows, base, grace in SCEN:
        ok, g, e, gs, es = check(rows, base, grace); ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: {g} 集計={gs}")
        if not ok: print("   期待", e, es)
    rnd = random.Random(20261027); rows = []
    for i in range(500):
        c = rnd.choice([None, d(rnd.randint(-5, 90)), d(rnd.randint(-5, 90))])
        dd = None if c is None else rnd.choice([None, None, d(max(0, (BASE - c).days - rnd.randint(0, 30)))])
        a = f"K-{i:03d}"; b = f"邸{i}"; e = rnd.randint(1, 200) * 10000; x = rnd.random()
        if x < 0.02: a = None
        elif x < 0.04: a = rows[i - 1][0] if i else a
        elif x < 0.06: c = rnd.choice(["２０２６/１０/５", "不明", 0, dt.datetime(2026, 10, 10, 9, 0)])
        elif x < 0.08: dd = rnd.choice(["済", "２０２６/１０/５"])
        elif x < 0.10: e = rnd.choice([None, "80万", 0, 1000.5, -5])
        elif x < 0.12: dd = d(rnd.randint(-5, 90)); c = c if c is not None else None
        rows.append((a, b, c, dd, e))
    ok, g, e, gs, es = check(rows, BASE, 7); ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: {len(rows)}行 全項目一致={ok} 集計={gs}")
    if not ok:
        for i, (x, y) in enumerate(zip(g, e)):
            if x != y: print("  行", i + 2, rows[i], x, y); break
        print("  集計", gs, es)
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
