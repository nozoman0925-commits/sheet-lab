"""記事『消費税の端数処理 ROUND/ROUNDUP/ROUNDDOWN の使い分け』をLibreOfficeで再計算し、整数計算の期待値と照合する。
使い方: python3 site/verify/round_tax.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random
from openpyxl import Workbook, load_workbook

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

# 税率10%。整数計算: 切り捨て n//10、切り上げ -(-n//10)、四捨五入 (n*2+10)//20  (n>=0)
def f_down(n): return n // 10
def f_up(n): return -(-n // 10)
def f_round(n): return (2 * n + 10) // 20
# 内税(税込金額Mから消費税額を逆算): M*10/110 = M/11
def inc_down(m): return m * 10 // 110
def inc_round(m): return (2 * m * 10 + 110) // 220
# 桁指定 ROUND(x,-2) 等 は Decimal で
from decimal import Decimal, ROUND_HALF_UP, ROUND_DOWN, ROUND_UP
def dec(x, digits, mode):
    q = Decimal(1).scaleb(-digits)
    return Decimal(x).quantize(q, rounding=mode) if digits >= 0 else (Decimal(x) / Decimal(10) ** -digits).quantize(Decimal(1), rounding=mode) * Decimal(10) ** -digits

def main():
    ng = 0
    random.seed(7)
    vals = [0, 1, 5, 9, 10, 14, 15, 16, 99, 100, 105, 995, 1005, 1234, 9999, 12345, 99995, 100000] + [random.randint(0, 2_000_000) for _ in range(500)]
    wb = Workbook(); ws = wb.active
    for i, n in enumerate(vals, 1):
        ws.cell(i, 1, n)
        ws.cell(i, 2, f"=ROUNDDOWN(A{i}*0.1,0)")
        ws.cell(i, 3, f"=ROUNDUP(A{i}*0.1,0)")
        ws.cell(i, 4, f"=ROUND(A{i}*0.1,0)")
        ws.cell(i, 5, f"=ROUNDDOWN(A{i}*10/110,0)")
        ws.cell(i, 6, f"=ROUND(A{i}*10/110,0)")
        ws.cell(i, 7, f"=ROUNDDOWN(A{i}/1.1,0)")  # 税抜の求め方(参考): 税込M/1.1 切り捨て
        ws.cell(i, 8, f"=A{i}-ROUNDDOWN(A{i}*10/110,0)")  # 税抜 = 税込 - 消費税額
    res = run_lo(wb).worksheets[0]
    bad = 0
    for i, n in enumerate(vals, 1):
        got = tuple(res.cell(i, c).value for c in range(2, 7))
        exp = (f_down(n), f_up(n), f_round(n), inc_down(n), inc_round(n))
        if got != exp: bad += 1; print("NG", n, got, exp)
        if res.cell(i, 8).value != n - inc_down(n): bad += 1; print("NG税抜", n)
    # 税抜の2通りの出し方(税込-切捨税額 と 税込/1.1切捨)が違う例を数える
    diff = [n for i, n in enumerate(vals, 1) if res.cell(i, 7).value != res.cell(i, 8).value]
    ex = [(n, res.cell(vals.index(n)+1, 7).value, res.cell(vals.index(n)+1, 8).value) for n in diff[:3]]
    print(f"税抜の2通りが食い違う件数: {len(diff)}/{len(vals)} 例 {ex}")
    for n in (10005, 11000, 1100):
        a = n // 1 ; print("例", n, "税額切捨", inc_down(n), "税額四捨五入", inc_round(n), "税抜(税込-税額)", n - inc_down(n), "税抜(/1.1切捨)", n * 10 // 11)
    print(f"税率10%の5式 x {len(vals)}件: 不一致 {bad}"); ng += bad

    # 桁指定とマイナス
    cases = [(1234.5678, 0), (1234.5678, 2), (1234.5678, -2), (1250, -2), (1349, -2), (-1234.5, 0), (-1250, -2), (0.5, 0), (-0.5, 0), (2.5, 0)]
    wb = Workbook(); ws = wb.active
    for i, (x, d) in enumerate(cases, 1):
        ws.cell(i, 1, x); ws.cell(i, 2, d)
        ws.cell(i, 3, f"=ROUND(A{i},B{i})"); ws.cell(i, 4, f"=ROUNDUP(A{i},B{i})"); ws.cell(i, 5, f"=ROUNDDOWN(A{i},B{i})"); ws.cell(i, 6, f"=INT(A{i})")
    res = run_lo(wb).worksheets[0]
    bad = 0
    for i, (x, d) in enumerate(cases, 1):
        # Excel/LibreOfficeのROUNDは0から遠ざかる方向の四捨五入(半分は絶対値が大きい方へ)
        exp = (float(dec(str(x), d, ROUND_HALF_UP)), float(dec(str(x), d, ROUND_UP)), float(dec(str(x), d, ROUND_DOWN)))
        got = tuple(float(res.cell(i, c).value) for c in (3, 4, 5))
        ok = got == exp; bad += not ok
        import math
        print(("OK " if ok else "NG ") + f"x={x} 桁={d}: ROUND={got[0]:g} ROUNDUP={got[1]:g} ROUNDDOWN={got[2]:g} INT={res.cell(i,6).value:g}")
    ng += bad
    # 表示形式だけ丸める場合の合計のずれ
    wb = Workbook(); ws = wb.active
    for i, v in enumerate([1.4, 1.4, 1.4], 1): ws.cell(i, 1, v); ws.cell(i, 2, f"=ROUND(A{i},0)")
    ws["A4"] = "=SUM(A1:A3)"; ws["B4"] = "=SUM(B1:B3)"
    r = run_lo(wb).worksheets[0]
    print("表示だけ丸め: 合計", round(r["A4"].value, 10), "ROUND後の合計", r["B4"].value)
    ng += (round(r["A4"].value, 10) != 4.2) + (r["B4"].value != 3)
    print("不一致", ng, "件"); sys.exit(1 if ng else 0)
main()
