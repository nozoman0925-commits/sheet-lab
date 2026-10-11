"""記事『VLOOKUPが#N/Aになる原因』の数式を実際に計算して確認する。"""
from openpyxl import Workbook
import subprocess, json, sys
from openpyxl import load_workbook

wb = Workbook(); ws = wb.active; ws.title = "t"
# 商品表 A:B (A=コード(文字列), B=商品名)
rows = [("A001", "ペン"), ("A002", "ノート"), ("A003", "消しゴム")]
for i, (k, v) in enumerate(rows, 1):
    ws.cell(row=i, column=1, value=k); ws.cell(row=i, column=2, value=v)
# 数値コード表 D:E (D=文字列の数字)
ws["D1"] = "1001"; ws["E1"] = "りんご"; ws["D2"] = "1002"; ws["E2"] = "みかん"
for c in ("D1", "D2"): ws[c].number_format = "@"
# ケース1: 末尾スペース
ws["G1"] = "A002 "
ws["H1"] = "=VLOOKUP(G1,A1:B3,2,FALSE)"            # 期待 #N/A
ws["I1"] = "=VLOOKUP(TRIM(G1),A1:B3,2,FALSE)"      # 期待 ノート
# ケース2: 数値と文字列
ws["G2"] = 1002
ws["H2"] = "=VLOOKUP(G2,D1:E2,2,FALSE)"            # 期待 #N/A
ws["I2"] = '=VLOOKUP(G2&"",D1:E2,2,FALSE)'         # 期待 みかん
# ケース3: 近似一致(4番目省略)・未整列
ws["K1"] = 30; ws["L1"] = "C"; ws["K2"] = 10; ws["L2"] = "A"; ws["K3"] = 20; ws["L3"] = "B"
ws["G3"] = 20
ws["H3"] = "=VLOOKUP(G3,K1:L3,2)"                  # 未整列: 正しい'B'とは限らない
ws["I3"] = "=VLOOKUP(G3,K1:L3,2,FALSE)"            # 期待 B
# ケース4: 検索列が左端にない
ws["N1"] = "ペン"; ws["O1"] = "A001"; ws["N2"] = "ノート"; ws["O2"] = "A002"
ws["G4"] = "A002"
ws["H4"] = "=VLOOKUP(G4,N1:O2,1,FALSE)"            # 検索値は左端列(N)にないので #N/A
ws["I4"] = "=INDEX(N1:N2,MATCH(G4,O1:O2,0))"       # 期待 ノート
# ケース5: IFERRORで隠す
ws["H5"] = '=IFERROR(VLOOKUP(G1,A1:B3,2,FALSE),"該当なし")'  # 期待 該当なし(原因は隠れる)
wb.save("/tmp/vlookup_test.xlsx")
