"""記事『請求書の支払期限(締め日・支払月・支払日)をExcelで出す式』をLibreOfficeで再計算して期待値(Pythonのcalendar)と照合する。
使い方: python3 site/verify/payment_due.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random, calendar
from datetime import datetime, date
from openpyxl import Workbook, load_workbook

WD = "日月火水木金土"
CHK = ('=IF(COUNTA(B{r}:E{r})=0,"",IF(NOT(ISNUMBER(B{r})),"請求日が空欄か日付でない",'
       'IF(OR(B{r}<1,B{r}>73050),"請求日が範囲外",'
       'IF(NOT(ISNUMBER(C{r})),"締め日は1〜31の数字(末日は31)",IF(OR(C{r}<1,C{r}>31,C{r}<>INT(C{r})),"締め日は1〜31の整数(末日は31)",'
       'IF(NOT(ISNUMBER(D{r})),"支払月は0〜12の数字(翌月は1)",IF(OR(D{r}<0,D{r}>12,D{r}<>INT(D{r})),"支払月は0〜12の整数(翌月は1)",'
       'IF(NOT(ISNUMBER(E{r})),"支払日は1〜31の数字(末日は31)",IF(OR(E{r}<1,E{r}>31,E{r}<>INT(E{r})),"支払日は1〜31の整数(末日は31)","")))))))))')

def build(rows):
    """rows: [(請求書名, 請求日, 締め日, 支払月, 支払日)] (セルの値そのまま)"""
    wb = Workbook(); ws = wb.active; ws.title = "支払期限"
    for i, h in enumerate(["請求", "請求日", "締め日(末日=31)", "支払月(翌月=1)", "支払日(末日=31)", "締め月", "支払期限", "曜日", "チェック"], 1):
        ws.cell(1, i, h)
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
        ws.cell(r, 2).number_format = "yyyy/mm/dd"
        ws.cell(r, 9, CHK.format(r=r))
        ws.cell(r, 6, f'=IF(OR(A{r}&B{r}&C{r}&D{r}&E{r}="",I{r}<>""),"",DATE(YEAR(B{r}),MONTH(B{r})+IF(DAY(B{r})>C{r},1,0),1))')
        ws.cell(r, 6).number_format = "yyyy/mm"
        ws.cell(r, 7, f'=IF(F{r}="","",DATE(YEAR(F{r}),MONTH(F{r})+D{r},MIN(E{r},DAY(EOMONTH(DATE(YEAR(F{r}),MONTH(F{r})+D{r},1),0)))))')
        ws.cell(r, 7).number_format = "yyyy/mm/dd"
        ws.cell(r, 8, f'=IF(G{r}="","",CHOOSE(WEEKDAY(G{r}),"日","月","火","水","木","金","土"))')
    return wb

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

def num(v): return isinstance(v, (int, float)) and not isinstance(v, bool)

def expected(rows):
    res = []
    for nm, b, c, d, e in rows:
        if all(x in (None, "") for x in (nm, b, c, d, e)): res.append(("", "", "", "")); continue
        if num(b) and (b < 1 or b > 73050): res.append(("", "", "", "請求日が範囲外")); continue
        if not isinstance(b, datetime): res.append(("", "", "", "請求日が空欄か日付でない")); continue
        if not num(c): res.append(("", "", "", "締め日は1〜31の数字(末日は31)")); continue
        if c < 1 or c > 31 or c != int(c): res.append(("", "", "", "締め日は1〜31の整数(末日は31)")); continue
        if not num(d): res.append(("", "", "", "支払月は0〜12の数字(翌月は1)")); continue
        if d < 0 or d > 12 or d != int(d): res.append(("", "", "", "支払月は0〜12の整数(翌月は1)")); continue
        if not num(e): res.append(("", "", "", "支払日は1〜31の数字(末日は31)")); continue
        if e < 1 or e > 31 or e != int(e): res.append(("", "", "", "支払日は1〜31の整数(末日は31)")); continue
        k = b.year * 12 + b.month - 1 + (1 if b.day > c else 0)
        cy, cm = divmod(k, 12); close = datetime(cy, cm + 1, 1)
        py, pm = divmod(k + int(d), 12); pm += 1
        day = min(int(e), calendar.monthrange(py, pm)[1])
        due = datetime(py, pm, day)
        res.append((close, due, WD[(due.weekday() + 1) % 7], ""))
    return res

def read(res, n):
    ws = res["支払期限"]
    return [tuple((ws.cell(r, c).value if ws.cell(r, c).value is not None else "") for c in (6, 7, 8, 9)) for r in range(2, n + 2)]

def same(got, exp):
    if len(got) != len(exp): return False
    for a, b in zip(got, exp):
        for x, y in zip(a, b):
            if isinstance(y, datetime):
                if not isinstance(x, datetime) or (x.year, x.month, x.day) != (y.year, y.month, y.day): return False
            elif x != y: return False
    return True

D = lambda y, m, d: datetime(y, m, d)
SCEN = [
 ("月末締め翌月末払い(10/23請求)", D(2026,10,23), 31, 1, 31),
 ("月末締めの当日(10/31請求)", D(2026,10,31), 31, 1, 31),
 ("20日締め翌月25日払い・締め日当日(10/20請求)", D(2026,10,20), 20, 1, 25),
 ("20日締め翌月25日払い・翌日(10/21請求)", D(2026,10,21), 20, 1, 25),
 ("25日締め翌月末払い・年またぎ(12/26請求)", D(2026,12,26), 25, 1, 31),
 ("うるう年でない2月末(2027/1/15請求)", D(2027,1,15), 31, 1, 31),
 ("うるう年の2月末(2028/1/15請求)", D(2028,1,15), 31, 1, 31),
 ("30日の月の末日(3/10請求)", D(2026,3,10), 31, 1, 31),
 ("支払日30日を指定して2月払い(1/20請求)", D(2026,1,20), 31, 1, 30),
 ("翌々月払い(11/20請求)", D(2026,11,20), 31, 2, 31),
 ("年またぎ(12/31請求・翌月10日払い)", D(2026,12,31), 31, 1, 10),
 ("当月払い(10/5請求・締めなし月末)", D(2026,10,5), 31, 0, 31),
 ("時刻つきの請求日(10/23 14:30)", datetime(2026,10,23,14,30), 31, 1, 31),
 ("請求日が文字の日付『2026/10/23』", "2026/10/23", 31, 1, 31),
 ("請求日が空欄", None, 31, 1, 31),
 ("請求日が範囲外の数値", 99999999, 31, 1, 31),
 ("締め日が0", D(2026,10,23), 0, 1, 31),
 ("締め日が32", D(2026,10,23), 32, 1, 31),
 ("締め日が全角『３１』", D(2026,10,23), "３１", 1, 31),
 ("締め日が『末』", D(2026,10,23), "末", 1, 31),
 ("支払月がマイナス", D(2026,10,23), 31, -1, 31),
 ("支払月が13", D(2026,10,23), 31, 13, 31),
 ("支払日が小数", D(2026,10,23), 31, 1, 15.5),
 ("支払日が空欄", D(2026,10,23), 31, 1, None),
 ("行が空欄", None, None, None, None),
]

def main():
    ng = 0
    for name, *v in SCEN:
        rows = [("t", *v)] if name != "行が空欄" else [(None, None, None, None, None)]
        got = read(run_lo(build(rows)), 1); exp = expected(rows)
        ok = same(got, exp); ng += not ok
        g = [x.strftime("%Y/%m/%d") if isinstance(x, datetime) else x for x in got[0]]
        print(("OK " if ok else "NG ") + f"{name}: 締め月/支払期限/曜日/チェック={g}")
        if not ok: print("   期待", exp)
    rnd = random.Random(20261023)
    rows = []
    for i in range(1000):
        b = datetime.fromordinal(date(2024, 1, 1).toordinal() + rnd.randint(0, 1460))
        c = rnd.choice([10, 15, 20, 25, 28, 30, 31, 31, 31]); d = rnd.choice([0, 1, 1, 1, 2, 3]); e = rnd.choice([5, 10, 15, 20, 25, 28, 29, 30, 31, 31])
        t = rnd.random()
        if t < 0.03: b = rnd.choice(["2026/10/23", None, "令和8年10月23日"])
        elif t < 0.05: c = rnd.choice([0, 32, "末", "３１", None, 15.5])
        elif t < 0.07: d = rnd.choice([-1, 13, "翌月", None])
        elif t < 0.09: e = rnd.choice([0, 40, "末", None, 2.5])
        rows.append((f"請求{i}", b, c, d, e))
    got = read(run_lo(build(rows)), len(rows)); exp = expected(rows)
    ok = same(got, exp); ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: 1000行 全項目一致={ok}")
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
