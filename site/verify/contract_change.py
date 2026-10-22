"""記事『追加工事・変更契約を反映した現在の契約金額をExcelで出す式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/contract_change.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random
from datetime import datetime, date
from openpyxl import Workbook, load_workbook

N = 1000
CK = ('=IF(A{r}&B{r}="","",IF(A{r}="","工事番号が空欄",IF(COUNTIF($A$2:$A${n},A{r})>1,"工事番号が重複",'
      'IF(NOT(ISNUMBER(B{r})),"当初契約額が数字でない",IF(OR(B{r}<0,B{r}<>INT(B{r})),"当初契約額は0以上の整数",'
      'IF(COUNTIFS(変更!$A$2:$A${n},A{r},変更!$E$2:$E${n},"?*")>0,"この工事の変更行にエラーがある",'
      'IF(B{r}+SUMIFS(変更!$D$2:$D${n},変更!$A$2:$A${n},A{r})<0,"変更の合計が大きすぎる(現契約額がマイナス)","")))))))')
CC = ('=IF(A{r}&B{r}&C{r}&D{r}="","",IF(A{r}="","工事番号が空欄",IF(COUNTIF(契約!$A$2:$A${n},A{r})=0,"契約表にない工事番号",'
      'IF(NOT(ISNUMBER(B{r})),"変更日が日付でない",IF(OR(B{r}<1,B{r}>73050),"変更日が範囲外",'
      'IF(NOT(ISNUMBER(D{r})),"増減額が数字でない",IF(D{r}<>INT(D{r}),"増減額は整数の円",IF(D{r}=0,"増減額が0","")))))))) ')

def build(contracts, changes):
    wb = Workbook(); ws = wb.active; ws.title = "契約"; wc = wb.create_sheet("変更")
    for i, h in enumerate(["工事番号", "当初契約額(税抜)", "変更の合計", "現在の契約額", "変更回数", "チェック"], 1): ws.cell(1, i, h)
    for i, h in enumerate(["工事番号", "変更日", "内容", "増減額(税抜・減額はマイナス)", "チェック"], 1): wc.cell(1, i, h)
    for r, row in enumerate(contracts, 2):
        for c, v in enumerate(row, 1): ws.cell(r, c, v)
        ws.cell(r, 6, CK.format(r=r, n=N))
        ws.cell(r, 3, f'=IF(OR(F{r}<>"",A{r}&B{r}=""),"",SUMIFS(変更!$D$2:$D${N},変更!$A$2:$A${N},A{r}))')
        ws.cell(r, 4, f'=IF(OR(F{r}<>"",A{r}&B{r}=""),"",B{r}+C{r})')
        ws.cell(r, 5, f'=IF(OR(F{r}<>"",A{r}&B{r}=""),"",COUNTIFS(変更!$A$2:$A${N},A{r}))')
    for r, row in enumerate(changes, 2):
        for c, v in enumerate(row, 1): wc.cell(r, c, v)
        wc.cell(r, 2).number_format = "yyyy/mm/dd"
        wc.cell(r, 5, CC.format(r=r, n=N).strip())
    return wb

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
key = lambda v: v.lower() if isinstance(v, str) else v

def exp_change(contracts, changes):
    ids = [key(c[0]) for c in contracts if c[0] not in (None, "")]
    out = []
    for a, b, c, d in changes:
        if all(x in (None, "") for x in (a, b, c, d)): out.append(""); continue
        if a in (None, ""): out.append("工事番号が空欄"); continue
        if key(a) not in ids: out.append("契約表にない工事番号"); continue
        if not isinstance(b, datetime): out.append("変更日が日付でない"); continue
        if not num(d): out.append("増減額が数字でない"); continue
        if d != int(d): out.append("増減額は整数の円"); continue
        out.append("増減額が0" if d == 0 else "")
    return out

def exp_contract(contracts, changes):
    ce = exp_change(contracts, changes)
    res = []
    for a, b in contracts:
        if a in (None, "") and b in (None, ""): res.append(("", "", "", "")); continue
        if a in (None, ""): res.append(("", "", "", "工事番号が空欄")); continue
        if sum(1 for c in contracts if c[0] not in (None, "") and key(c[0]) == key(a)) > 1: res.append(("", "", "", "工事番号が重複")); continue
        if not num(b): res.append(("", "", "", "当初契約額が数字でない")); continue
        if b < 0 or b != int(b): res.append(("", "", "", "当初契約額は0以上の整数")); continue
        mine = [(ch, e) for ch, e in zip(changes, ce) if ch[0] not in (None, "") and key(ch[0]) == key(a)]
        if any(e for _, e in mine): res.append(("", "", "", "この工事の変更行にエラーがある")); continue
        s = sum(ch[3] for ch, _ in mine if num(ch[3]))
        if b + s < 0: res.append(("", "", "", "変更の合計が大きすぎる(現契約額がマイナス)")); continue
        res.append((s, b + s, len(mine), ""))
    return res, ce

def read(wb, nc, nh):
    ws, wc = wb["契約"], wb["変更"]
    g = lambda w, r, c: (w.cell(r, c).value if w.cell(r, c).value is not None else "")
    return ([tuple(g(ws, r, c) for c in (3, 4, 5, 6)) for r in range(2, nc + 2)], [g(wc, r, 5) for r in range(2, nh + 2)])

def check(contracts, changes):
    ec, ech = exp_contract(contracts, changes)
    gc, gch = read(run_lo(build(contracts, changes)), len(contracts), len(changes))
    return gc == ec and gch == ech, gc, ec, gch, ech

D = lambda m, d: datetime(2026, m, d)
BASE = [("A-001", 3000000), ("A-002", 1500000)]
SCEN = [
 ("追加工事1回(+200,000)", BASE, [("A-001", D(9,10), "追加工事", 200000)]),
 ("追加と減額の2回(+200,000と−50,000)", BASE, [("A-001", D(9,10), "追加", 200000), ("A-001", D(9,20), "減額", -50000)]),
 ("変更なしの工事", BASE, [("A-001", D(9,10), "追加", 200000)]),
 ("大文字小文字の違い『a-001』", BASE, [("a-001", D(9,10), "追加", 100000)]),
 ("存在しない工事番号『A-009』", BASE, [("A-009", D(9,10), "追加", 100000)]),
 ("全角の工事番号『Ａ-００１』", BASE, [("Ａ-００１", D(9,10), "追加", 100000)]),
 ("末尾にスペース『A-001␣』", BASE, [("A-001 ", D(9,10), "追加", 100000)]),
 ("増減額が文字『20万』", BASE, [("A-001", D(9,10), "追加", "20万")]),
 ("増減額が全角『２００，０００』", BASE, [("A-001", D(9,10), "追加", "２００，０００")]),
 ("増減額が空欄", BASE, [("A-001", D(9,10), "追加", None)]),
 ("増減額が小数", BASE, [("A-001", D(9,10), "追加", 1000.5)]),
 ("増減額が0", BASE, [("A-001", D(9,10), "追加", 0)]),
 ("変更日が文字『9/10』", BASE, [("A-001", "9/10", "追加", 100000)]),
 ("変更日が空欄", BASE, [("A-001", None, "追加", 100000)]),
 ("工事番号が空欄の変更行", BASE, [(None, D(9,10), "追加", 100000)]),
 ("減額が大きく契約額がマイナス", BASE, [("A-002", D(9,10), "大幅減額", -2000000)]),
 ("契約表の工事番号が重複", [("A-001", 3000000), ("A-001", 100)], [("A-001", D(9,10), "追加", 1)]),
 ("当初契約額が文字『300万』", [("A-001", "300万")], []),
 ("当初契約額が空欄", [("A-001", None)], []),
 ("契約表の工事番号が空欄", [(None, 1000)], []),
 ("契約表の行が空欄", [(None, None)], [(None, None, None, None)]),
 ("契約行を消した(A-002が表にない)", [("A-001", 3000000)], [("A-002", D(9,10), "追加", 100000)]),
]

def main():
    ng = 0
    for name, c, h in SCEN:
        ok, gc, ec, gch, ech = check(c, h); ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: 契約={gc} 変更チェック={gch}")
        if not ok: print("   期待", ec, ech)
    rnd = random.Random(20261023)
    contracts = []
    for i in range(200):
        b = rnd.randint(0, 20) * 500000 + rnd.choice([0, 0, 12000])
        contracts.append((f"J-{i:03d}", b))
    t = [(0.03, "dup"), (0.05, "empty")]
    for i in range(len(contracts)):
        x = rnd.random()
        if x < 0.02: contracts[i] = (contracts[i - 1][0] if i else "J-000", contracts[i][1])
        elif x < 0.04: contracts[i] = (contracts[i][0], rnd.choice(["300万", None, -5, 100.5]))
        elif x < 0.05: contracts[i] = (None, contracts[i][1])
    changes = []
    for i in range(600):
        a = f"J-{rnd.randint(0, 204):03d}"
        dt = datetime.fromordinal(date(2026, 1, 1).toordinal() + rnd.randint(0, 300))
        amt = rnd.choice([-1, 1]) * rnd.randint(1, 40) * 10000
        x = rnd.random()
        if x < 0.02: a = rnd.choice(["j-" + a[2:], a + " ", "Ｊ-" + a[2:], None])
        elif x < 0.04: dt = rnd.choice(["9/10", None])
        elif x < 0.06: amt = rnd.choice(["20万", None, 1000.5, 0, "２００"])
        changes.append((a, dt, "変更", amt))
    ok, gc, ec, gch, ech = check(contracts, changes); ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: 契約200行・変更600行 全項目一致={ok}")
    if not ok:
        for i, (a, b) in enumerate(zip(gc, ec)):
            if a != b: print("  契約", i + 2, contracts[i], a, b); break
        for i, (a, b) in enumerate(zip(gch, ech)):
            if a != b: print("  変更", i + 2, changes[i], a, b); break
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
