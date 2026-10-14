"""note原稿『エクセルで重複に印を付ける』の式をLibreOfficeで再計算して期待値と照合する。
使い方: python3 site/verify/dup_mark.py"""
import subprocess, sys, tempfile, os, re
from openpyxl import Workbook, load_workbook

DATA = [
    "山田商事", "佐藤工業", "山田商事", "山田商事 ", "山田　商事",
    "山田*", "abc", "ABC", "鈴木電機", "佐藤工業", "Q?", "QX",
]
N = len(DATA); LAST = N + 1

def wild_to_re(p):
    out = ""; i = 0
    while i < len(p):
        c = p[i]
        if c == "~" and i + 1 < len(p): out += re.escape(p[i + 1]); i += 2; continue
        out += ".*" if c == "*" else "." if c == "?" else re.escape(c); i += 1
    return re.compile("^" + out + "$", re.I | re.S)

def tidy(s): return re.sub(r" +", " ", s.replace("　", " ")).strip(" ")

def main():
    wb = Workbook(); ws = wb.active
    for i, a in enumerate(DATA, 2):
        ws.cell(i, 1, a)
        ws[f"B{i}"] = f"=COUNTIF($A$2:$A${LAST},A{i})"
        ws[f"C{i}"] = f'=IF(COUNTIF($A$2:A{i},A{i})>1,"重複","")'
        ws[f"D{i}"] = f"=SUMPRODUCT(--($A$2:$A${LAST}=A{i}))"
        ws[f"E{i}"] = f'=TRIM(SUBSTITUTE(A{i},"　"," "))'
        ws[f"F{i}"] = f"=SUMPRODUCT(--($E$2:$E${LAST}=E{i}))"
        ws[f"G{i}"] = f'=IF(SUMPRODUCT(--($E$2:E{i}=E{i}))>1,"重複","")'
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "o"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    ws = load_workbook(os.path.join(out, "t.xlsx"), data_only=True).active
    ng = 0
    tid = [tidy(a) for a in DATA]
    for k, a in enumerate(DATA):
        r = k + 2
        b = sum(1 for x in DATA if wild_to_re(a).match(x))           # COUNTIF: ワイルドカード・大文字小文字無視
        c = "重複" if sum(1 for x in DATA[:k + 1] if wild_to_re(a).match(x)) > 1 else ""
        dd = sum(1 for x in DATA if x.lower() == a.lower())          # =比較: 大文字小文字無視・ワイルドカードなし
        f = sum(1 for x in tid if x.lower() == tid[k].lower())
        g = "重複" if sum(1 for x in tid[:k + 1] if x.lower() == tid[k].lower()) > 1 else ""
        got = (ws[f"B{r}"].value, ws[f"C{r}"].value or "", ws[f"D{r}"].value, ws[f"F{r}"].value, ws[f"G{r}"].value or "")
        exp = (b, c, dd, f, g)
        ok = got == exp; ng += not ok
        print(("OK " if ok else "NG ") + f"{a!r}: COUNTIF={got[0]} 印={got[1]!r} SUMPRODUCT={got[2]} 整形後={got[3]} 整形後の印={got[4]!r} (期待 {exp})")
    print("不一致", ng, "件 /", N * 5, "件"); sys.exit(1 if ng else 0)
main()
