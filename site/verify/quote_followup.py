"""記事『見積の提出後の回答待ちを管理し、期限切れ・受注額・受注率をExcelで出す式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/quote_followup.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random
from datetime import datetime, date, timedelta
from openpyxl import Workbook, load_workbook

N = 1000
G = ('=IF(A{r}&B{r}&C{r}&D{r}="","",IF(A{r}="","見積番号が空欄",IF(COUNTIF($A$2:$A${n},A{r})>1,"見積番号が重複",'
     'IF(NOT(ISNUMBER($K$1)),"基準日が日付でない",IF(OR(NOT(ISNUMBER($K$2)),$K$2<1,$K$2<>INT($K$2)),"有効日数は1以上の整数",'
     'IF(NOT(ISNUMBER(B{r})),"提出日が日付でない",IF(OR(B{r}<1,B{r}>73050),"提出日が範囲外",IF(B{r}>$K$1,"提出日が基準日より後",'
     'IF(NOT(ISNUMBER(C{r})),"金額が数字でない",IF(OR(C{r}<0,C{r}<>INT(C{r})),"金額は0以上の整数",'
     'IF(AND(D{r}<>"受注",D{r}<>"失注",D{r}<>""),"結果は受注・失注・空欄のどれか","")))))))))))')
E = '=IF(OR(G{r}<>"",A{r}&B{r}&C{r}&D{r}=""),"",B{r}+$K$2)'
F = ('=IF(OR(G{r}<>"",A{r}&B{r}&C{r}&D{r}=""),"",IF(D{r}="受注","受注",IF(D{r}="失注","失注",'
     'IF(E{r}<$K$1,"期限切れ",IF(E{r}-$K$1<=7,"7日以内","回答待ち")))))')
SUM = {"K5": '=SUMIFS($C$2:$C${n},$F$2:$F${n},"受注")',
       "K6": '=SUMIFS($C$2:$C${n},$F$2:$F${n},"期限切れ")+SUMIFS($C$2:$C${n},$F$2:$F${n},"7日以内")+SUMIFS($C$2:$C${n},$F$2:$F${n},"回答待ち")',
       "K7": '=COUNTIF($F$2:$F${n},"期限切れ")',
       "K8": '=IF(COUNTIF($F$2:$F${n},"受注")+COUNTIF($F$2:$F${n},"失注")=0,"",COUNTIF($F$2:$F${n},"受注")/(COUNTIF($F$2:$F${n},"受注")+COUNTIF($F$2:$F${n},"失注")))'}

def build(rows, base, days):
    wb = Workbook(); ws = wb.active; ws.title = "見積"
    for i, h in enumerate(["見積番号", "提出日", "金額(税抜)", "結果", "有効期限", "状態", "チェック"], 1): ws.cell(1, i, h)
    ws["J1"] = "基準日"; ws["K1"] = base; ws["K1"].number_format = "yyyy/mm/dd"
    ws["J2"] = "有効日数"; ws["K2"] = days
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
            if isinstance(v, datetime): ws.cell(r, c).number_format = "yyyy/mm/dd"
        ws.cell(r, 7, G.format(r=r, n=N)); ws.cell(r, 5, E.format(r=r)); ws.cell(r, 5).number_format = "yyyy/mm/dd"
        ws.cell(r, 6, F.format(r=r))
    for k, f in SUM.items(): ws[k] = f.format(n=N)
    return wb

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

isdt = lambda v: isinstance(v, datetime)
num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
key = lambda v: v.lower() if isinstance(v, str) else v
empty = lambda v: v in (None, "")

def expect(rows, base, days):
    out = []
    for a, b, c, d in rows:
        if all(empty(x) for x in (a, b, c, d)): out.append(("", "", "")); continue
        err = None
        if empty(a): err = "見積番号が空欄"
        elif sum(1 for x in rows if not empty(x[0]) and key(x[0]) == key(a)) > 1: err = "見積番号が重複"
        elif not isdt(base): err = "基準日が日付でない"
        elif not num(days) or days < 1 or days != int(days): err = "有効日数は1以上の整数"
        elif b == 0 and num(b): err = "提出日が範囲外"
        elif not isdt(b): err = "提出日が日付でない"
        elif b > base: err = "提出日が基準日より後"
        elif not num(c): err = "金額が数字でない"
        elif c < 0 or c != int(c): err = "金額は0以上の整数"
        elif d not in (None, "", "受注", "失注"): err = "結果は受注・失注・空欄のどれか"
        if err: out.append(("", "", err)); continue
        due = b.date() + timedelta(days=days)
        if d == "受注": st = "受注"
        elif d == "失注": st = "失注"
        else:
            diff = (due - base.date()).days
            st = "期限切れ" if diff < 0 else "7日以内" if diff <= 7 else "回答待ち"
        out.append((due, st, ""))
    return out

def exp_sum(rows, ex):
    won = sum(r[2] for r, e in zip(rows, ex) if e[1] == "受注")
    wait = sum(r[2] for r, e in zip(rows, ex) if e[1] in ("期限切れ", "7日以内", "回答待ち"))
    late = sum(1 for e in ex if e[1] == "期限切れ")
    nw = sum(1 for e in ex if e[1] == "受注"); nl = sum(1 for e in ex if e[1] == "失注")
    return (won, wait, late, "" if nw + nl == 0 else nw / (nw + nl))

def read(wb, n):
    ws = wb["見積"]; res = []
    for r in range(2, n + 2):
        v = [ws.cell(r, c).value for c in (5, 6, 7)]
        v = ["" if x is None else x for x in v]
        if isinstance(v[0], datetime): v[0] = v[0].date()
        res.append(tuple(v))
    s = tuple("" if ws[k].value is None else ws[k].value for k in ("K5", "K6", "K7", "K8"))
    return res, s

def same_sum(a, b):
    return all((x == y) or (num(x) and num(y) and abs(x - y) < 1e-9) for x, y in zip(a, b))

def check(rows, base, days=30):
    e = expect(rows, base, days); es = exp_sum(rows, e)
    g, gs = read(run_lo(build(rows, base, days)), len(rows))
    return g == e and same_sum(gs, es), g, e, gs, es

D = lambda y, m, d: datetime(y, m, d)
BASE = D(2026, 10, 26)
S = lambda name, rows, base=BASE, days=30: (name, rows, base, days)
SCEN = [
 S("提出から10日・結果なし(回答待ち)", [("Q-001", D(2026,10,16), 1000000, None)]),
 S("提出から23日=期限まで7日(7日以内)", [("Q-001", D(2026,10,3), 1000000, None)]),
 S("提出から22日=期限まで8日(回答待ち)", [("Q-001", D(2026,10,4), 1000000, None)]),
 S("期限が今日(7日以内)", [("Q-001", D(2026,9,26), 1000000, None)]),
 S("期限が昨日(期限切れ)", [("Q-001", D(2026,9,25), 1000000, None)]),
 S("提出が今日(回答待ち)", [("Q-001", D(2026,10,26), 500000, None)]),
 S("受注(期限が過ぎていても受注)", [("Q-001", D(2026,8,1), 2000000, "受注")]),
 S("失注", [("Q-001", D(2026,8,1), 2000000, "失注")]),
 S("受注と失注と回答待ちの集計", [("Q-001", D(2026,8,1), 2000000, "受注"), ("Q-002", D(2026,9,1), 800000, "失注"), ("Q-003", D(2026,10,20), 300000, None), ("Q-004", D(2026,9,1), 100000, None), ("Q-005", D(2026,10,1), 400000, "受注")]),
 S("受注率の分母が0(結果がすべて空欄)", [("Q-001", D(2026,10,20), 300000, None)]),
 S("有効日数を14日に変更", [("Q-001", D(2026,10,10), 300000, None)], days=14),
 S("見積番号が空欄", [(None, D(2026,10,20), 300000, None)]),
 S("見積番号が重複(大文字小文字違い)", [("Q-001", D(2026,10,20), 300000, None), ("q-001", D(2026,10,21), 100000, None)]),
 S("提出日が文字『10/20』", [("Q-001", "10/20", 300000, None)]),
 S("提出日が空欄", [("Q-001", None, 300000, None)]),
 S("提出日が0", [("Q-001", 0, 300000, None)]),
 S("提出日が基準日より後(日付の逆転)", [("Q-001", D(2026,11,5), 300000, None)]),
 S("金額が文字『30万』", [("Q-001", D(2026,10,20), "30万", None)]),
 S("金額が全角『３００，０００』", [("Q-001", D(2026,10,20), "３００，０００", None)]),
 S("金額が空欄", [("Q-001", D(2026,10,20), None, None)]),
 S("金額が小数", [("Q-001", D(2026,10,20), 1000.5, None)]),
 S("金額がマイナス", [("Q-001", D(2026,10,20), -1000, None)]),
 S("結果が『受注済』", [("Q-001", D(2026,10,20), 300000, "受注済")]),
 S("結果の末尾にスペース『受注␣』", [("Q-001", D(2026,10,20), 300000, "受注 ")]),
 S("結果が『○』", [("Q-001", D(2026,10,20), 300000, "○")]),
 S("基準日が空欄", [("Q-001", D(2026,10,20), 300000, None)], None),
 S("基準日が文字", [("Q-001", D(2026,10,20), 300000, None)], "今日"),
 S("有効日数が0", [("Q-001", D(2026,10,20), 300000, None)], days=0),
 S("有効日数が小数", [("Q-001", D(2026,10,20), 300000, None)], days=30.5),
 S("有効日数が空欄", [("Q-001", D(2026,10,20), 300000, None)], days=None),
 S("空行をはさむ", [(None, None, None, None), ("Q-001", D(2026,10,20), 300000, None)]),
]

def main():
    ng = 0
    for name, rows, base, days in SCEN:
        ok, g, e, gs, es = check(rows, base, days); ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: {g} 集計={gs}")
        if not ok: print("   期待", e, es)
    rnd = random.Random(20261026); rows = []
    for i in range(500):
        b = datetime.fromordinal(date(2026, 1, 1).toordinal() + rnd.randint(0, 298))
        c = rnd.randint(1, 60) * 50000; d = rnd.choice([None, None, None, "受注", "失注"])
        a = f"Q-{i:03d}"; x = rnd.random()
        if x < 0.02: a = None
        elif x < 0.04: a = rows[i - 1][0] if i else a
        elif x < 0.06: b = rnd.choice(["10/20", None, 0, datetime(2026, 12, 1)])
        elif x < 0.08: c = rnd.choice(["30万", None, -5, 1000.5, "３０００００"])
        elif x < 0.10: d = rnd.choice(["受注済", "受注 ", "○", 1])
        rows.append((a, b, c, d))
    ok, g, e, gs, es = check(rows, D(2026, 10, 26)); ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: {len(rows)}行 全項目一致={ok} 集計={gs}")
    if not ok:
        for i, (x, y) in enumerate(zip(g, e)):
            if x != y: print("  行", i + 2, rows[i], x, y); break
        print("  集計", gs, es)
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
