"""記事『電話番号のハイフンを一括で削除する方法』の数式をLibreOfficeで再計算して期待値と照合する。
使い方: python3 site/verify/phone_hyphen.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os
from openpyxl import Workbook, load_workbook

# (入力, 式の種類ごとの期待値)
CASES = [
    # 入力,             ハイフンのみ削除, ハイフン・括弧・スペース全削除, 全角対応(ASC+全角ハイフン類)
    ("03-1234-5678",    "0312345678", "0312345678", "0312345678"),
    ("090-1234-5678",   "09012345678", "09012345678", "09012345678"),
    ("(03)1234-5678",   "(03)12345678", "0312345678", "0312345678"),
    ("03 1234 5678",    "03 1234 5678", "0312345678", "0312345678"),
    ("０３－１２３４－５６７８", "０３－１２３４－５６７８", "０３－１２３４－５６７８", "0312345678"),
    ("03ー1234ー5678",  "03ー1234ー5678", "03ー1234ー5678", "03ｰ1234ｰ5678"),  # 長音符(ー)は削除対象外。ASCで半角の｢ｰ｣になって残る
]
F1 = '=SUBSTITUTE(A{r},"-","")'
F2 = '=SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(A{r},"-",""),"(",""),")","")," ",""),"　","")'
F3 = '=SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(ASC(A{r}),"-",""),"(",""),")","")," ",""),"－","")'

def main():
    wb = Workbook(); ws = wb.active
    for r, (inp, *_ ) in enumerate(CASES, 1):
        ws.cell(r, 1, inp); ws.cell(r, 1).number_format = "@"
        ws.cell(r, 2, F1.format(r=r)); ws.cell(r, 3, F2.format(r=r)); ws.cell(r, 4, F3.format(r=r))
    # 数値で入っている電話番号(先頭0は既に消えている)
    ws.cell(10, 1, 312345678); ws.cell(10, 2, '=SUBSTITUTE(A10,"-","")'); ws.cell(10, 3, '=ISNUMBER(B10)')
    d = tempfile.mkdtemp()
    src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    res = load_workbook(os.path.join(out, "t.xlsx"), data_only=True).active
    ng = 0
    for r, (inp, *exp) in enumerate(CASES, 1):
        for c, e in zip((2, 3, 4), exp):
            got = res.cell(r, c).value
            ok = got == e; ng += not ok
            print(("OK " if ok else "NG ") + f"{inp!r} col{c}: 期待{e!r} 結果{got!r}")
    got = res.cell(10, 2).value, res.cell(10, 3).value
    print("数値セル:", got)
    # SUBSTITUTEは数値でも文字列を返す(=結果は文字列。ISNUMBERはFALSE)
    if got != ("312345678", False): ng += 1; print("NG 数値セルの結果が想定と違う")
    print("不一致", ng, "件"); sys.exit(1 if ng else 0)

main()
