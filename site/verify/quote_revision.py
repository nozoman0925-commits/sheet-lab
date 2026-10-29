"""記事『見積書の改訂版(第2版・第3版)から最新版だけを拾うExcelの式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/quote_revision.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random, datetime as dt
from openpyxl import Workbook, load_workbook

N = 1000
G = ('=IF(A{r}&B{r}&C{r}&D{r}="","",IF(A{r}="","工事番号が空欄",IF(B{r}="","版が空欄",IF(NOT(ISNUMBER(B{r})),"版が数字でない",'
     'IF(OR(B{r}<1,B{r}<>INT(B{r})),"版は1以上の整数",IF(C{r}="","見積額が空欄",IF(NOT(ISNUMBER(C{r})),"見積額が数字でない",'
     'IF(OR(C{r}<1,C{r}<>INT(C{r})),"見積額は1以上の整数",IF(D{r}="","提出日が空欄",IF(NOT(ISNUMBER(D{r})),"提出日が日付でない",'
     'IF(D{r}<>INT(D{r}),"提出日が日付でない",IF(COUNTIFS($A$2:$A${n},A{r},$B$2:$B${n},B{r})>1,"同じ工事の同じ版が重複",'
     'IF(COUNTIFS($A$2:$A${n},A{r},$B$2:$B${n},"<"&B{r},$D$2:$D${n},">"&D{r})>0,"前の版より提出日が前",""))))))))))))')
H = '=IF(AND(G{r}="",A{r}&B{r}&C{r}&D{r}<>""),B{r},"")'
F = '=IF(OR(G{r}<>"",A{r}&B{r}&C{r}&D{r}=""),"",IF(B{r}=_xlfn.MAXIFS($H$2:$H${n},$A$2:$A${n},A{r}),"最新版","旧版"))'
I = ('=IF(F{r}="最新版",C{r}-SUMIFS($C$2:$C${n},$A$2:$A${n},A{r},$H$2:$H${n},_xlfn.MINIFS($H$2:$H${n},$A$2:$A${n},A{r})),"")')
SUM = {"K4": '=COUNTIF($F$2:$F${n},"最新版")', "K5": '=SUMIF($F$2:$F${n},"最新版",$C$2:$C${n})',
       "K6": '=COUNTIF($F$2:$F${n},"旧版")', "K7": '=SUM($I$2:$I${n})', "K8": '=SUMPRODUCT(--(LEN($G$2:$G${n})>0))'}

def build(rows):
    wb = Workbook(); ws = wb.active; ws.title = "見積履歴"
    for i, h in enumerate(["工事番号", "版", "見積額", "提出日", "", "状態", "チェック", "有効な版", "初版からの増減"], 1): ws.cell(1, i, h)
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1): ws.cell(r, c, v)
        ws.cell(r, 7, G.format(r=r, n=N)); ws.cell(r, 8, H.format(r=r)); ws.cell(r, 6, F.format(r=r, n=N)); ws.cell(r, 9, I.format(r=r, n=N))
    for k, f in SUM.items(): ws[k] = f.format(n=N)
    return wb

def run_lo(wb):
    t = tempfile.mkdtemp(); src = os.path.join(t, "t.xlsx"); wb.save(src)
    out = os.path.join(t, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
empty = lambda v: v in (None, "")
key = lambda v: v.lower() if isinstance(v, str) else v
def ser(v):
    if isinstance(v, dt.datetime): return (v - dt.datetime(1899, 12, 30)).total_seconds() / 86400
    if isinstance(v, dt.date): return (v - dt.date(1899, 12, 30)).days
    return v
isdate = lambda v: (isinstance(v, dt.date) or num(v)) and ser(v) == int(ser(v))
dnum = lambda v: isinstance(v, (dt.date, dt.datetime)) or num(v)

def expect(rows):
    chk = []
    for i, (a, b, c, d) in enumerate(rows):
        if all(empty(x) for x in (a, b, c, d)): chk.append(""); continue
        err = ""
        if empty(a): err = "工事番号が空欄"
        elif empty(b): err = "版が空欄"
        elif not num(b): err = "版が数字でない"
        elif b < 1 or b != int(b): err = "版は1以上の整数"
        elif empty(c): err = "見積額が空欄"
        elif not num(c): err = "見積額が数字でない"
        elif c < 1 or c != int(c): err = "見積額は1以上の整数"
        elif empty(d): err = "提出日が空欄"
        elif not dnum(d): err = "提出日が日付でない"
        elif not isdate(d): err = "提出日が日付でない"
        elif sum(1 for x in rows if not empty(x[0]) and key(x[0]) == key(a) and num(x[1]) and x[1] == b) > 1: err = "同じ工事の同じ版が重複"
        elif any(not empty(x[0]) and key(x[0]) == key(a) and num(x[1]) and x[1] < b and dnum(x[3]) and ser(x[3]) > ser(d) for x in rows): err = "前の版より提出日が前"
        chk.append(err)
    valid = {i: r[1] for i, r in enumerate(rows) if chk[i] == "" and not all(empty(x) for x in r)}
    status, diff = [], []
    for i, r in enumerate(rows):
        if i not in valid: status.append(""); diff.append(""); continue
        vs = [valid[j] for j in valid if key(rows[j][0]) == key(r[0])]
        if r[1] == max(vs):
            status.append("最新版")
            first = min(vs); diff.append(r[2] - sum(rows[j][2] for j in valid if key(rows[j][0]) == key(r[0]) and valid[j] == first))
        else: status.append("旧版"); diff.append("")
    summ = (status.count("最新版"), sum(r[2] for r, s in zip(rows, status) if s == "最新版"), status.count("旧版"),
            sum(x for x in diff if x != ""), sum(1 for x in chk if x))
    return list(zip(status, chk, diff)), summ

def cv(v): return "" if v is None else v

def check(rows):
    e, es = expect(rows)
    ws = run_lo(build(rows))["見積履歴"]
    g = [tuple(cv(ws.cell(r, c).value) for c in (6, 7, 9)) for r in range(2, len(rows) + 2)]
    gs = tuple(cv(ws[k].value) for k in ("K4", "K5", "K6", "K7", "K8"))
    return g == e and gs == es, g, e, gs, es

D = lambda n: dt.date(2026, 10, 1) + dt.timedelta(days=n)
S = lambda name, rows: (name, rows)
SCEN = [
 S("1版だけ(最新版)", [("A-001", 1, 1000000, D(0))]),
 S("2版(第1版が旧版、第2版が最新版)", [("A-001", 1, 1000000, D(0)), ("A-001", 2, 1150000, D(7))]),
 S("行の順番が逆(第2版が先)", [("A-001", 2, 1150000, D(7)), ("A-001", 1, 1000000, D(0))]),
 S("3版・金額が下がる(値引き)", [("A-001", 1, 1000000, D(0)), ("A-001", 2, 1150000, D(7)), ("A-001", 3, 980000, D(12))]),
 S("版が飛ぶ(1版と3版)", [("A-001", 1, 500000, D(0)), ("A-001", 3, 560000, D(9))]),
 S("2つの工事が混ざる", [("A-001", 1, 1000000, D(0)), ("A-002", 1, 300000, D(1)), ("A-001", 2, 1100000, D(5)), ("A-002", 2, 280000, D(6))]),
 S("同じ日に第1版と第2版を出す", [("A-001", 1, 700000, D(3)), ("A-001", 2, 720000, D(3))]),
 S("工事番号の大文字小文字違いは同じ工事", [("A-001", 1, 700000, D(0)), ("a-001", 2, 720000, D(4))]),
 S("工事番号が空欄", [(None, 1, 700000, D(0))]),
 S("版が空欄", [("A-001", None, 700000, D(0))]),
 S("版が全角『２』", [("A-001", "２", 700000, D(0))]),
 S("版が0", [("A-001", 0, 700000, D(0))]),
 S("版が小数1.5", [("A-001", 1.5, 700000, D(0))]),
 S("版が負", [("A-001", -1, 700000, D(0))]),
 S("見積額が空欄", [("A-001", 1, None, D(0))]),
 S("見積額が『70万』", [("A-001", 1, "70万", D(0))]),
 S("見積額が0", [("A-001", 1, 0, D(0))]),
 S("見積額が小数", [("A-001", 1, 1000.5, D(0))]),
 S("提出日が空欄", [("A-001", 1, 700000, None)]),
 S("提出日が全角文字『２０２６/１０/５』", [("A-001", 1, 700000, "２０２６/１０/５")]),
 S("提出日に時刻つき", [("A-001", 1, 700000, dt.datetime(2026, 10, 5, 12, 0))]),
 S("同じ工事の同じ版が2行(重複)", [("A-001", 1, 700000, D(0)), ("A-001", 1, 710000, D(1))]),
 S("提出日の逆転(第2版が第1版より前)", [("A-001", 1, 700000, D(10)), ("A-001", 2, 720000, D(3))]),
 S("提出日の逆転で最新版が決まらない行は止まる", [("A-001", 1, 700000, D(0)), ("A-001", 2, 720000, D(5)), ("A-001", 3, 740000, D(2))]),
 S("エラー行は他の工事の判定に影響しない", [("A-001", 1, 700000, D(0)), ("A-002", "x", 100, D(0)), ("A-002", 2, 120000, D(3))]),
 S("空行をはさむ", [(None, None, None, None), ("A-001", 1, 700000, D(0))]),
 S("集計(3工事・最新3・旧版3・チェック1)", [("A-001", 1, 1000000, D(0)), ("A-001", 2, 1150000, D(7)), ("A-002", 1, 300000, D(1)), ("A-002", 2, 280000, D(6)),
                                          ("A-002", 3, 270000, D(9)), ("A-003", 1, 90000, D(2)), (None, 1, 5, D(2))]),
]

def main():
    ng = 0
    for name, rows in SCEN:
        ok, g, e, gs, es = check(rows); ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: {g}")
        if not ok: print("   期待", e, es, "\n   実際", g, gs)
    rnd = random.Random(20261030); rows = []; ver = {}
    for i in range(500):
        k = f"A-{rnd.randint(1, 120):03d}"; ver[k] = ver.get(k, 0) + rnd.choice([1, 1, 1, 2])
        row = [k, ver[k], rnd.randint(10, 300) * 10000, D(ver[k] * 4 + rnd.randint(0, 3))]; x = rnd.random()
        if x < 0.02: row[0] = None
        elif x < 0.04: row[1] = rnd.choice([None, "２", 0, 1.5, -1])
        elif x < 0.06: row[2] = rnd.choice([None, "70万", 0, 1000.5])
        elif x < 0.08: row[3] = rnd.choice([None, "２０２６/１０/５", dt.datetime(2026, 10, 5, 9, 0)])
        elif x < 0.10: row[3] = D(rnd.randint(-30, 0))
        rows.append(tuple(row))
    rows.reverse() if rnd.random() < 0.5 else rnd.shuffle(rows)
    ok, g, e, gs, es = check(rows); ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: {len(rows)}行 全項目一致={ok} 集計={gs}")
    if not ok:
        for i, (x, y) in enumerate(zip(g, e)):
            if x != y: print("  行", i + 2, rows[i], x, y); break
        print("  集計", gs, es)
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
