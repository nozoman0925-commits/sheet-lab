"""記事『作業日報から工事ごとの人工と労務費を集計する式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/daily_labor.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random, datetime, math
from openpyxl import Workbook, load_workbook

N = 3000  # 式の範囲 2:3000
MAXD = 2  # 同じ人・同じ日の合計の上限(人工)

def build(jobs, workers, rows):
    """jobs: [番号], workers: [(氏名, 日当)], rows: [(日付, 番号, 氏名, 人工)] (セルの値そのまま)"""
    wb = Workbook(); led = wb.active; led.title = "台帳"; wk = wb.create_sheet("作業員"); dr = wb.create_sheet("日報")
    for i, h in enumerate(["工事番号", "人工合計", "労務費"], 1): led.cell(1, i, h)
    for i, h in enumerate(["氏名", "日当"], 1): wk.cell(1, i, h)
    for i, h in enumerate(["日付", "工事番号", "作業員", "人工", "整えた番号", "日当", "労務費", "チェック"], 1): dr.cell(1, i, h)
    for r, no in enumerate(jobs, 2):
        led.cell(r, 1, no)
        led.cell(r, 2, f'=IF(A{r}="","",SUMIFS(日報!$D$2:$D${N},日報!$E$2:$E${N},A{r},日報!$D$2:$D${N},">0"))')
        led.cell(r, 3, f'=IF(A{r}="","",SUMIFS(日報!$G$2:$G${N},日報!$E$2:$E${N},A{r}))')
    for r, (nm, w) in enumerate(workers, 2):
        wk.cell(r, 1, nm); wk.cell(r, 2, w)
    for r, (d, no, nm, h) in enumerate(rows, 2):
        dr.cell(r, 1, d); dr.cell(r, 2, no); dr.cell(r, 3, nm); dr.cell(r, 4, h)
        dr.cell(r, 5, f'=ASC(TRIM(B{r}))')
        dr.cell(r, 6, f'=IF(TRIM(C{r})="","",IFERROR(INDEX(作業員!$B$2:$B$200,MATCH(TRIM(C{r}),作業員!$A$2:$A$200,0)),""))')
        dr.cell(r, 7, f'=IF(AND(ISNUMBER(D{r}),D{r}>0,ISNUMBER(F{r})),ROUND(D{r}*F{r},0),0)')
        dr.cell(r, 8, (f'=IF(AND(B{r}="",C{r}="",D{r}=""),"",IF(NOT(ISNUMBER(D{r})),"人工が未入力か数値でない",IF(D{r}<=0,"人工が0以下",'
                       f'IF(B{r}="","工事番号が空欄",IF(ISNA(MATCH(E{r},台帳!$A$2:$A${N},0)),"工事番号なし",IF(F{r}="","作業員が未登録",'
                       f'IF(AND(A{r}<>"",SUMIFS($D$2:$D${N},$A$2:$A${N},A{r},$C$2:$C${N},C{r})>{MAXD}),"同日の合計が{MAXD}人工超","")))))))'))
    led["E1"], led["F1"] = "未紐づけの労務費", f'=SUM(日報!G2:G{N})-SUM(C2:C{N})'
    return wb

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

def fw(s):
    return "".join(chr(ord(c) + 0xFEE0) if "!" <= c <= "~" and c != "-" else ("－" if c == "-" else c) for c in s)

def norm(v):
    if v is None: return ""
    return "".join(chr(ord(c) - 0xFEE0) if "！" <= c <= "～" else ("-" if c == "－" else c) for c in str(v).strip().replace("　", " ").strip())

def isnum(v): return isinstance(v, (int, float)) and not isinstance(v, bool)

def expected(jobs, workers, rows):
    wage = {nm: w for nm, w in workers}; jobset = set(jobs)
    hours = {j: 0 for j in jobs}; cost = {j: 0 for j in jobs}; flags = []; unlinked = 0; total_cost = 0
    for d, no, nm, h in rows:
        k = norm(no); nmt = (nm or "").strip(); w = wage.get(nmt)
        lab = math.floor(h * w + 0.5) if isnum(h) and h > 0 and w is not None else 0
        total_cost += lab
        if no in (None, "") and nm in (None, "") and h in (None, ""): flags.append(""); continue
        if not isnum(h): flags.append("人工が未入力か数値でない")
        elif h <= 0: flags.append("人工が0以下")
        elif no in (None, ""): flags.append("工事番号が空欄")
        elif k not in jobset: flags.append("工事番号なし")
        elif w is None: flags.append("作業員が未登録")
        else:
            day = sum(x[3] for x in rows if d is not None and x[0] == d and x[2] == nm and isnum(x[3]))
            flags.append(f"同日の合計が{MAXD}人工超" if d is not None and day > MAXD else "")
        if k in jobset:
            if isnum(h) and h > 0: hours[k] += h
            cost[k] += lab
    linked = sum(cost.values())
    return hours, cost, flags, total_cost - linked

D = datetime.date
W = [("職人A", 18000), ("職人B", 16000), ("職人C", 15001)]
SCEN = [
 ("通常(1工事に3人・半日含む)", ["K-001"], [(D(2026,10,1),"K-001","職人A",1),(D(2026,10,1),"K-001","職人B",0.5),(D(2026,10,2),"K-001","職人C",1)]),
 ("1日に2現場(0.5+0.5)", ["K-002","K-003"], [(D(2026,10,1),"K-002","職人A",0.5),(D(2026,10,1),"K-003","職人A",0.5)]),
 ("存在しない工事番号", ["K-004"], [(D(2026,10,1),"K-040","職人A",1)]),
 ("工事番号が空欄", ["K-005"], [(D(2026,10,1),None,"職人A",1)]),
 ("全角・スペースの工事番号", ["K-006"], [(D(2026,10,1),"Ｋ－006","職人A",1),(D(2026,10,2)," K-006 ","職人B",1)]),
 ("人工が文字(全角の1・「半日」)", ["K-007"], [(D(2026,10,1),"K-007","職人A","１"),(D(2026,10,2),"K-007","職人B","半日")]),
 ("人工が空欄・0・マイナス", ["K-008"], [(D(2026,10,1),"K-008","職人A",None),(D(2026,10,2),"K-008","職人B",0),(D(2026,10,3),"K-008","職人C",-1)]),
 ("未登録の作業員", ["K-009"], [(D(2026,10,1),"K-009","職人Z",1)]),
 ("同じ人・同じ日の合計が2人工超", ["K-010","K-011"], [(D(2026,10,1),"K-010","職人A",1.5),(D(2026,10,1),"K-011","職人A",1)]),
 ("端数が出る日当(15,001円×0.5)", ["K-012"], [(D(2026,10,1),"K-012","職人C",0.5)]),
 ("日報なし", ["K-013"], []),
]

def main():
    ng = 0
    for name, jobs, rows in SCEN:
        res = run_lo(build(jobs, W, rows)); led, dr = res["台帳"], res["日報"]
        hrs, cost, flags, unl = expected(jobs, W, rows)
        gh = {norm(led.cell(r,1).value): (led.cell(r,2).value or 0) for r in range(2, len(jobs)+2)}
        gc = {norm(led.cell(r,1).value): (led.cell(r,3).value or 0) for r in range(2, len(jobs)+2)}
        gf = [(dr.cell(r, 8).value or "") for r in range(2, len(rows)+2)]
        gu = led["F1"].value or 0
        ok = gh == hrs and gc == cost and gf == flags and gu == unl
        ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: 人工{gh} 労務費{gc} チェック{gf} 未紐づけ{gu}")
        if not ok: print("   期待", hrs, cost, flags, unl)
    rnd = random.Random(20261021)
    jobs = [f"K-{i:03d}" for i in range(1, 301)]
    workers = [(f"職人{chr(65+i)}", 12000 + 500 * i + (1 if i % 3 == 0 else 0)) for i in range(20)]
    rows = []
    pairs = [(w[0], D(2026,8,1) + datetime.timedelta(days=k)) for w in workers for k in range(90)]
    rnd.shuffle(pairs)
    for w, d in pairs[:1500]:
        t = rnd.random(); j = rnd.choice(jobs); h = rnd.choice([0.5, 1, 1, 1, 1])
        if t < 0.82: rows.append((d, j, w, h))
        elif t < 0.87: rows.append((d, "K-9%02d" % rnd.randint(0, 99), w, h))
        elif t < 0.90: rows.append((d, None, w, h))
        elif t < 0.94: rows.append((d, fw(j), w, h))
        elif t < 0.96: rows.append((d, " " + j + " ", w, h))
        elif t < 0.98: rows.append((d, j, w, rnd.choice(["１", "半日", None, 0, -1])))
        else: rows.append((d, j, "未登録の人", h))
    for w, d in pairs[:40]:  # 同じ人・同じ日に2現場(合計が2人工を超える行を意図的に混ぜる)
        rows.append((d, rnd.choice(jobs), w, 1.5))
    res = run_lo(build(jobs, workers, rows)); led, dr = res["台帳"], res["日報"]
    hrs, cost, flags, unl = expected(jobs, workers, rows)
    gh = {norm(led.cell(r,1).value): (led.cell(r,2).value or 0) for r in range(2, 302)}
    gc = {norm(led.cell(r,1).value): (led.cell(r,3).value or 0) for r in range(2, 302)}
    gf = [(dr.cell(r, 8).value or "") for r in range(2, len(rows)+2)]
    ok = gh == hrs and gc == cost and gf == flags and (led["F1"].value or 0) == unl
    ng += not ok
    from collections import Counter
    print(("OK " if ok else "NG ") + f"ランダム照合: 工事300件・作業員20人・日報約1,540行 一致={ok}", dict(Counter(flags)))
    if not ok:
        print([ (j, gh[j], hrs[j]) for j in jobs if gh[j]!=hrs[j]][:5], sum(1 for a,b in zip(gf,flags) if a!=b), led["F1"].value, unl)
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
