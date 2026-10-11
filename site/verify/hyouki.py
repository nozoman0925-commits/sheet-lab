"""表記ゆれの統一(記事: hyouki-yure-kansu)と、COPILOT関数記事の対比例を LibreOffice で再計算して確認する。"""
import subprocess, json, sys
from openpyxl import Workbook, load_workbook
wb = Workbook(); ws = wb.active
A = ["(株)山田商事", "（株）山田商事", "㈱山田商事", "株式会社 山田商事", "Ａ１２３-ｂｃ", "ﾔﾏﾀﾞ ﾀﾛｳ", "山田　太郎", " 山田商事 "]
for i, v in enumerate(A, 1): ws.cell(row=i, column=1, value=v)
for i in range(1, len(A)+1):
    ws.cell(row=i, column=2, value=f'=ASC(A{i})')            # 全角→半角
    ws.cell(row=i, column=3, value=f'=JIS(A{i})')            # 半角→全角
    ws.cell(row=i, column=4, value=f'=TRIM(SUBSTITUTE(A{i},"　"," "))')
    ws.cell(row=i, column=5, value=f'=SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(A{i},"(株)","株式会社"),"（株）","株式会社"),"㈱","株式会社")')
    ws.cell(row=i, column=6, value=f'=EXACT(E{i},E1)')
ws["H1"] = '=COUNTIF(E1:E4,"株式会社山田商事")'
ws["H2"] = '=SUMPRODUCT(--(E1:E4="株式会社山田商事"))'
wb.save("hyouki.xlsx")
out = subprocess.run(["python3", "/mnt/skills/public/xlsx/scripts/recalc.py", "hyouki.xlsx", "90"], capture_output=True, text=True).stdout
print(json.loads(out)["status"])
ws = load_workbook("hyouki.xlsx", data_only=True).active
for r in range(1, len(A)+1):
    print([ws.cell(row=r, column=c).value for c in range(1, 7)])
print("H1,H2 =", ws["H1"].value, ws["H2"].value)
