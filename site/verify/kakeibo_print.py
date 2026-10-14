"""記事『家計簿をエクセルで作ってA4・1枚に印刷する方法』の式と印刷設定をLibreOfficeで確認する。
使い方: python3 site/verify/kakeibo_print.py"""
import subprocess, sys, tempfile, os, random, re
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.properties import PageSetupProperties

random.seed(5)
CATS = ["食費", "日用品", "交通費", "光熱費", "通信費", "趣味"]
ROWS = 40  # 明細行数(2〜41行)
data = [(random.choice(CATS), random.randint(100, 5000)) for _ in range(30)]  # 30行入力、10行は空欄

def build(path, fit):
    wb = Workbook(); ws = wb.active
    ws.append(["日付", "費目", "内容", "支出", "残高"])
    for i in range(ROWS):
        r = i + 2
        if i < len(data):
            ws.cell(r, 1, f"2026/10/{i % 28 + 1}"); ws.cell(r, 2, data[i][0]); ws.cell(r, 3, "メモ"); ws.cell(r, 4, data[i][1])
        # 残高 = 予算 - 累計(H1が月の予算)。空欄行は空文字にする
        ws.cell(r, 5, f'=IF(D{r}="","",$H$1-SUM($D$2:D{r}))')
    ws["G1"] = "月の予算"; ws["H1"] = 100000
    for k, c in enumerate(CATS):
        ws.cell(k + 3, 7, c); ws.cell(k + 3, 8, f'=SUMIFS($D$2:$D$41,$B$2:$B$41,G{k+3})')
    ws["G9"] = "合計"; ws["H9"] = "=SUM(D2:D41)"
    for col, w in zip("ABCDEFGH", (12, 10, 40, 12, 12, 4, 14, 12)):
        ws.column_dimensions[col].width = w
    # 幅を意図的に広げ、設定なしでは1ページに収まらない状態を作る
    ws.column_dimensions["C"].width = 70 if not fit else 40
    ws.print_area = "A1:H41"
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.orientation = "portrait"
    if fit:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    wb.save(path)

def pages(xlsx, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, xlsx], check=True, capture_output=True)
    pdf = os.path.join(outdir, os.path.basename(xlsx).replace(".xlsx", ".pdf"))
    info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
    return int(re.search(r"Pages:\s+(\d+)", info).group(1)), re.search(r"Page size:\s+(.*)", info).group(1)

def main():
    d = tempfile.mkdtemp(); ng = 0; n = 0
    # 1. 式の確認
    p = os.path.join(d, "calc.xlsx"); build(p, True)
    out = os.path.join(d, "o"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, p], check=True, capture_output=True)
    ws = load_workbook(os.path.join(out, "calc.xlsx"), data_only=True).active
    run = 0
    for i in range(ROWS):
        r = i + 2
        if i < len(data): run += data[i][1]; exp = 100000 - run
        else: exp = None
        got = ws.cell(r, 5).value
        got = None if got in ("", None) else got
        n += 1; ng += got != exp
    for k, c in enumerate(CATS):
        e = sum(v for cc, v in data if cc == c); g = ws.cell(k + 3, 8).value
        n += 1; ng += g != e; print(("OK " if g == e else "NG ") + f"{c} 合計 {g} (期待 {e})")
    tot = sum(v for _, v in data); g = ws["H9"].value; n += 1; ng += g != tot
    print(f"合計 {g} (期待 {tot}) / 最終行の残高 {ws.cell(31, 5).value}(期待 {100000 - tot})")
    # 2. 印刷ページ数
    p0 = os.path.join(d, "nofit.xlsx"); build(p0, False)
    p1 = os.path.join(d, "fit.xlsx"); build(p1, True)
    o2 = os.path.join(d, "pdf"); os.mkdir(o2)
    n0, s0 = pages(p0, o2); n1, s1 = pages(p1, o2)
    print(f"拡大縮小なし(列Cを広く): {n0} ページ / {s0}")
    print(f"1ページに収める設定: {n1} ページ / {s1}")
    n += 1; ng += n1 != 1
    print(f"集計 {n} 件、不一致 {ng} 件"); sys.exit(1 if ng else 0)
main()
