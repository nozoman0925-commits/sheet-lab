"""記事『見積番号を自動で振る』をLibreOfficeで再計算して期待値と照合する。
使い方: python3 site/verify/estimate_number.py"""
import subprocess, sys, tempfile, os, datetime, random
from openpyxl import Workbook, load_workbook

def build_dates():
    d = datetime.date
    base = [d(2026,9,28), d(2026,9,30), d(2026,10,1), d(2026,10,1), d(2026,10,15), d(2026,10,31),
            d(2026,11,1), d(2026,11,2), d(2026,12,31), d(2027,1,1), d(2027,1,1), d(2024,2,29),
            d(2024,2,29), d(2024,3,1)]
    return base
def exp_numbers(dates):
    cnt = {}; out_month = []; out_serial = []
    for i, x in enumerate(dates, 1):
        k = (x.year, x.month); cnt[k] = cnt.get(k, 0) + 1
        out_month.append(f"Q{x.year}{x.month:02d}-{cnt[k]:03d}")
        out_serial.append(f"Q-{i:04d}")
    return out_month, out_serial
def run(dates, blank_rows=()):
    wb = Workbook(); ws = wb.active
    for i, x in enumerate(dates, 2):
        if i in blank_rows: continue
        ws.cell(i, 1, x).number_format = "yyyy/mm/dd"
        # B: 月ごとに001から振り直す / C: 通し番号 / D: 重複チェック
        ws.cell(i, 2, f'=IF(A{i}="","","Q"&TEXT(A{i},"yyyymm")&"-"&TEXT(COUNTIFS($A$2:A{i},">="&DATE(YEAR(A{i}),MONTH(A{i}),1),$A$2:A{i},"<="&EOMONTH(A{i},0)),"000"))')
        ws.cell(i, 3, f'=IF(A{i}="","","Q-"&TEXT(ROW()-1,"0000"))')
        ws.cell(i, 4, f'=IF(B{i}="","",IF(COUNTIF($B$2:$B${len(dates)+1},B{i})>1,"重複","OK"))')
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "o"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True).active
def main():
    ng = 0; n = 0
    dates = build_dates(); em, es = exp_numbers(dates)
    r = run(dates)
    for i, x in enumerate(dates):
        got = (r.cell(i+2, 2).value, r.cell(i+2, 3).value, r.cell(i+2, 4).value)
        exp = (em[i], es[i], "OK"); ok = got == exp; ng += not ok; n += 3
        print(("OK " if ok else "NG ") + f"{x} -> {got[0]} / {got[1]} / {got[2]}")
        if not ok: print("  期待", exp)
    # 空行がある場合(3行目を空にする)
    r2 = run(dates, blank_rows=(3,))
    d2 = [x for j, x in enumerate(dates) if j+2 != 3]
    em2, _ = exp_numbers(d2)
    got2 = [r2.cell(j+2, 2).value for j in range(len(dates)) if j+2 != 3]
    ok = got2 == em2; ng += not ok; n += len(em2)
    print("空行あり:", "OK" if ok else "NG", r2.cell(3, 2).value in (None, ""))
    # 並べ替え(新しい順)で番号が変わることを確認(月別は日付順に並べないと崩れる)
    rev = list(reversed(dates)); emr, _ = exp_numbers(rev)
    r3 = run(rev); gotr = [r3.cell(j+2, 2).value for j in range(len(rev))]
    ok = gotr == emr; ng += not ok; n += len(emr)
    print("逆順に並べた表:", "OK" if ok else "NG", "(式は上から数えるので、並べ替えると番号が変わる)", gotr[:3])
    # 同じ月に1000件: 連番が4桁になる
    big = [datetime.date(2026, 10, 1)] * 1001
    emb, _ = exp_numbers(big); r4 = run(big)
    gotb = [r4.cell(j+2, 2).value for j in range(len(big))]
    ok = gotb == emb; ng += not ok; n += len(emb)
    print("1001件:", "OK" if ok else "NG", gotb[998], gotb[999], gotb[1000])
    # ランダム500件(2026-2027を昇順)
    random.seed(7)
    rd = sorted(datetime.date(2026,1,1)+datetime.timedelta(random.randrange(0,730)) for _ in range(500))
    emr2, _ = exp_numbers(rd); r5 = run(rd)
    gotr2 = [r5.cell(j+2, 2).value for j in range(len(rd))]
    ok = gotr2 == emr2; ng += not ok; n += len(rd)
    print("ランダム500件:", "OK" if ok else "NG")
    print("検証", n, "件 不一致", ng); sys.exit(1 if ng else 0)
main()
