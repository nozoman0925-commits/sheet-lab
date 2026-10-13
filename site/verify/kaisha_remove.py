"""記事『会社名から「株式会社」「(株)」を取り除いて社名だけにする』の数式をLibreOfficeで再計算して期待値と照合する。
使い方: python3 site/verify/kaisha_remove.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os
from openpyxl import Workbook, load_workbook

# 入力, B列(全部の法人格表記を削除+スペース整理), C列(要確認判定付き), G列(先頭と末尾だけ削除。D〜Fは途中経過)
CASES = [
    ("株式会社山田",         "山田",       "山田",       "山田"),
    ("山田株式会社",         "山田",       "山田",       "山田"),
    ("(株)山田",             "山田",       "山田",       "山田"),
    ("山田(株)",             "山田",       "山田",       "山田"),
    ("（株）山田",           "山田",       "山田",       "山田"),
    ("㈱山田",               "山田",       "山田",       "山田"),
    ("有限会社山田商店",     "山田商店",   "山田商店",   "山田商店"),
    ("株式会社 山田",        "山田",       "山田",       "山田"),
    ("株式会社　山田　",     "山田",       "山田",       "山田"),
    ("山田商事",             "山田商事",   "山田商事",   "山田商事"),
    ("株式会社",             "",           "要確認",     "要確認"),
    ("日本株式会社サービス", "日本サービス", "日本サービス", "日本株式会社サービス"),
    ("",                     "",           "",           ""),
]
CORP = ["株式会社", "有限会社", "(株)", "（株）", "㈱", "(有)", "（有）", "㈲"]

def chain(ref):
    s = ref + '&""'
    for w in CORP: s = f'SUBSTITUTE({s},"{w}","")'
    return s

def main():
    wb = Workbook(); ws = wb.active
    for r, (inp, *_) in enumerate(CASES, 1):
        ws.cell(r, 1, inp or None); ws.cell(r, 1).number_format = "@"
        ws.cell(r, 2, f'=TRIM(SUBSTITUTE({chain(f"A{r}")},"　"," "))')
        ws.cell(r, 3, f'=IF(A{r}="","",IF(B{r}="","要確認",B{r}))')
        # 先頭・末尾の法人格だけ削除(D=スペース整理、E=先頭、F=末尾、G=結果)
        ws.cell(r, 4, f'=TRIM(SUBSTITUTE(A{r}&"","　"," "))')
        ws.cell(r, 5, f'=IF(OR(LEFT(D{r},4)="株式会社",LEFT(D{r},4)="有限会社"),MID(D{r},5,99),IF(OR(LEFT(D{r},3)="(株)",LEFT(D{r},3)="（株）",LEFT(D{r},3)="(有)",LEFT(D{r},3)="（有）"),MID(D{r},4,99),IF(OR(LEFT(D{r},1)="㈱",LEFT(D{r},1)="㈲"),MID(D{r},2,99),D{r})))')
        ws.cell(r, 6, f'=IF(OR(RIGHT(E{r},4)="株式会社",RIGHT(E{r},4)="有限会社"),LEFT(E{r},LEN(E{r})-4),IF(OR(RIGHT(E{r},3)="(株)",RIGHT(E{r},3)="（株）",RIGHT(E{r},3)="(有)",RIGHT(E{r},3)="（有）"),LEFT(E{r},LEN(E{r})-3),IF(OR(RIGHT(E{r},1)="㈱",RIGHT(E{r},1)="㈲"),LEFT(E{r},LEN(E{r})-1),E{r})))')
        ws.cell(r, 7, f'=IF(A{r}="","",IF(TRIM(F{r})="","要確認",TRIM(F{r})))')
    dd = tempfile.mkdtemp(); src = os.path.join(dd, "t.xlsx"); wb.save(src)
    out = os.path.join(dd, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    res = load_workbook(os.path.join(out, "t.xlsx"), data_only=True).active
    ng = 0
    for r, (inp, *exp) in enumerate(CASES, 1):
        for c, e in zip((2, 3, 7), exp):
            got = res.cell(r, c).value or ""
            ok = got == e; ng += not ok
            print(("OK " if ok else "NG ") + f"{inp!r} col{c}: 期待{e!r} 結果{got!r}")
    print("不一致", ng, "件"); sys.exit(1 if ng else 0)
main()
