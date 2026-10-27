"""記事『工事の実行予算と原価実績を比べて予算超過・原価の先行を見つけるExcelの式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/budget_actual.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random
from openpyxl import Workbook, load_workbook

N = 1000
G = ('=IF(A{r}&B{r}&C{r}&D{r}&E{r}="","",IF(A{r}="","工事番号が空欄",IF(COUNTIF($A$2:$A${n},A{r})>1,"工事番号が重複",'
     'IF(C{r}="","実行予算が空欄",IF(NOT(ISNUMBER(C{r})),"実行予算が数字でない",IF(OR(C{r}<1,C{r}<>INT(C{r})),"実行予算は1以上の整数",'
     'IF(D{r}="","原価実績が空欄",IF(NOT(ISNUMBER(D{r})),"原価実績が数字でない",IF(OR(D{r}<0,D{r}<>INT(D{r})),"原価実績は0以上の整数",'
     'IF(E{r}="","進捗率が空欄",IF(NOT(ISNUMBER(E{r})),"進捗率が数字でない",IF(OR(E{r}<0,E{r}>1),"進捗率は0〜100%",""))))))))))))')
F = '=IF(OR(G{r}<>"",A{r}&B{r}&C{r}&D{r}&E{r}=""),"",IF(D{r}>C{r},"予算超過",IF(ROUND(D{r}/C{r}-E{r}-N($N$2),6)>0,"要注意","順調")))'
H = '=IF(F{r}="","",D{r}/C{r})'
I = '=IF(F{r}="","",C{r}-D{r})'
J = '=IF(OR(F{r}="",N(E{r})=0),"",D{r}/E{r})'
K = '=IF(J{r}="","",J{r}-C{r})'
SUM = {"N5": '=COUNTIF($F$2:$F${n},"予算超過")', "N6": '=-SUMIF($F$2:$F${n},"予算超過",$I$2:$I${n})',
       "N7": '=COUNTIF($F$2:$F${n},"要注意")', "N8": '=COUNTIF($F$2:$F${n},"順調")'}

def build(rows, tol):
    wb = Workbook(); ws = wb.active; ws.title = "予算実績"
    for i, h in enumerate(["工事番号", "工事名", "実行予算", "原価実績", "進捗率", "状態", "チェック", "消化率", "予算残", "完成時の見込原価", "見込差額"], 1): ws.cell(1, i, h)
    ws["M2"] = "許容ポイント"; ws["N2"] = tol
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1): ws.cell(r, c, v)
        ws.cell(r, 7, G.format(r=r, n=N)); ws.cell(r, 6, F.format(r=r))
        ws.cell(r, 8, H.format(r=r)); ws.cell(r, 9, I.format(r=r)); ws.cell(r, 10, J.format(r=r)); ws.cell(r, 11, K.format(r=r))
    for k, f in SUM.items(): ws[k] = f.format(n=N)
    return wb

def run_lo(wb):
    t = tempfile.mkdtemp(); src = os.path.join(t, "t.xlsx"); wb.save(src)
    out = os.path.join(t, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
empty = lambda v: v in (None, "")
key = lambda v: v.lower() if isinstance(v, str) else v
BLANK = ("", "", "", "", "", "")

def expect(rows, tol):
    t = tol if num(tol) else 0
    out = []
    for a, b, c, d, e in rows:
        if all(empty(x) for x in (a, b, c, d, e)): out.append(BLANK); continue
        err = None
        if empty(a): err = "工事番号が空欄"
        elif sum(1 for x in rows if not empty(x[0]) and key(x[0]) == key(a)) > 1: err = "工事番号が重複"
        elif empty(c): err = "実行予算が空欄"
        elif not num(c): err = "実行予算が数字でない"
        elif c < 1 or c != int(c): err = "実行予算は1以上の整数"
        elif empty(d): err = "原価実績が空欄"
        elif not num(d): err = "原価実績が数字でない"
        elif d < 0 or d != int(d): err = "原価実績は0以上の整数"
        elif empty(e): err = "進捗率が空欄"
        elif not num(e): err = "進捗率が数字でない"
        elif e < 0 or e > 1: err = "進捗率は0〜100%"
        if err: out.append(("", "", "", "", "", err)); continue
        if d > c: st = "予算超過"
        elif round(d / c - e - t, 6) > 0: st = "要注意"
        else: st = "順調"
        fc = d / e if e else ""
        out.append((st, d / c, c - d, fc, (fc - c) if e else "", ""))
    return out

def exp_sum(rows, ex):
    over = [r for r, x in zip(rows, ex) if x[0] == "予算超過"]
    return (len(over), sum(r[3] - r[2] for r in over), sum(1 for x in ex if x[0] == "要注意"), sum(1 for x in ex if x[0] == "順調"))

def read(wb, n):
    ws = wb["予算実績"]; res = []
    for r in range(2, n + 2):
        res.append(tuple("" if ws.cell(r, c).value is None else ws.cell(r, c).value for c in (6, 8, 9, 10, 11, 7)))
    return res, tuple("" if ws[k].value is None else ws[k].value for k in ("N5", "N6", "N7", "N8"))

def same(x, y):
    if len(x) != len(y): return False
    for p, q in zip(x, y):
        if num(p) and num(q):
            if abs(p - q) > 1e-6: return False
        elif p != q: return False
    return True

def check(rows, tol):
    e = expect(rows, tol); es = exp_sum(rows, e)
    g, gs = read(run_lo(build(rows, tol)), len(rows))
    ok = all(same(a, b) for a, b in zip(g, e)) and same(gs, es)
    return ok, g, e, gs, es

S = lambda name, rows, tol=0.1: (name, rows, tol)
SCEN = [
 S("順調(消化率40%・進捗50%)", [("K-001", "山田邸", 1000000, 400000, 0.5)]),
 S("進捗と消化率が同じ(順調)", [("K-001", "山田邸", 1000000, 500000, 0.5)], 0),
 S("許容ちょうど(消化率60%・進捗50%・許容10ポイント=順調)", [("K-001", "山田邸", 1000000, 600000, 0.5)]),
 S("許容を超えた(消化率61%・進捗50%・許容10ポイント=要注意)", [("K-001", "山田邸", 1000000, 610000, 0.5)]),
 S("許容0で消化率が1円分だけ先行(要注意)", [("K-001", "山田邸", 1000000, 500001, 0.5)], 0),
 S("予算ちょうど(実績=予算・進捗100%=順調)", [("K-001", "山田邸", 1000000, 1000000, 1)]),
 S("予算ちょうどで進捗80%(要注意)", [("K-001", "山田邸", 1000000, 1000000, 0.8)]),
 S("予算超過(実績が予算を1円超える)", [("K-001", "山田邸", 1000000, 1000001, 1)]),
 S("予算超過・進捗30%", [("K-001", "山田邸", 1000000, 1300000, 0.3)]),
 S("着工前(進捗0%・実績0=順調・見込は空)", [("K-001", "山田邸", 1000000, 0, 0)]),
 S("進捗0%で実績5%(許容10ポイント内=順調・見込は空)", [("K-001", "山田邸", 1000000, 50000, 0)]),
 S("進捗0%で実績5%(許容0=要注意・見込は空)", [("K-001", "山田邸", 1000000, 50000, 0)], 0),
 S("進捗100%・予算内(順調・見込=実績)", [("K-001", "山田邸", 1000000, 900000, 1)]),
 S("許容ポイントが空欄(0として扱う)", [("K-001", "山田邸", 1000000, 510000, 0.5)], None),
 S("許容ポイントが文字(0として扱う)", [("K-001", "山田邸", 1000000, 510000, 0.5)], "10%"),
 S("工事番号が空欄", [(None, "山田邸", 1000000, 400000, 0.5)]),
 S("工事番号が重複(大文字小文字違い)", [("K-001", "山田邸", 1000000, 400000, 0.5), ("k-001", "佐藤邸", 800000, 100000, 0.2)]),
 S("実行予算が空欄", [("K-001", "山田邸", None, 400000, 0.5)]),
 S("実行予算が文字『100万』", [("K-001", "山田邸", "100万", 400000, 0.5)]),
 S("実行予算が全角『１，０００，０００』", [("K-001", "山田邸", "１，０００，０００", 400000, 0.5)]),
 S("実行予算が0", [("K-001", "山田邸", 0, 400000, 0.5)]),
 S("実行予算が小数", [("K-001", "山田邸", 1000.5, 400, 0.5)]),
 S("原価実績が空欄", [("K-001", "山田邸", 1000000, None, 0.5)]),
 S("原価実績が文字『40万』", [("K-001", "山田邸", 1000000, "40万", 0.5)]),
 S("原価実績がマイナス", [("K-001", "山田邸", 1000000, -1, 0.5)]),
 S("原価実績が小数", [("K-001", "山田邸", 1000000, 400000.5, 0.5)]),
 S("進捗率が空欄", [("K-001", "山田邸", 1000000, 400000, None)]),
 S("進捗率が文字『半分』", [("K-001", "山田邸", 1000000, 400000, "半分")]),
 S("進捗率を『50』と入れた(5000%)", [("K-001", "山田邸", 1000000, 400000, 50)]),
 S("進捗率がマイナス", [("K-001", "山田邸", 1000000, 400000, -0.1)]),
 S("進捗率が100%超(1.2)", [("K-001", "山田邸", 1000000, 400000, 1.2)]),
 S("集計(予算超過2件・要注意1件・順調2件)", [("K-001", "山田邸", 1000000, 1200000, 0.9), ("K-002", "佐藤邸", 500000, 520000, 1), ("K-003", "鈴木邸", 800000, 500000, 0.4), ("K-004", "高橋邸", 600000, 100000, 0.3), ("K-005", "田中邸", 900000, 300000, 0.5)]),
 S("空行をはさむ", [(None, None, None, None, None), ("K-001", "山田邸", 1000000, 400000, 0.5)]),
]

def main():
    ng = 0
    for name, rows, tol in SCEN:
        ok, g, e, gs, es = check(rows, tol); ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: {g} 集計={gs}")
        if not ok: print("   期待", e, es)
    rnd = random.Random(20261028); rows = []
    for i in range(500):
        a = f"K-{i:03d}"; b = f"邸{i}"; c = rnd.randint(10, 500) * 10000
        e = rnd.choice([0, 1, round(rnd.random(), 2)])
        d = int(c * rnd.uniform(0, 1.15)); x = rnd.random()
        if x < 0.02: a = None
        elif x < 0.04: a = rows[i - 1][0] if i else a
        elif x < 0.06: c = rnd.choice([None, "100万", 0, 1000.5, "１０００"])
        elif x < 0.08: d = rnd.choice([None, "40万", -5, 1000.5])
        elif x < 0.12: e = rnd.choice([None, "半分", 50, -0.1, 1.2])
        rows.append((a, b, c, d, e))
    ok, g, e, gs, es = check(rows, 0.1); ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: {len(rows)}行 全項目一致={ok} 集計={gs}")
    if not ok:
        for i, (x, y) in enumerate(zip(g, e)):
            if not same(x, y): print("  行", i + 2, rows[i], x, y); break
        print("  集計", gs, es)
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
