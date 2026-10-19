"""記事『工事の工程表で遅れている工程を自動で表示する式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/schedule_delay.py   (soffice が必要。不一致なら終了コード1)
基準日はTODAY()ではなく固定のセル(K1)に入れて検証する(本番ではK1に =TODAY() を入れる)。"""
import subprocess, sys, tempfile, os, random, datetime
from openpyxl import Workbook, load_workbook

N = 2000
D = datetime.date

def build(projects, tasks, today):
    """projects: [番号], tasks: [(番号, 工種, 予定開始, 予定終了, 実績終了)]  日付は date / 文字列 / None"""
    wb = Workbook(); led = wb.active; led.title = "台帳"; sc = wb.create_sheet("工程")
    for i, h in enumerate(["工事番号", "遅れ工程数", "遅れ日数の合計"], 1):
        led.cell(1, i, h)
    for i, h in enumerate(["工事番号", "工種", "予定開始", "予定終了", "実績終了", "整えた番号", "チェック", "遅れ日数", "状態"], 1):
        sc.cell(1, i, h)
    for r, no in enumerate(projects, 2):
        led.cell(r, 1, no)
        led.cell(r, 2, f'=IF(A{r}="","",COUNTIFS(工程!$F$2:$F${N},A{r},工程!$H$2:$H${N},">0"))')
        led.cell(r, 3, f'=IF(A{r}="","",SUMIFS(工程!$H$2:$H${N},工程!$F$2:$F${N},A{r}))')
    for r, (no, kind, s, e, a) in enumerate(tasks, 2):
        sc.cell(r, 1, no); sc.cell(r, 2, kind); sc.cell(r, 3, s); sc.cell(r, 4, e); sc.cell(r, 5, a)
        sc.cell(r, 6, f'=ASC(TRIM(A{r}))')
        sc.cell(r, 7, (f'=IF(AND(A{r}="",C{r}="",D{r}=""),"",IF(A{r}="","工事番号が空欄",IF(ISNA(MATCH(F{r},台帳!$A$2:$A${N},0)),"工事番号なし",'
                       f'IF(OR(NOT(ISNUMBER(C{r})),NOT(ISNUMBER(D{r}))),"予定日が未入力か日付でない",IF(D{r}<C{r},"日付が逆転",'
                       f'IF(AND(E{r}<>"",NOT(ISNUMBER(E{r}))),"実績日が日付でない",""))))))'))
        sc.cell(r, 8, f'=IF(G{r}<>"","",IF(AND(A{r}="",C{r}="",D{r}=""),"",IF(E{r}<>"",MAX(0,E{r}-D{r}),MAX(0,$K$1-D{r}))))')
        sc.cell(r, 9, (f'=IF(AND(A{r}="",C{r}="",D{r}=""),"",IF(G{r}<>"","要確認",IF(E{r}<>"",IF(H{r}>0,"完了(遅れ)","完了"),'
                       f'IF(H{r}>0,"遅れ",IF($K$1<C{r},"着工前","進行中")))))'))
    sc["K1"] = today
    sc["J1"] = "基準日"
    sc["K2"] = f'=COUNTIF(I2:I{N},"要確認")'
    return wb

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

def norm(v):
    if v is None: return ""
    return "".join(chr(ord(c) - 0xFEE0) if "！" <= c <= "～" else ("-" if c == "－" else c) for c in str(v).strip().replace("　", " ").strip())

def fw(s):
    return "".join(chr(ord(c) + 0xFEE0) if "!" <= c <= "~" and c != "-" else ("－" if c == "-" else c) for c in s)

def isd(v): return isinstance(v, D)

def expected(projects, tasks, today):
    keys = {norm(p) for p in projects}
    rows = []; late = {k: [0, 0] for k in keys}; bad = 0
    for no, kind, s, e, a in tasks:
        if no in (None, "") and s is None and e is None: rows.append(("", "", "")); continue
        chk = ""
        if no in (None, ""): chk = "工事番号が空欄"
        elif norm(no) not in keys: chk = "工事番号なし"
        elif not isd(s) or not isd(e): chk = "予定日が未入力か日付でない"
        elif e < s: chk = "日付が逆転"
        elif a is not None and not isd(a): chk = "実績日が日付でない"
        if chk:
            rows.append((chk, "", "要確認")); bad += 1; continue
        days = max(0, ((a if a is not None else today) - e).days)
        if a is not None: st = "完了(遅れ)" if days > 0 else "完了"
        elif days > 0: st = "遅れ"
        elif today < s: st = "着工前"
        else: st = "進行中"
        rows.append(("", days, st))
        if days > 0: late[norm(no)][0] += 1
        late[norm(no)][1] += days
    return rows, late, bad

TODAY = D(2026, 10, 20)
SCEN = [
 ("進行中(予定終了の前)", ["K-001"], [("K-001","大工",D(2026,10,15),D(2026,10,25),None)]),
 ("予定終了日の当日は遅れではない", ["K-002"], [("K-002","大工",D(2026,10,10),D(2026,10,20),None)]),
 ("1日超過は遅れ", ["K-003"], [("K-003","大工",D(2026,10,10),D(2026,10,19),None)]),
 ("着工前", ["K-004"], [("K-004","左官",D(2026,10,25),D(2026,10,30),None)]),
 ("完了・予定内", ["K-005"], [("K-005","塗装",D(2026,10,10),D(2026,10,19),D(2026,10,18))]),
 ("完了・予定より遅れて完了", ["K-006"], [("K-006","塗装",D(2026,10,10),D(2026,10,20),D(2026,10,25))]),
 ("月をまたぐ遅れ(基準日11/2・予定終了10/30)", ["K-007"], [("K-007","電気",D(2026,10,20),D(2026,10,30),None)]),
 ("存在しない工事番号", ["K-008"], [("K-080","電気",D(2026,10,1),D(2026,10,10),None)]),
 ("工事番号が空欄", ["K-009"], [(None,"電気",D(2026,10,1),D(2026,10,10),None)]),
 ("日付の逆転(終了が開始より前)", ["K-010"], [("K-010","設備",D(2026,10,10),D(2026,10,1),None)]),
 ("予定終了が空欄", ["K-011"], [("K-011","設備",D(2026,10,1),None,None)]),
 ("日付の欄に文字('未定')", ["K-012"], [("K-012","設備","未定",D(2026,10,10),None)]),
 ("実績の欄に文字", ["K-013"], [("K-013","設備",D(2026,10,1),D(2026,10,10),"済")]),
 ("全角・前後スペースの番号", ["K-014"], [("Ｋ－014","解体",D(2026,10,1),D(2026,10,10),None),(" K-014 ","解体",D(2026,10,1),D(2026,10,12),None)]),
 ("まったくの空行", ["K-015"], [(None,None,None,None,None)]),
]

def main():
    ng = 0
    for name, projects, tasks in SCEN:
        today = D(2026, 11, 2) if name.startswith("月をまたぐ") else TODAY
        res = run_lo(build(projects, tasks, today)); led, sc = res["台帳"], res["工程"]
        rows, late, bad = expected(projects, tasks, today)
        got = [((sc.cell(r,7).value or ""), (sc.cell(r,8).value if sc.cell(r,8).value is not None else ""), (sc.cell(r,9).value or "")) for r in range(2, len(tasks)+2)]
        gl = {norm(led.cell(r,1).value): [led.cell(r,2).value, led.cell(r,3).value] for r in range(2, len(projects)+2)}
        ok = got == rows and gl == late and sc["K2"].value == bad
        ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: チェック/遅れ日数/状態{got} 台帳{gl} 要確認{sc['K2'].value}")
        if not ok: print("   期待", rows, late, bad)
    rnd = random.Random(20261020)
    projects = [f"K-{i:03d}" for i in range(1, 301)]
    kinds = ["解体", "大工", "電気", "設備", "左官", "塗装"]
    tasks = []
    for _ in range(1000):
        p = rnd.choice(projects); s = D(2026,9,15) + datetime.timedelta(days=rnd.randint(0, 60)); e = s + datetime.timedelta(days=rnd.randint(1, 20))
        t = rnd.random(); a = None
        if rnd.random() < 0.5: a = e + datetime.timedelta(days=rnd.randint(-3, 8))
        if t < 0.76: tasks.append((p, rnd.choice(kinds), s, e, a))
        elif t < 0.80: tasks.append(("K-9%02d" % rnd.randint(0, 99), "大工", s, e, a))
        elif t < 0.84: tasks.append((None, "大工", s, e, a))
        elif t < 0.88: tasks.append((fw(p), "大工", s, e, a))
        elif t < 0.91: tasks.append((" " + p + " ", "大工", s, e, a))
        elif t < 0.94: tasks.append((p, "大工", e, s if s != e else s - datetime.timedelta(days=1), a))
        elif t < 0.97: tasks.append((p, "大工", s, None, a))
        else: tasks.append((p, "大工", "未定", e, "済" if a else None))
    res = run_lo(build(projects, tasks, TODAY)); led, sc = res["台帳"], res["工程"]
    rows, late, bad = expected(projects, tasks, TODAY)
    got = [((sc.cell(r,7).value or ""), (sc.cell(r,8).value if sc.cell(r,8).value is not None else ""), (sc.cell(r,9).value or "")) for r in range(2, 1002)]
    gl = {norm(led.cell(r,1).value): [led.cell(r,2).value, led.cell(r,3).value] for r in range(2, 302)}
    ok = got == rows and gl == late and sc["K2"].value == bad
    ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: 工事300件・工程1000行 一致={ok} 要確認{bad}件")
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
