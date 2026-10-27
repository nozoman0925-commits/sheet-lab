"""記事『未入金の請求を経過日数別に分けて集計するExcelの式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/receivable_aging.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random
from datetime import date, timedelta
from openpyxl import Workbook, load_workbook

N = 1000
BASE = date(2026, 10, 31)
G = ('=IF(A{r}&B{r}&C{r}&D{r}&E{r}="","",IF(NOT(ISNUMBER($N$2)),"基準日が日付でない",IF(OR($N$2<1,$N$2<>INT($N$2)),"基準日が日付でない",'
     'IF(A{r}="","請求番号が空欄",IF(COUNTIF($A$2:$A${n},A{r})>1,"請求番号が重複",'
     'IF(C{r}="","請求額が空欄",IF(NOT(ISNUMBER(C{r})),"請求額が数字でない",IF(OR(C{r}<1,C{r}<>INT(C{r})),"請求額は1以上の整数",'
     'IF(D{r}="","入金期日が空欄",IF(NOT(ISNUMBER(D{r})),"入金期日が日付でない",IF(OR(D{r}<1,D{r}<>INT(D{r})),"入金期日が日付でない",'
     'IF(E{r}="","",IF(NOT(ISNUMBER(E{r})),"入金額が数字でない",IF(OR(E{r}<0,E{r}<>INT(E{r})),"入金額は0以上の整数",'
     'IF(E{r}>C{r},"入金額が請求額を超えている","")))))))))))))))')
F = ('=IF(OR(G{r}<>"",A{r}&B{r}&C{r}&D{r}&E{r}=""),"",IF(C{r}-N(E{r})=0,"入金済み",IF(D{r}>=$N$2,"期日内",'
     'IF($N$2-D{r}<=30,"超過30日以内",IF($N$2-D{r}<=60,"超過31〜60日","超過61日以上")))))')
H = '=IF(OR(F{r}="",F{r}="入金済み"),"",C{r}-N(E{r}))'
I = '=IF(OR(F{r}="",F{r}="入金済み",F{r}="期日内"),"",$N$2-D{r})'
SUM = {"N5": '=SUMIF($F$2:$F${n},"期日内",$H$2:$H${n})', "N6": '=SUMIF($F$2:$F${n},"超過30日以内",$H$2:$H${n})',
       "N7": '=SUMIF($F$2:$F${n},"超過31〜60日",$H$2:$H${n})', "N8": '=SUMIF($F$2:$F${n},"超過61日以上",$H$2:$H${n})',
       "N9": '=SUM(N6:N8)'}

def serial(v):
    return (v - date(1899, 12, 30)).days if isinstance(v, date) else v

def build(rows, base):
    wb = Workbook(); ws = wb.active; ws.title = "未入金"
    for i, h in enumerate(["請求番号", "取引先", "請求額", "入金期日", "入金額", "区分", "チェック", "未入金額", "超過日数"], 1): ws.cell(1, i, h)
    ws["M2"] = "基準日"; ws["N2"] = base
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1): ws.cell(r, c, v)
        ws.cell(r, 7, G.format(r=r, n=N)); ws.cell(r, 6, F.format(r=r))
        ws.cell(r, 8, H.format(r=r)); ws.cell(r, 9, I.format(r=r))
    for k, f in SUM.items(): ws[k] = f.format(n=N)
    return wb

def run_lo(wb):
    t = tempfile.mkdtemp(); src = os.path.join(t, "t.xlsx"); wb.save(src)
    out = os.path.join(t, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
empty = lambda v: v in (None, "")
key = lambda v: v.lower() if isinstance(v, str) else v

def expect(rows, base):
    b = serial(base); out = []
    for a, nm, c, d, e in rows:
        d = serial(d)
        if all(empty(x) for x in (a, nm, c, d, e)): out.append(("", "", "", "")); continue
        err = None
        if not num(b) or b < 1 or b != int(b): err = "基準日が日付でない"
        elif empty(a): err = "請求番号が空欄"
        elif sum(1 for x in rows if not empty(x[0]) and key(x[0]) == key(a)) > 1: err = "請求番号が重複"
        elif empty(c): err = "請求額が空欄"
        elif not num(c): err = "請求額が数字でない"
        elif c < 1 or c != int(c): err = "請求額は1以上の整数"
        elif empty(d): err = "入金期日が空欄"
        elif not num(d) or d < 1 or d != int(d): err = "入金期日が日付でない"
        elif not empty(e) and not num(e): err = "入金額が数字でない"
        elif not empty(e) and (e < 0 or e != int(e)): err = "入金額は0以上の整数"
        elif not empty(e) and e > c: err = "入金額が請求額を超えている"
        if err: out.append(("", "", "", err)); continue
        rest = c - (0 if empty(e) else e)
        if rest == 0: out.append(("入金済み", "", "", "")); continue
        if d >= b: out.append(("期日内", rest, "", "")); continue
        late = b - d
        st = "超過30日以内" if late <= 30 else "超過31〜60日" if late <= 60 else "超過61日以上"
        out.append((st, rest, late, ""))
    return out

def exp_sum(ex):
    s = lambda k: sum(x[1] for x in ex if x[0] == k)
    return (s("期日内"), s("超過30日以内"), s("超過31〜60日"), s("超過61日以上"), s("超過30日以内") + s("超過31〜60日") + s("超過61日以上"))

def read(wb, n):
    ws = wb["未入金"]; res = []
    for r in range(2, n + 2):
        res.append(tuple("" if ws.cell(r, c).value is None else ws.cell(r, c).value for c in (6, 8, 9, 7)))
    return res, tuple("" if ws[k].value is None else ws[k].value for k in ("N5", "N6", "N7", "N8", "N9"))

def same(x, y):
    if len(x) != len(y): return False
    for p, q in zip(x, y):
        if num(p) and num(q):
            if abs(p - q) > 1e-6: return False
        elif p != q: return False
    return True

def check(rows, base):
    e = expect(rows, base); es = exp_sum(e)
    g, gs = read(run_lo(build(rows, base)), len(rows))
    ok = all(same(a, b) for a, b in zip(g, e)) and same(gs, es)
    return ok, g, e, gs, es

D = lambda n: BASE - timedelta(days=n)   # 基準日のn日前
def R(due, paid=None, amt=100000, no="S-001"): return [(no, "A社", amt, due, paid)]
S = lambda name, rows, base=BASE: (name, rows, base)
SCEN = [
 S("期日の前日(期日内)", R(BASE + timedelta(days=1))),
 S("期日が基準日と同じ(期日内)", R(BASE)),
 S("期日から1日超過(30日以内)", R(D(1))),
 S("期日から30日超過ちょうど(30日以内)", R(D(30))),
 S("期日から31日超過(31〜60日)", R(D(31))),
 S("期日から60日超過ちょうど(31〜60日)", R(D(60))),
 S("期日から61日超過(61日以上)", R(D(61))),
 S("期日から400日超過(61日以上)", R(D(400))),
 S("全額入金済み(超過していても入金済み)", R(D(90), 100000)),
 S("一部入金(61日超過・未入金額=残り)", R(D(90), 40000)),
 S("入金額が空欄(0として扱う)", R(D(10), None)),
 S("入金額が0", R(D(10), 0)),
 S("年またぎ(期日2025/12/20・基準日2026/01/10)", R(date(2025, 12, 20)), date(2026, 1, 10)),
 S("請求番号が空欄", [(None, "A社", 100000, D(10), None)]),
 S("請求番号が重複(大文字小文字違い)", [("S-001", "A社", 100000, D(10), None), ("s-001", "B社", 50000, D(5), None)]),
 S("請求額が空欄", [("S-001", "A社", None, D(10), None)]),
 S("請求額が文字『10万』", [("S-001", "A社", "10万", D(10), None)]),
 S("請求額が全角『１００，０００』", [("S-001", "A社", "１００，０００", D(10), None)]),
 S("請求額が0", [("S-001", "A社", 0, D(10), None)]),
 S("請求額が小数", [("S-001", "A社", 1000.5, D(10), None)]),
 S("入金期日が空欄", [("S-001", "A社", 100000, None, None)]),
 S("入金期日が文字『10/05』", [("S-001", "A社", 100000, "10/05", None)]),
 S("入金期日が全角文字『２０２６/１０/０５』", [("S-001", "A社", 100000, "２０２６/１０/０５", None)]),
 S("入金額が文字『4万』", R(D(10), "4万")),
 S("入金額がマイナス", R(D(10), -1)),
 S("入金額が請求額超過(二重入金)", R(D(10), 100001)),
 S("基準日が空欄", R(D(10)), None),
 S("基準日が文字", R(D(10)), "今日"),
 S("集計(期日内2・30日以内・31〜60日・61日以上・入金済み)", [
    ("S-001", "A社", 100000, BASE + timedelta(days=5), None), ("S-002", "B社", 200000, BASE, 50000),
    ("S-003", "C社", 300000, D(10), None), ("S-004", "D社", 400000, D(45), 100000),
    ("S-005", "E社", 500000, D(100), None), ("S-006", "F社", 600000, D(200), 600000)]),
 S("空行をはさむ", [(None, None, None, None, None), ("S-001", "A社", 100000, D(10), None)]),
]

def main():
    ng = 0
    for name, rows, base in SCEN:
        ok, g, e, gs, es = check(rows, base); ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: {g} 集計={gs}")
        if not ok: print("   期待", e, es)
    rnd = random.Random(20261028); rows = []
    for i in range(500):
        a = f"S-{i:03d}"; c = rnd.randint(1, 500) * 1000
        d = BASE + timedelta(days=rnd.randint(-150, 40))
        e = rnd.choice([None, 0, c, rnd.randint(0, c)]); x = rnd.random()
        if x < 0.02: a = None
        elif x < 0.04: a = rows[i - 1][0] if i else a
        elif x < 0.06: c = rnd.choice([None, "10万", 0, 1000.5, "１０００"])
        elif x < 0.08: d = rnd.choice([None, "10/05", "２０２６/１０/０５"])
        elif x < 0.12: e = rnd.choice(["4万", -5, 1000.5, (c or 0) + 1 if num(c) else 1])
        rows.append((a, f"取引先{i}", c, d, e))
    ok, g, e, gs, es = check(rows, BASE); ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: {len(rows)}行 全項目一致={ok} 集計={gs}")
    if not ok:
        for i, (x, y) in enumerate(zip(g, e)):
            if not same(x, y): print("  行", i + 2, rows[i], x, y); break
        print("  集計", gs, es)
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
