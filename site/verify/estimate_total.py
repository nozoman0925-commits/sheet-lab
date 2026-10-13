"""記事『見積書の合計・税込計算をExcelで自動にする式』をLibreOfficeで再計算して期待値(整数計算)と照合する。
使い方: python3 site/verify/estimate_total.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os
from decimal import Decimal
from openpyxl import Workbook, load_workbook

# --- 1) 見積書の10ケース ---
# 明細: (数量, 単価)。単価は小数も可。諸経費率(%)、値引き(円)
CASES = [
 ("単純3行",           [(1, 100000), (2, 35000), (10, 1200)], 0, 0),
 ("諸経費10%",         [(1, 100000), (2, 35000), (10, 1200)], 10, 0),
 ("諸経費10%+値引き",  [(1, 100000), (2, 35000), (10, 1200)], 10, 5000),
 ("小数の数量",        [(2.5, 3333), (0.5, 12345)], 0, 0),
 ("諸経費の端数",      [(1, 12345), (3, 6789)], 7, 0),
 ("税の端数",          [(1, 1005)], 0, 0),
 ("値引きで税端数",    [(1, 99999)], 5, 1234),
 ("明細が空",          [], 10, 0),
 ("値引き=小計+諸経費", [(1, 10000)], 10, 11000),
 ("大きな金額",        [(12, 987654), (3, 456789)], 15, 100000),
]

def half_up(num, den):  # 非負の分数 num/den を四捨五入(整数)
    return (2 * num + den) // (2 * den)

def expected(lines, pct, disc):
    sub = 0
    for q, p in lines:
        v = Decimal(str(q)) * Decimal(str(p))
        sub += int(v.to_integral_value(rounding="ROUND_HALF_UP"))
    exp = half_up(sub * pct, 100)
    taxable = sub + exp - disc
    tax = taxable * 10 // 100 if taxable >= 0 else None
    return sub, exp, taxable, tax, (taxable + tax if tax is not None else None)

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

def main():
    ng = 0
    wb = Workbook(); wb.remove(wb.active)
    for name, lines, pct, disc in CASES:
        ws = wb.create_sheet(name[:28])
        ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "品名", "数量", "単価", "金額"
        for i in range(2, 22):
            ws[f"D{i}"] = f'=IF(OR(B{i}="",C{i}=""),0,ROUND(B{i}*C{i},0))'
        for i, (q, p) in enumerate(lines, 2):
            ws[f"B{i}"], ws[f"C{i}"] = q, p
        ws["F2"], ws["G2"] = "小計", "=SUM(D2:D21)"
        ws["F3"], ws["G3"] = "諸経費率", pct / 100
        ws["F4"], ws["G4"] = "諸経費", "=ROUND(G2*G3,0)"
        ws["F5"], ws["G5"] = "値引き", disc
        ws["F6"], ws["G6"] = "課税対象額", "=G2+G4-G5"
        ws["F7"], ws["G7"] = "消費税", "=ROUNDDOWN(G6*0.1,0)"
        ws["F8"], ws["G8"] = "税込合計", "=G6+G7"
    res = run_lo(wb)
    for name, lines, pct, disc in CASES:
        ws = res[name[:28]]
        got = tuple(ws[c].value for c in ("G2", "G4", "G6", "G7", "G8"))
        exp = expected(lines, pct, disc)
        if exp[2] < 0:  # マイナスはこの検証では対象外(別途「注意点」で扱う)
            exp = (exp[0], exp[1], exp[2], None, None)
            got = got[:3] + (None, None)
        ok = got == exp; ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: 小計{got[0]} 諸経費{got[1]} 課税{got[2]} 税{got[3]} 税込{got[4]} (期待 {exp})")

    # --- 2) 端数処理の網羅 ---
    wb2 = Workbook(); ws = wb2.active
    # 消費税: ROUNDDOWN(N*0.1,0) と 整数 N//10 を N=0..30000 で照合
    N = 30000
    for n in range(N + 1):
        ws.cell(n + 1, 1, n); ws.cell(n + 1, 2, f"=ROUNDDOWN(A{n+1}*0.1,0)")
    # 諸経費: ROUND(S*率,0) と 整数四捨五入を S=0..2000 x 率1..20% で照合
    ws2 = wb2.create_sheet("k")
    r = 1
    for pct in range(1, 21):
        for s in range(0, 2001):
            ws2.cell(r, 1, s); ws2.cell(r, 2, pct / 100); ws2.cell(r, 3, f"=ROUND(A{r}*B{r},0)"); ws2.cell(r, 4, pct); r += 1
    res2 = run_lo(wb2)
    bad_tax = [n for n in range(N + 1) if res2.worksheets[0].cell(n + 1, 2).value != n // 10]
    bad_ex = 0
    w2 = res2["k"]
    for rr in range(1, r):
        s, pct, got = w2.cell(rr, 1).value, w2.cell(rr, 4).value, w2.cell(rr, 3).value
        if got != half_up(s * pct, 100): bad_ex += 1
    print(f"消費税(ROUNDDOWN(N*0.1,0)) N=0..{N}: 不一致 {len(bad_tax)} 件 {bad_tax[:5]}")
    print(f"諸経費(ROUND(S*率,0)) S=0..2000 x 率1..20%: {r-1} 件 不一致 {bad_ex} 件")
    ng += len(bad_tax) + bad_ex
    print("不一致", ng, "件"); sys.exit(1 if ng else 0)
main()
