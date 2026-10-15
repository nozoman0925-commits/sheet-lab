"""note原稿用: 年度(4月始まり)別の連番・取引先別の連番・次の番号をLibreOfficeで確認する。
使い方: python3 site/verify/number_sets.py"""
import subprocess, sys, tempfile, os, datetime
from openpyxl import Workbook, load_workbook
D = datetime.date
ROWS = [(D(2026,3,30),"A社"),(D(2026,3,31),"B社"),(D(2026,4,1),"A社"),(D(2026,4,1),"C社"),(D(2026,9,15),"A社"),
        (D(2027,1,5),"B社"),(D(2027,3,31),"A社"),(D(2027,4,1),"B社"),(D(2024,2,29),"C社"),(D(2024,4,1),"A社")]
def main():
    wb=Workbook(); ws=wb.active
    n=len(ROWS)+1
    for i,(d,c) in enumerate(ROWS,2):
        ws.cell(i,1,d).number_format="yyyy/mm/dd"; ws.cell(i,2,c)
        fy=f"(YEAR(A{i})-(MONTH(A{i})<4))"
        ws.cell(i,3,f'=IF(A{i}="","","Q"&{fy}&"-"&TEXT(COUNTIFS($A$2:A{i},">="&DATE({fy},4,1),$A$2:A{i},"<="&DATE({fy}+1,3,31)),"000"))')
        ws.cell(i,4,f'=IF(B{i}="","",B{i}&"-"&TEXT(COUNTIF($B$2:B{i},B{i}),"000"))')
        ws.cell(i,5,f'="Q"&TEXT(A{i},"yyyymm")&"-"&TEXT(COUNTIF($F$2:$F${n},"Q"&TEXT(A{i},"yyyymm")&"-*")+1,"000")')
    # 次の番号の確認用: 既存の番号(F列)
    existing=["Q202610-001","Q202610-002","Q202610-003","Q202609-001"]
    for i,v in enumerate(existing,2): ws.cell(i,6,v)
    ws.cell(20,1,D(2026,10,20)); ws.cell(20,5,'="Q"&TEXT(A20,"yyyymm")&"-"&TEXT(COUNTIF($F$2:$F$12,"Q"&TEXT(A20,"yyyymm")&"-*")+1,"000")')
    ws.cell(21,1,D(2026,11,2)); ws.cell(21,5,'="Q"&TEXT(A21,"yyyymm")&"-"&TEXT(COUNTIF($F$2:$F$12,"Q"&TEXT(A21,"yyyymm")&"-*")+1,"000")')
    d=tempfile.mkdtemp(); src=os.path.join(d,"t.xlsx"); wb.save(src); out=os.path.join(d,"o"); os.mkdir(out)
    subprocess.run(["soffice","--headless","--convert-to","xlsx","--outdir",out,src],check=True,capture_output=True)
    r=load_workbook(os.path.join(out,"t.xlsx"),data_only=True).active
    fyc={}; cc={}; ng=0; k=0
    for i,(d_,c) in enumerate(ROWS,2):
        fy=d_.year-(1 if d_.month<4 else 0)
        fyc[fy]=fyc.get(fy,0)+1; cc[c]=cc.get(c,0)+1
        exp=(f"Q{fy}-{fyc[fy]:03d}", f"{c}-{cc[c]:03d}")
        got=(r.cell(i,3).value, r.cell(i,4).value); ok=got==exp; ng+=not ok; k+=2
        print(("OK " if ok else "NG ")+f"{d_} {c}: 年度別={got[0]} 取引先別={got[1]}")
        if not ok: print("  期待",exp)
    nx=(r.cell(20,5).value, r.cell(21,5).value)
    ok=nx==("Q202610-004","Q202611-001"); ng+=not ok; k+=2
    print("次の番号:", nx, "OK" if ok else "NG")
    print("検証",k,"件 不一致",ng); sys.exit(1 if ng else 0)
main()
