"""記事『完工月ごとの完工高・粗利・粗利率をExcelで集計する式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/monthly_completion.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random
from datetime import date, timedelta
from openpyxl import Workbook, load_workbook

N = 1000
START = date(2026, 4, 1)
G = ('=IF(A{r}&B{r}&C{r}&D{r}&E{r}="","",IF(A{r}="","工事番号が空欄",IF(COUNTIF($A$2:$A${n},A{r})>1,"工事番号が重複",'
     'IF(C{r}="","契約金額が空欄",IF(NOT(ISNUMBER(C{r})),"契約金額が数字でない",IF(OR(C{r}<1,C{r}<>INT(C{r})),"契約金額は1以上の整数",'
     'IF(D{r}="",IF(E{r}="","","原価が空欄"),IF(NOT(ISNUMBER(D{r})),"原価が数字でない",IF(OR(D{r}<0,D{r}<>INT(D{r})),"原価は0以上の整数",'
     'IF(E{r}="","",IF(NOT(ISNUMBER(E{r})),"完了日が日付でない",IF(OR(E{r}<1,E{r}<>INT(E{r})),"完了日が日付でない",""))))))))))))')
F = '=IF(OR(G{r}<>"",A{r}&B{r}&C{r}&D{r}&E{r}=""),"",IF(E{r}="","進行中","完工"))'
H = '=IF(F{r}="完工",DATE(YEAR(E{r}),MONTH(E{r}),1),"")'
I = '=IF(F{r}="完工",C{r}-D{r},"")'
J = '=IF(F{r}="完工",(C{r}-D{r})/C{r},"")'

def serial(v):
    return (v - date(1899, 12, 30)).days if isinstance(v, date) else v

def build(rows, start):
    wb = Workbook(); ws = wb.active; ws.title = "完工"
    for i, h in enumerate(["工事番号", "工事名", "契約金額", "原価合計", "完了日", "区分", "チェック", "完工月", "粗利", "粗利率"], 1): ws.cell(1, i, h)
    ws["S2"] = start
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1): ws.cell(r, c, v)
        ws.cell(r, 7, G.format(r=r, n=N)); ws.cell(r, 6, F.format(r=r))
        ws.cell(r, 8, H.format(r=r)); ws.cell(r, 9, I.format(r=r)); ws.cell(r, 10, J.format(r=r))
    ws["L2"] = '=IF(ISNUMBER($S$2),DATE(YEAR($S$2),MONTH($S$2),1),"")'
    for r in range(3, 14): ws.cell(r, 12, f'=IF(L{r-1}="","",EDATE(L{r-1},1))')
    for r in range(2, 14):
        ws.cell(r, 13, f'=IF(L{r}="","",COUNTIF($H$2:$H${N},L{r}))')
        ws.cell(r, 14, f'=IF(L{r}="","",SUMIFS($C$2:$C${N},$H$2:$H${N},L{r}))')
        ws.cell(r, 15, f'=IF(L{r}="","",SUMIFS($I$2:$I${N},$H$2:$H${N},L{r}))')
        ws.cell(r, 16, f'=IF(L{r}="","",IF(N{r}=0,"",O{r}/N{r}))')
    ws["M14"] = '=IF(L2="","",SUM(M2:M13))'; ws["N14"] = '=IF(L2="","",SUM(N2:N13))'; ws["O14"] = '=IF(L2="","",SUM(O2:O13))'
    ws["P14"] = '=IF(L2="","",IF(N14=0,"",O14/N14))'
    ws["M15"] = f'=IF(L2="","",COUNT($H$2:$H${N})-M14)'
    ws["M16"] = f'=COUNTIF($F$2:$F${N},"進行中")'
    ws["M17"] = f'=SUMPRODUCT(--(LEN($G$2:$G${N})>0))'
    return wb

def run_lo(wb):
    t = tempfile.mkdtemp(); src = os.path.join(t, "t.xlsx"); wb.save(src)
    out = os.path.join(t, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
empty = lambda v: v in (None, "")
key = lambda v: v.lower() if isinstance(v, str) else v

def first_of_month(sv):
    d = date(1899, 12, 30) + timedelta(days=int(sv))
    return (date(d.year, d.month, 1) - date(1899, 12, 30)).days

def add_month(sv, k):
    d = date(1899, 12, 30) + timedelta(days=int(sv))
    m = d.month - 1 + k
    return (date(d.year + m // 12, m % 12 + 1, 1) - date(1899, 12, 30)).days

def expect(rows):
    out = []
    for a, nm, c, d, e in rows:
        e = serial(e)
        if all(empty(x) for x in (a, nm, c, d, e)): out.append(("", "", "", "", "")); continue
        err = None
        if empty(a): err = "工事番号が空欄"
        elif sum(1 for x in rows if not empty(x[0]) and key(x[0]) == key(a)) > 1: err = "工事番号が重複"
        elif empty(c): err = "契約金額が空欄"
        elif not num(c): err = "契約金額が数字でない"
        elif c < 1 or c != int(c): err = "契約金額は1以上の整数"
        elif empty(d): err = "原価が空欄" if not empty(e) else None
        elif not num(d): err = "原価が数字でない"
        elif d < 0 or d != int(d): err = "原価は0以上の整数"
        if err is None and not empty(d) and not empty(e) and (not num(e) or e < 1 or e != int(e)): err = "完了日が日付でない"
        if err: out.append(("", "", "", "", err)); continue
        if empty(e): out.append(("進行中", "", "", "", "")); continue
        out.append(("完工", first_of_month(e), c - d, (c - d) / c, ""))
    return out

def expect_summary(rows, ex, start):
    start = serial(start)
    if not num(start): return None
    s0 = first_of_month(start); months = []; m = s0
    for k in range(12): months.append(add_month(s0, k))
    tab = []
    for mo in months:
        idx = [i for i, x in enumerate(ex) if x[0] == "完工" and x[1] == mo]
        cnt = len(idx); sales = sum(rows[i][2] for i in idx); prof = sum(ex[i][2] for i in idx)
        tab.append((mo, cnt, sales, prof, prof / sales if sales else ""))
    tc = sum(t[1] for t in tab); ts = sum(t[2] for t in tab); tp = sum(t[3] for t in tab)
    total = (tc, ts, tp, tp / ts if ts else "")
    allc = sum(1 for x in ex if x[0] == "完工")
    return tab, total, allc - tc

def same(p, q):
    if num(p) and num(q): return abs(p - q) < 1e-9
    return p == q

def norm(v):
    if v is None: return ""
    if hasattr(v, "year"): return (v.date() if hasattr(v, "date") else v) and (( v.date() if hasattr(v, "date") else v) - date(1899, 12, 30)).days
    return v

def check(rows, start):
    ex = expect(rows); es = expect_summary(rows, ex, start)
    ws = run_lo(build(rows, start))["完工"]
    ok = True; got = []
    for i, e in enumerate(ex):
        r = i + 2
        g = tuple(norm(ws.cell(r, c).value) for c in (6, 8, 9, 10, 7)); got.append(g)
        if not all(same(a, b) for a, b in zip(g, e)): ok = False
    gsum = None
    if es is None:
        gs = [norm(ws.cell(r, c).value) for r in range(2, 14) for c in range(12, 17)] + [norm(ws[k].value) for k in ("M14", "N14", "O14", "P14", "M15")]
        if any(v != "" for v in gs): ok = False
    else:
        tab, total, outside = es
        for k, t in enumerate(tab):
            g = tuple(norm(ws.cell(2 + k, c).value) for c in range(12, 17))
            if not all(same(a, b) for a, b in zip(g, t)): ok = False; gsum = (k, g, t)
        gt = tuple(norm(ws[x].value) for x in ("M14", "N14", "O14", "P14"))
        if not all(same(a, b) for a, b in zip(gt, total)): ok = False; gsum = (gt, total)
        if not same(norm(ws["M15"].value), outside): ok = False; gsum = ("期間外", ws["M15"].value, outside)
    if norm(ws["M16"].value) != sum(1 for x in ex if x[0] == "進行中"): ok = False; gsum = "進行中件数"
    if norm(ws["M17"].value) != sum(1 for x in ex if x[4] != ""): ok = False; gsum = "チェック件数"
    summ = None if es is None else (es[1][:3], es[2])
    return ok, got, ex, summ, gsum

D = lambda y, m, d: date(y, m, d)
def JB(done, cost=60000, amt=100000, no="K-001"): return [(no, "A邸", amt, cost, done)]
S = lambda name, rows, start=START: (name, rows, start)
SCEN = [
 S("月末日(10/31)に完工→10月", JB(D(2026, 10, 31))),
 S("月初(11/01)に完工→11月", JB(D(2026, 11, 1))),
 S("4月1日(期間の最初の日)に完工", JB(D(2026, 4, 1))),
 S("3月31日(期間の前日)に完工→期間外", JB(D(2026, 3, 31))),
 S("翌年3月31日(期間の最後の日)に完工", JB(D(2027, 3, 31))),
 S("翌年4月1日(期間の翌日)に完工→期間外", JB(D(2027, 4, 1))),
 S("うるう日(2028/02/29)に完工", JB(D(2028, 2, 29)), D(2028, 1, 1)),
 S("完了日が空欄=進行中(集計に入らない)", JB(None)),
 S("進行中で原価も空欄(まだ入っていない)", JB(None, None)),
 S("原価が0の完工(粗利=契約金額)", JB(D(2026, 10, 5), 0)),
 S("赤字の完工(原価>契約金額)", JB(D(2026, 10, 5), 130000)),
 S("契約金額が空欄", [("K-001", "A邸", None, 60000, D(2026, 10, 5))]),
 S("契約金額が文字『10万』", [("K-001", "A邸", "10万", 60000, D(2026, 10, 5))]),
 S("契約金額が全角『１００，０００』", [("K-001", "A邸", "１００，０００", 60000, D(2026, 10, 5))]),
 S("契約金額が0", [("K-001", "A邸", 0, 60000, D(2026, 10, 5))]),
 S("契約金額が小数", [("K-001", "A邸", 1000.5, 600, D(2026, 10, 5))]),
 S("完工済みで原価が空欄(粗利が膨らむ誤り)", JB(D(2026, 10, 5), None)),
 S("原価が文字『6万』", JB(D(2026, 10, 5), "6万")),
 S("原価がマイナス", JB(D(2026, 10, 5), -1)),
 S("原価が小数", JB(D(2026, 10, 5), 600.5)),
 S("完了日が文字『10/05』", JB("10/05")),
 S("完了日が全角『２０２６/１０/０５』", JB("２０２６/１０/０５")),
 S("完了日が0", JB(0)),
 S("工事番号が空欄", [(None, "A邸", 100000, 60000, D(2026, 10, 5))]),
 S("工事番号が重複(大文字小文字違い)", [("K-001", "A邸", 100000, 60000, D(2026, 10, 5)), ("k-001", "B邸", 50000, 30000, D(2026, 10, 6))]),
 S("開始月が月の途中(4/15→4月に切り下げ)", JB(D(2026, 4, 2)), D(2026, 4, 15)),
 S("開始月が空欄(集計は空)", JB(D(2026, 10, 5)), None),
 S("開始月が文字(集計は空)", JB(D(2026, 10, 5)), "今期"),
 S("空行をはさむ", [(None, None, None, None, None), ("K-001", "A邸", 100000, 60000, D(2026, 10, 5))]),
 S("集計(10月に2件・11月に1件・進行中1件・期間外1件・エラー1件)", [
    ("K-001", "A邸", 1000000, 700000, D(2026, 10, 5)), ("K-002", "B邸", 500000, 350000, D(2026, 10, 31)),
    ("K-003", "C邸", 800000, 560000, D(2026, 11, 1)), ("K-004", "D邸", 900000, 100000, None),
    ("K-005", "E邸", 300000, 200000, D(2027, 6, 1)), ("K-006", "F邸", 400000, None, D(2026, 10, 9))]),
]

def main():
    ng = 0
    for name, rows, start in SCEN:
        ok, g, e, s, gs = check(rows, start); ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: {g[:3]} 集計={s}")
        if not ok: print("   期待", e[:3], gs)
    rnd = random.Random(20261029); rows = []
    for i in range(500):
        a = f"K-{i:03d}"; c = rnd.randint(1, 500) * 10000
        cost = int(c * rnd.uniform(0.4, 1.1)); d = START + timedelta(days=rnd.randint(-30, 400))
        if rnd.random() < 0.2: d = None
        x = rnd.random()
        if x < 0.02: a = None
        elif x < 0.04: a = rows[i - 1][0] if i else a
        elif x < 0.06: c = rnd.choice([None, "10万", 0, 1000.5, "１０００"])
        elif x < 0.09: cost = rnd.choice([None, "6万", -5, 600.5])
        elif x < 0.12: d = rnd.choice(["10/05", "２０２６/１０/０５", 0])
        rows.append((a, f"邸{i}", c, cost, d))
    ok, g, e, s, gs = check(rows, START); ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: {len(rows)}行 全項目一致={ok} 集計={s}")
    if not ok:
        for i, (x, y) in enumerate(zip(g, e)):
            if not all(same(p, q) for p, q in zip(x, y)): print("  行", i + 2, rows[i], x, y); break
        print("  ", gs)
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
