"""note有料原稿『月間表の日付・曜日・土日の式集』の式をLibreOfficeで再計算して照合する。
使い方: python3 site/verify/month_table.py"""
import subprocess, sys, tempfile, os, datetime, calendar
from openpyxl import Workbook, load_workbook
YM=[(2024,2),(2023,2),(2026,10),(2026,12),(2025,1),(2100,2),(2000,2),(2026,4),(2027,3),(2028,2),(2026,8),(2026,11)]
WD="日月火水木金土"
def main():
    wb=Workbook(); ws=wb.active
    for i,(y,m) in enumerate(YM,2):
        ws.cell(i,1,y); ws.cell(i,2,m)
        ws.cell(i,3,f"=DATE(A{i},B{i},1)")
        # 31日分の日付(空欄処理つき)を使い、土日数と平日数を数える
        ws.cell(i,4,f"=NETWORKDAYS(C{i},EOMONTH(C{i},0))")                      # 平日数
        ws.cell(i,5,f"=DAY(EOMONTH(C{i},0))-D{i}")                               # 土日数
        ws.cell(i,6,f'=SUMPRODUCT(--(WEEKDAY(C{i}+ROW($1:$31)-1,2)>=6),--(ROW($1:$31)<=DAY(EOMONTH(C{i},0))))')  # 土日数(別式)
        ws.cell(i,7,f'=IF(WEEKDAY(C{i},2)>=6,"休","平")')                        # 1日が土日か
        ws.cell(i,8,f'=ROUNDUP((DAY(EOMONTH(C{i},0))+WEEKDAY(C{i},2)-1)/7,0)')  # 月曜始まりの週の行数
        ws.cell(i,9,f'=TEXT(C{i},"yyyy/mm/dd")')
    d=tempfile.mkdtemp(); src=os.path.join(d,"t.xlsx"); wb.save(src)
    out=os.path.join(d,"o"); os.mkdir(out)
    subprocess.run(["soffice","--headless","--convert-to","xlsx","--outdir",out,src],check=True,capture_output=True)
    r=load_workbook(os.path.join(out,"t.xlsx"),data_only=True).active
    ng=0;n=0
    for i,(y,m) in enumerate(YM,2):
        nd=calendar.monthrange(y,m)[1]
        days=[datetime.date(y,m,k) for k in range(1,nd+1)]
        we=sum(1 for x in days if x.weekday()>=5)
        exp=(nd-we,we,we,"休" if days[0].weekday()>=5 else "平",
             len(calendar.Calendar(0).monthdayscalendar(y,m)),f"{y:04d}/{m:02d}/01")
        got=tuple(r.cell(i,c).value for c in (4,5,6,7,8,9))
        ok=got==exp; ng+=not ok; n+=6
        print(("OK " if ok else "NG ")+f"{y}/{m}: 平日{got[0]} 土日{got[1]}/{got[2]} 1日={got[3]} 週行数{got[4]}"+("" if ok else f" 期待{exp}"))
    print("検証",n,"件 不一致",ng,"月"); sys.exit(1 if ng else 0)
main()
