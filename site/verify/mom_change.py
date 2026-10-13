"""記事『エクセル家計簿で前月比の増減を費目別に出す方法』の数式をLibreOfficeで再計算して期待値と照合する。
使い方: python3 site/verify/mom_change.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os
from openpyxl import Workbook, load_workbook

# (費目, 先月B, 今月C, 増減額D, 増減率E(None=空文字), 判定F)
CASES = [
    ("食費",     30000, 36000,  6000,  0.2,      "増"),
    ("光熱費",   36000, 30000, -6000, -1/6,      "減"),
    ("通信費",   8000,  8000,      0,  0.0,      "同"),
    ("教養娯楽", 0,     5000,   5000,  None,     "増"),   # 先月0円: 率は出さない
    ("医療",     0,     0,         0,  None,     "同"),
    ("衣服美容", 5000,  0,     -5000, -1.0,      "減"),
    ("交通費",   None,  3000,   3000,  None,     "増"),   # 先月が空欄は0扱い
    ("住居",     80000, 80000,     0,  0.0,      "同"),
    ("日用品",   1234,  2345,   1111,  1111/1234, "増"),
    ("その他",   None,  None,      0,  None,     "同"),
]
D = '=C{r}-B{r}'
E = '=IF(B{r}=0,"",D{r}/B{r})'
F = '=IF(D{r}>0,"増",IF(D{r}<0,"減","同"))'
TOTAL_E = '=IF(B{r}=0,"",D{r}/B{r})'

def main():
    wb = Workbook(); ws = wb.active
    n = len(CASES)
    for r, (name, b, c, *_ ) in enumerate(CASES, 2):
        ws.cell(r, 1, name)
        if b is not None: ws.cell(r, 2, b)
        if c is not None: ws.cell(r, 3, c)
        ws.cell(r, 4, D.format(r=r)); ws.cell(r, 5, E.format(r=r)); ws.cell(r, 6, F.format(r=r))
    t = n + 2
    ws.cell(t, 1, "合計")
    ws.cell(t, 2, f"=SUM(B2:B{t-1})"); ws.cell(t, 3, f"=SUM(C2:C{t-1})")
    ws.cell(t, 4, D.format(r=t)); ws.cell(t, 5, E.format(r=t)); ws.cell(t, 6, F.format(r=t))
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    res = load_workbook(os.path.join(out, "t.xlsx"), data_only=True).active
    ng = 0
    rows = list(CASES) + [("合計", sum(x[1] or 0 for x in CASES), sum(x[2] or 0 for x in CASES), None, None, None)]
    tb, tc = rows[-1][1], rows[-1][2]
    rows[-1] = ("合計", tb, tc, tc - tb, (tc - tb) / tb, "増" if tc > tb else "減" if tc < tb else "同")
    for r, (name, b, c, ed, ee, ef) in enumerate(rows, 2):
        gd, ge, gf = res.cell(r, 4).value, res.cell(r, 5).value, res.cell(r, 6).value
        ok_d = gd == ed
        ok_e = (ge in (None, "")) if ee is None else (ge is not None and ge != "" and abs(ge - ee) < 1e-9)
        ok_f = gf == ef
        ok = ok_d and ok_e and ok_f; ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: 先月{b} 今月{c} → 増減額{gd!r}(期待{ed!r}) 増減率{ge!r}(期待{ee!r}) 判定{gf!r}(期待{ef!r})")
    print("不一致", ng, "件"); sys.exit(1 if ng else 0)
main()
