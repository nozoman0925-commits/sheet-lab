"""note原稿『勤務時間のExcel式集』の数式をLibreOfficeで再計算して期待値と照合する。
使い方: python3 site/verify/work_hours.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os
from openpyxl import Workbook, load_workbook

# (開始 "h:mm", 終了 "h:mm", 休憩分)
CASES = [("9:00","18:30",60),("9:00","17:00",0),("22:00","6:00",60),("23:30","0:15",0),("9:00","9:00",0),
         ("8:45","12:00",0),("13:00","17:20",15),("0:00","23:59",0),("6:00","14:45",45),("10:10","10:09",0)]
SUM_CASES = [["7:45","8:15","0:30"], ["12:30","13:45"], ["0:20","0:20","0:20"], ["8:00"]*5]

def mins(t): h,m=t.split(":"); return int(h)*60+int(m)
def hm(m): return f"{m//60}:{m%60:02d}"

def run(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice","--headless","--convert-to","xlsx","--outdir",out,src], check=True, capture_output=True)
    return load_workbook(os.path.join(out,"t.xlsx"), data_only=True).active

def main():
    ng = 0
    wb = Workbook(); ws = wb.active
    # A開始 B終了 C休憩分 D勤務(時刻値) E分(整数) F小数時間 G文字表示[h]:mm H15分切捨て(分)
    for r,(s,e,b) in enumerate(CASES,2):
        h,m = s.split(":"); ws.cell(r,1,f"=TIME({h},{m},0)")
        h,m = e.split(":"); ws.cell(r,2,f"=TIME({h},{m},0)")
        ws.cell(r,3,b)
        ws.cell(r,4,f"=MOD(B{r}-A{r},1)-C{r}/1440")
        ws.cell(r,5,f"=ROUND(D{r}*1440,0)")
        ws.cell(r,6,f"=ROUND(D{r}*24,2)")
        ws.cell(r,7,f'=TEXT(D{r},"[h]:mm")')
        ws.cell(r,8,f"=FLOOR(E{r},15)")
    res = run(wb)
    for r,(s,e,b) in enumerate(CASES,2):
        span = (mins(e)-mins(s))%1440   # 同時刻は0(MODの仕様)
        w = span-b
        exp_min = w
        got = [res.cell(r,c).value for c in (5,6,7,8)]
        exp_dec = round(w/60+1e-12,2)
        exp = [exp_min, exp_dec, hm(w) if w>=0 else None, (w//15)*15]
        ok = got[0]==exp[0] and abs(got[1]-exp[1])<1e-9 and (got[2]==exp[2] if w>=0 else True) and got[3]==exp[3]
        ng += not ok
        print(("OK " if ok else "NG ")+f"{s}〜{e} 休憩{b}: 分{got[0]} 小数{got[1]} 表示{got[2]} 15分切捨{got[3]} (期待 {exp})")
    # 合計: [h]:mm と h:mm の違い
    wb = Workbook(); ws = wb.active
    for r,items in enumerate(SUM_CASES,1):
        for c,t in enumerate(items,1):
            h,m=t.split(":"); ws.cell(r,c,f"=TIME({h},{m},0)")
        last = ws.cell(r,len(items)).column_letter
        ws.cell(r,10,f"=SUM(A{r}:{last}{r})")
        ws.cell(r,11,f'=TEXT(J{r},"[h]:mm")'); ws.cell(r,12,f'=TEXT(J{r},"h:mm")')
    res = run(wb)
    for r,items in enumerate(SUM_CASES,1):
        tot = sum(mins(t) for t in items)
        k, l = res.cell(r,11).value, res.cell(r,12).value
        ok = k==hm(tot) and l==hm(tot%1440)
        ng += not ok
        print(("OK " if ok else "NG ")+f"合計{items}: [h]:mm={k!r}(期待{hm(tot)}) h:mm={l!r}(期待{hm(tot%1440)})")
    print("不一致", ng, "件"); sys.exit(1 if ng else 0)
main()
