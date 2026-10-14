"""note原稿用: 10行の小さな明細で、費目別合計・前方一致・空白除去・取りこぼし差額をLibreOfficeで確認する。
使い方: python3 site/verify/category_sum_mini.py"""
import subprocess, sys, tempfile, os
from openpyxl import Workbook, load_workbook
ROWS = [("食費",1200),("食費 ",800),("食費",450),("日用品",980),("食費(外食)",1500),("交通費",320),("日用品　",600),("交通費",540),("",700),("食費",2100)]
CATS = ["食費","日用品","交通費"]
def main():
    wb=Workbook(); ws=wb.active
    for i,(c,v) in enumerate(ROWS,2):
        if c: ws.cell(i,1,c)
        ws.cell(i,2,v)
    last=len(ROWS)+1
    for k,c in enumerate(CATS):
        r=k+2; ws.cell(r,4,c)
        ws.cell(r,5,f'=SUMIFS($B$2:$B${last},$A$2:$A${last},D{r})')
        ws.cell(r,6,f'=SUMIFS($B$2:$B${last},$A$2:$A${last},D{r}&"*")')
        ws.cell(r,7,f'=SUMPRODUCT((TRIM(SUBSTITUTE($A$2:$A${last},"　"," "))=D{r})*$B$2:$B${last})')
    ws["E6"]=f'=SUM(B2:B{last})-SUM(E2:E4)'
    ws["F6"]=f'=SUM(B2:B{last})-SUM(F2:F4)'
    ws["G6"]=f'=SUM(B2:B{last})-SUM(G2:G4)'
    d=tempfile.mkdtemp(); src=os.path.join(d,"t.xlsx"); wb.save(src); out=os.path.join(d,"o"); os.mkdir(out)
    subprocess.run(["soffice","--headless","--convert-to","xlsx","--outdir",out,src],check=True,capture_output=True)
    res=load_workbook(os.path.join(out,"t.xlsx"),data_only=True).active
    ng=0
    tot=sum(v for _,v in ROWS)
    exp={}
    for k,c in enumerate(CATS):
        r=k+2
        e=(sum(v for cc,v in ROWS if cc==c), sum(v for cc,v in ROWS if cc.startswith(c)), sum(v for cc,v in ROWS if cc.replace("　"," ").strip()==c))
        got=tuple(res.cell(r,col).value for col in (5,6,7)); ok=got==e; ng+=not ok
        print(("OK " if ok else "NG ")+f"{c}: 完全一致{got[0]} 前方一致{got[1]} 空白除去{got[2]} (期待{e})")
        exp[c]=e
    for col,idx in ((5,0),(6,1),(7,2)):
        e=tot-sum(exp[c][idx] for c in CATS); g=res.cell(6,col).value; ng+=g!=e
        print(f"取りこぼし差額 列{col}: {g} (期待{e})")
    print("全明細合計",tot,"不一致",ng); sys.exit(1 if ng else 0)
main()
