"""記事『名簿の全角スペースと半角スペースを一括削除』の数式をLibreOfficeで再計算して期待値と照合する。
使い方: python3 site/verify/space_remove.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os
from openpyxl import Workbook, load_workbook

# 入力, TRIM, 半角のみ削除, 半角+全角削除, 全角のみTRIM相当(全角→半角→TRIM)
CASES = [
    ("山田 太郎",        "山田 太郎", "山田太郎", "山田太郎", "山田 太郎"),
    ("山田　太郎",       "山田　太郎", "山田　太郎", "山田太郎", "山田 太郎"),
    ("  山田太郎  ",     "山田太郎", "山田太郎", "山田太郎", "山田太郎"),
    ("山田   太郎",      "山田 太郎", "山田太郎", "山田太郎", "山田 太郎"),
    ("　山田　太郎　",   "　山田　太郎　", "　山田　太郎　", "山田太郎", "山田 太郎"),
    ("山田 　太郎",      "山田 　太郎", "山田　太郎", "山田太郎", "山田 太郎"),
    ("株式会社 サンプル　商事", "株式会社 サンプル　商事", "株式会社サンプル　商事", "株式会社サンプル商事", "株式会社 サンプル 商事"),
    ("",                 "", "", "", ""),
]
F = [
    '=TRIM(A{r})',
    '=SUBSTITUTE(A{r}," ","")',
    '=SUBSTITUTE(SUBSTITUTE(A{r}," ",""),"　","")',
    '=TRIM(SUBSTITUTE(A{r},"　"," "))',
]

def main():
    wb = Workbook(); ws = wb.active
    for r, (inp, *_) in enumerate(CASES, 1):
        ws.cell(r, 1, inp or None); ws.cell(r, 1).number_format = "@"
        for c, f in enumerate(F, 2): ws.cell(r, c, f.format(r=r))
    # 全角スペースをCODE/UNICODEで見分ける
    ws.cell(12, 1, "山田　太郎"); ws.cell(12, 2, "=_xlfn.UNICODE(MID(A12,3,1))")
    ws.cell(13, 1, "山田 太郎"); ws.cell(13, 2, "=_xlfn.UNICODE(MID(A13,3,1))")
    ws.cell(14, 1, "山田　太郎"); ws.cell(14, 2, "=LEN(A14)-LEN(SUBSTITUTE(A14,\"　\",\"\"))")
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    res = load_workbook(os.path.join(out, "t.xlsx"), data_only=True).active
    ng = 0
    for r, (inp, *exp) in enumerate(CASES, 1):
        for c, e in zip(range(2, 6), exp):
            got = res.cell(r, c).value or ""
            ok = got == e; ng += not ok
            print(("OK " if ok else "NG ") + f"{inp!r} col{c}: 期待{e!r} 結果{got!r}")
    for r, e in ((12, 12288), (13, 32), (14, 1)):
        got = res.cell(r, 2).value; ok = got == e; ng += not ok
        print(("OK " if ok else "NG ") + f"row{r}: 期待{e!r} 結果{got!r}")
    print("不一致", ng, "件"); sys.exit(1 if ng else 0)
main()
