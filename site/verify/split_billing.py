"""記事『工事ごとの分割請求の残額をExcelで管理する式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/split_billing.py   (soffice が必要。不一致なら終了コード1)"""
import subprocess, sys, tempfile, os, random, datetime
from openpyxl import Workbook, load_workbook

N_ROWS = 2000  # 式の範囲 2:2000

def build(projects, bills):
    """projects: [(番号, 契約日, 契約金額)], bills: [(請求日, 番号(セルの値そのまま), 請求額)]"""
    wb = Workbook(); led = wb.active; led.title = "台帳"; bl = wb.create_sheet("請求")
    for i, h in enumerate(["工事番号", "契約日", "契約金額(税抜)", "請求済み", "残額", "状態"], 1):
        led.cell(1, i, h)
    for i, h in enumerate(["請求日", "工事番号", "請求額(税抜)", "整えた番号", "契約金額", "チェック"], 1):
        bl.cell(1, i, h)
    for r, (no, d, amt) in enumerate(projects, 2):
        led.cell(r, 1, no); led.cell(r, 2, d); led.cell(r, 3, amt)
        led.cell(r, 4, f'=IF(A{r}="","",SUMIFS(請求!$C$2:$C${N_ROWS},請求!$D$2:$D${N_ROWS},A{r}))')
        led.cell(r, 5, f'=IF(A{r}="","",C{r}-D{r})')
        led.cell(r, 6, f'=IF(A{r}="","",IF(E{r}<0,"請求超過",IF(E{r}=0,"請求完了","")))')
    for r, (d, no, amt) in enumerate(bills, 2):
        bl.cell(r, 1, d); bl.cell(r, 2, no); bl.cell(r, 3, amt)
    for r in range(2, len(bills) + 2):
        bl.cell(r, 4, f'=ASC(TRIM(B{r}))')
        bl.cell(r, 5, f'=IF(D{r}="","",IFERROR(INDEX(台帳!$C$2:$C${N_ROWS},MATCH(D{r},台帳!$A$2:$A${N_ROWS},0)),""))')
        bl.cell(r, 6, (f'=IF(AND(B{r}="",C{r}=""),"",IF(B{r}="","工事番号が空欄",IF(E{r}="","工事番号なし",'
                       f'IF(AND(A{r}<>"",INDEX(台帳!$B$2:$B${N_ROWS},MATCH(D{r},台帳!$A$2:$A${N_ROWS},0))>A{r}),"契約日より前の日付",""))))'))
    led["H1"], led["I1"] = "未紐づけの請求額", f'=SUM(請求!C2:C{N_ROWS})-SUM(D2:D{N_ROWS})'
    return wb

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

def fw(s):  # 半角→全角(英数とハイフン)
    return "".join(chr(ord(c) + 0xFEE0) if "!" <= c <= "~" and c != "-" else ("－" if c == "-" else c) for c in s)

def norm(v):
    if v is None: return ""
    return "".join(chr(ord(c) - 0xFEE0) if "！" <= c <= "～" else ("-" if c == "－" else c) for c in str(v).strip().replace("　", " ").strip())

def expected(projects, bills):
    keys = {norm(no): (d, amt) for no, d, amt in projects}
    billed = {k: 0 for k in keys}; flags = []; unlinked = 0
    for d, no, amt in bills:
        k = norm(no)
        if no in (None, "") and amt in (None, ""): flags.append(""); continue
        if no in (None, ""): flags.append("工事番号が空欄"); unlinked += amt or 0; continue
        if k not in keys: flags.append("工事番号なし"); unlinked += amt or 0; continue
        billed[k] += amt or 0
        flags.append("契約日より前の日付" if d is not None and keys[k][0] > d else "")
    return billed, flags, unlinked

D = datetime.date
SCEN = [  # 名前, projects, bills
 ("分割3回(着手・中間・完了前)", [("K-001", D(2026,10,1), 1000000)],
   [(D(2026,10,1),"K-001",300000),(D(2026,10,20),"K-001",400000),(D(2026,11,5),"K-001",100000)]),
 ("ちょうど全額請求", [("K-002", D(2026,10,1), 500000)], [(D(2026,10,2),"K-002",200000),(D(2026,11,2),"K-002",300000)]),
 ("請求超過", [("K-003", D(2026,10,1), 500000)], [(D(2026,10,2),"K-003",300000),(D(2026,11,2),"K-003",300000)]),
 ("存在しない工事番号", [("K-004", D(2026,10,1), 800000)], [(D(2026,10,2),"K-040",100000)]),
 ("工事番号が空欄", [("K-005", D(2026,10,1), 800000)], [(D(2026,10,2),None,100000)]),
 ("全角の工事番号", [("K-006", D(2026,10,1), 800000)], [(D(2026,10,2),"Ｋ－006",100000),(D(2026,10,3)," K-006 ",50000)]),
 ("日付の逆転", [("K-007", D(2026,10,10), 800000)], [(D(2026,10,2),"K-007",100000)]),
 ("請求なし", [("K-008", D(2026,10,1), 800000)], []),
]

def main():
    ng = 0
    for name, projects, bills in SCEN:
        res = run_lo(build(projects, bills))
        led, bl = res["台帳"], res["請求"]
        billed, flags, unlinked = expected(projects, bills)
        got_billed = {norm(led.cell(r, 1).value): (led.cell(r, 4).value or 0) for r in range(2, len(projects) + 2)}
        got_flags = [(bl.cell(r, 6).value or "") for r in range(2, len(bills) + 2)]
        got_unl = led["I1"].value
        ok = got_billed == billed and got_flags == flags and got_unl == unlinked
        ng += not ok
        rem = [(norm(led.cell(r,1).value), led.cell(r,5).value, led.cell(r,6).value) for r in range(2, len(projects)+2)]
        print(("OK " if ok else "NG ") + f"{name}: 請求済み{got_billed} 残額/状態{rem} チェック{got_flags} 未紐づけ{got_unl}")
        if not ok: print("   期待", billed, flags, unlinked)
    # --- ランダム照合: 工事300件・請求1000行(誤入力を混ぜる) ---
    rnd = random.Random(20261019)
    projects = [(f"K-{i:03d}", D(2026,10,1) + datetime.timedelta(days=rnd.randint(0, 60)), rnd.randint(1, 500) * 10000) for i in range(1, 301)]
    bills = []
    for _ in range(1000):
        t = rnd.random(); p = rnd.choice(projects)
        d = p[1] + datetime.timedelta(days=rnd.randint(-5, 90)); amt = rnd.randint(1, 100) * 10000
        if t < 0.80: bills.append((d, p[0], amt))
        elif t < 0.86: bills.append((d, "K-9%02d" % rnd.randint(0, 99), amt))
        elif t < 0.91: bills.append((d, None, amt))
        elif t < 0.96: bills.append((d, fw(p[0]), amt))
        else: bills.append((d, " " + p[0] + " ", amt))
    res = run_lo(build(projects, bills)); led, bl = res["台帳"], res["請求"]
    billed, flags, unlinked = expected(projects, bills)
    gb = {norm(led.cell(r,1).value): (led.cell(r,4).value or 0) for r in range(2, 302)}
    gf = [(bl.cell(r, 6).value or "") for r in range(2, 1002)]
    ok_rem = all(led.cell(r,5).value == projects[r-2][2] - billed[projects[r-2][0]] for r in range(2, 302))
    ok = gb == billed and gf == flags and led["I1"].value == unlinked and ok_rem
    ng += not ok
    print(("OK " if ok else "NG ") + f"ランダム照合: 工事300件・請求1000行 請求済み・残額・チェック・未紐づけ額が期待値と一致={ok}")
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
