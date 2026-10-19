"""記事『見積の最新版の金額を工事台帳に引いて転記ズレを見つける式』をLibreOfficeで再計算して期待値(Python)と照合する。
使い方: python3 site/verify/quote_to_ledger.py   (soffice が必要。不一致なら終了コード1)
MAXIFS は Excel 2019 / Microsoft 365 以降の関数(ファイルでは _xlfn.MAXIFS)。"""
import subprocess, sys, tempfile, os, random
from openpyxl import Workbook, load_workbook

N = 2000

def build(projects, quotes):
    """projects: [(工事番号, 転記済みの金額 or None)], quotes: [(工事番号, 版, 税抜合計)]"""
    wb = Workbook(); led = wb.active; led.title = "台帳"; q = wb.create_sheet("見積")
    for i, h in enumerate(["工事番号", "最新版", "見積金額(税抜)", "転記済みの金額", "状態", "差額"], 1):
        led.cell(1, i, h)
    for i, h in enumerate(["工事番号", "版", "税抜合計", "整えた番号", "チェック", "有効", "最新版"], 1):
        q.cell(1, i, h)
    for r, (no, tr) in enumerate(projects, 2):
        led.cell(r, 1, no); led.cell(r, 4, tr)
        led.cell(r, 2, f'=IF(A{r}="","",IF(COUNTIFS(見積!$D$2:$D${N},A{r},見積!$F$2:$F${N},1)=0,"見積なし",_xlfn.MAXIFS(見積!$B$2:$B${N},見積!$D$2:$D${N},A{r},見積!$F$2:$F${N},1)))')
        led.cell(r, 3, f'=IF(A{r}="","",IF(B{r}="見積なし","",IF(COUNTIFS(見積!$D$2:$D${N},A{r},見積!$G$2:$G${N},1)>1,"同じ版が重複",SUMIFS(見積!$C$2:$C${N},見積!$D$2:$D${N},A{r},見積!$G$2:$G${N},1))))')
        led.cell(r, 5, (f'=IF(A{r}="","",IF(NOT(ISNUMBER(C{r})),"見積を確認",IF(D{r}="","未転記",IF(NOT(ISNUMBER(D{r})),"転記額が数値でない",IF(D{r}=C{r},"一致","ズレ")))))'))
        led.cell(r, 6, f'=IF(E{r}="ズレ",D{r}-C{r},"")')
    for r, (no, ver, amt) in enumerate(quotes, 2):
        q.cell(r, 1, no); q.cell(r, 2, ver); q.cell(r, 3, amt)
        q.cell(r, 4, f'=ASC(TRIM(A{r}))')
        q.cell(r, 5, (f'=IF(AND(A{r}="",B{r}="",C{r}=""),"",IF(A{r}="","工事番号が空欄",IF(ISNA(MATCH(D{r},台帳!$A$2:$A${N},0)),"工事番号なし",'
                      f'IF(NOT(ISNUMBER(B{r})),"版が1以上の整数でない",IF(OR(B{r}<1,B{r}<>INT(B{r})),"版が1以上の整数でない",IF(NOT(ISNUMBER(C{r})),"金額が数値でない",IF(C{r}<0,"金額がマイナス","")))))))'))
        q.cell(r, 6, f'=IF(AND(A{r}="",B{r}="",C{r}=""),0,IF(E{r}="",1,0))')
        q.cell(r, 7, f'=IF(F{r}=1,IF(B{r}=_xlfn.MAXIFS($B$2:$B${N},$D$2:$D${N},D{r},$F$2:$F${N},1),1,0),0)')
    q["I1"] = "要確認の行数"; q["J1"] = f'=SUMPRODUCT(--(E2:E{N}<>""))'
    return wb

def run_lo(wb):
    d = tempfile.mkdtemp(); src = os.path.join(d, "t.xlsx"); wb.save(src)
    out = os.path.join(d, "out"); os.mkdir(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return load_workbook(os.path.join(out, "t.xlsx"), data_only=True)

def norm(v):
    if v is None: return ""
    return "".join(chr(ord(c) - 0xFEE0) if "！" <= c <= "～" else ("-" if c == "－" else c) for c in str(v).strip().replace("　", " ").strip())

def fw(s):
    return "".join(chr(ord(c) + 0xFEE0) if "!" <= c <= "~" and c != "-" else ("－" if c == "-" else c) for c in s)

def isn(v): return isinstance(v, (int, float)) and not isinstance(v, bool)

def expected(projects, quotes):
    keys = {norm(p) for p, _ in projects}
    qrows = []; valid = {k: [] for k in keys}; bad = 0
    for no, ver, amt in quotes:
        if no in (None, "") and ver is None and amt is None: qrows.append(""); continue
        if no in (None, ""): chk = "工事番号が空欄"
        elif norm(no) not in keys: chk = "工事番号なし"
        elif not isn(ver) or ver < 1 or ver != int(ver): chk = "版が1以上の整数でない"
        elif not isn(amt): chk = "金額が数値でない"
        elif amt < 0: chk = "金額がマイナス"
        else: chk = ""
        qrows.append(chk)
        if chk: bad += 1
        else: valid[norm(no)].append((ver, amt))
    lrows = []
    for no, tr in projects:
        if no in (None, ""): lrows.append(("", "", "", "")); continue
        v = valid[norm(no)]
        if not v: lrows.append(("見積なし", "", "見積を確認", "")); continue
        mx = max(x for x, _ in v); top = [a for x, a in v if x == mx]
        if len(top) > 1: lrows.append((mx, "同じ版が重複", "見積を確認", "")); continue
        c = top[0]
        if tr is None: st, df = "未転記", ""
        elif not isn(tr): st, df = "転記額が数値でない", ""
        elif tr == c: st, df = "一致", ""
        else: st, df = "ズレ", tr - c
        lrows.append((mx, c, st, df))
    return qrows, lrows, bad

def cmp(projects, quotes):
    res = run_lo(build(projects, quotes)); led, q = res["台帳"], res["見積"]
    eq, el, bad = expected(projects, quotes)
    gq = [(q.cell(r, 5).value or "") for r in range(2, len(quotes) + 2)]
    gl = [tuple(("" if led.cell(r, c).value is None else led.cell(r, c).value) for c in (2, 3, 5, 6)) for r in range(2, len(projects) + 2)]
    ok = gq == eq and gl == [tuple(x) for x in el] and q["J1"].value == bad
    return ok, (gq, gl, q["J1"].value), (eq, el, bad)

SCEN = [
 ("見積が1版だけ・転記額が一致", [("K-001", 1000000)], [("K-001", 1, 1000000)]),
 ("見積を改訂(版2)・古い版の金額が転記されたまま", [("K-002", 1000000)], [("K-002", 1, 1000000), ("K-002", 2, 1250000)]),
 ("版が行の順と逆(版2が先の行)", [("K-003", 1250000)], [("K-003", 2, 1250000), ("K-003", 1, 1000000)]),
 ("未転記(転記額が空欄)", [("K-004", None)], [("K-004", 1, 800000)]),
 ("見積が1行もない工事", [("K-005", 500000)], []),
 ("同じ版を2行入力(重複)", [("K-006", 900000)], [("K-006", 1, 900000), ("K-006", 1, 900000)]),
 ("存在しない工事番号の見積", [("K-007", None)], [("K-070", 1, 300000)]),
 ("工事番号が空欄の見積", [("K-008", None)], [(None, 1, 300000)]),
 ("金額欄に文字('300,000円'と入力)", [("K-009", None)], [("K-009", 1, "300,000円")]),
 ("版が小数・ゼロ・文字", [("K-010", None)], [("K-010", 1.5, 100), ("K-010", 0, 100), ("K-010", "第2版", 100)]),
 ("金額がマイナス(値引きを別行で入れた)", [("K-011", None)], [("K-011", 1, 500000), ("K-011", 2, -50000)]),
 ("全角・前後スペースの番号", [("K-012", 700000)], [("Ｋ－012", 1, 600000), (" K-012 ", 2, 700000)]),
 ("転記額が文字('120万')", [("K-013", "120万")], [("K-013", 1, 1200000)]),
 ("まったくの空行", [("K-014", None)], [(None, None, None)]),
]

def main():
    ng = 0
    for name, projects, quotes in SCEN:
        ok, got, exp = cmp(projects, quotes)
        ng += not ok
        print(("OK " if ok else "NG ") + f"{name}: 台帳{got[1]} 見積チェック{got[0]} 要確認{got[2]}")
        if not ok: print("   期待", exp[1], exp[0], exp[2])
    rnd = random.Random(20261020)
    names = [f"K-{i:03d}" for i in range(1, 301)]
    projects = []
    for p in names:
        t = rnd.random()
        projects.append((p, None if t < .1 else ("要確認" if t < .13 else rnd.randint(100, 3000) * 1000)))
    quotes = []
    for _ in range(1000):
        p = rnd.choice(names); ver = rnd.randint(1, 4); amt = rnd.randint(100, 3000) * 1000; t = rnd.random()
        if t < .8: quotes.append((p, ver, amt))
        elif t < .83: quotes.append(("K-9%02d" % rnd.randint(0, 99), ver, amt))
        elif t < .86: quotes.append((None, ver, amt))
        elif t < .89: quotes.append((fw(p), ver, amt))
        elif t < .92: quotes.append((" " + p + " ", ver, amt))
        elif t < .94: quotes.append((p, ver, "約" + str(amt)))
        elif t < .96: quotes.append((p, ver + .5, amt))
        elif t < .98: quotes.append((p, ver, -amt))
        else: quotes.append((p, 0, amt))
    # 転記額: 約半数を最新版の金額に合わせる(残りは古い版・別の金額のまま)
    _, el, _ = expected(projects, quotes)
    projects = [(p, (el[i][1] if (isinstance(el[i][1], int) and rnd.random() < .5 and tr is not None and isinstance(tr, int)) else tr)) for i, (p, tr) in enumerate(projects)]
    ok, got, exp = cmp(projects, quotes)
    ng += not ok
    st = {}
    for r in exp[1]: st[r[2]] = st.get(r[2], 0) + 1
    print(("OK " if ok else "NG ") + f"ランダム照合: 工事300件・見積1,000行 一致={ok} 要確認{exp[2]}行 台帳の状態{st}")
    print("不一致あり" if ng else "全件一致")
    sys.exit(1 if ng else 0)

if __name__ == "__main__":
    main()
